"""Render a reply draft as an editable Word document. Placeholders are highlighted in yellow."""
from __future__ import annotations

import re
from datetime import date
from pathlib import Path

from .content import ReplyDraft, ph

_PLACEHOLDER = re.compile(r"(\[\[.*?\]\])")


def _runs(paragraph, text: str, bold: bool = False) -> None:
    from docx.enum.text import WD_COLOR_INDEX

    for part in _PLACEHOLDER.split(text):
        if not part:
            continue
        run = paragraph.add_run(part[2:-2] if part.startswith("[[") else part)
        run.bold = bold
        if part.startswith("[["):
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW


def render_reply(draft: ReplyDraft, path: Path, client_name: str, client_gstin: str, client_address: str,
                 officer: str, letter_date: date, signatory: str = "") -> Path:
    from docx import Document
    from docx.shared import Mm, Pt

    doc = Document()
    sec = doc.sections[0]
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Mm(20))
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    head = doc.add_paragraph()
    head.add_run(client_name).bold = True
    head.add_run(f"\nGSTIN: {client_gstin}")
    if client_address:
        head.add_run(f"\n{client_address}")
    doc.add_paragraph(f"Date: {letter_date:%d-%m-%Y}")
    to = doc.add_paragraph("To\nThe Proper Officer\n")
    _runs(to, officer or ph("designation, range / division and address of the officer"))
    subj = doc.add_paragraph()
    subj.add_run("Subject: ").bold = True
    _runs(subj, draft.subject, bold=True)
    doc.add_paragraph("Sir / Madam,")
    for s in draft.sections:
        if s.heading:
            doc.add_heading(s.heading, level=2)
        for text in s.paragraphs:
            _runs(doc.add_paragraph(), text)
        if s.table:
            headers, rows = s.table
            t = doc.add_table(rows=1, cols=len(headers))
            t.style = "Table Grid"
            for i, h in enumerate(headers):
                t.rows[0].cells[i].text = h
                t.rows[0].cells[i].paragraphs[0].runs[0].bold = True
            for r in rows:
                cells = t.add_row().cells
                for i, v in enumerate(r):
                    cells[i].text = str(v)
            doc.add_paragraph()
    close = doc.add_paragraph("Yours faithfully,\n")
    close.add_run(f"For {client_name}\n\n\n")
    _runs(close, (signatory or ph("name of the authorised signatory")) + "\nAuthorised signatory")
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(".tmp.docx")
    doc.save(tmp)
    tmp.replace(path)
    return path
