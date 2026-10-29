from decimal import Decimal

from openpyxl import load_workbook

from fintax_gst.cli import main
from fintax_gst.core.models import Category, Source
from fintax_gst.engine.summary import dashboard
from fintax_gst.reports.excel_report import COL, LINE_COLUMNS, SHEET_NAMES
from fintax_gst.samples import CLIENT, SUPPLIERS
from fintax_gst.services.clients import ClientService
from fintax_gst.services.imports import ImportService
from fintax_gst.services.reconciliation import ReconciliationService
from fintax_gst.services.reports import export_run


def _run(db, files):
    c = ClientService(db).create(**CLIENT)
    svc = ImportService(db)
    for key, src in (("books_march", Source.BOOKS), ("books_april", Source.BOOKS), ("gstr2b_april", Source.GSTR2B)):
        svc.import_file(c.id, files[key], src)
    return ReconciliationService(db).run(c.id, "2025-04")


def _col_values(ws, header):
    idx = [n for n, *_ in LINE_COLUMNS].index(header)
    out = []
    for row in ws.iter_rows(min_row=5, values_only=True):
        if row[0] is None or row[0] == "Total (visible rows)" or not isinstance(row[0], int):
            continue
        out.append(row[idx])
    return out


def test_excel_report_structure_and_totals(db, sample_files, tmp_path):
    run_id = _run(db, sample_files)
    xlsx, pdf = export_run(db, run_id, tmp_path)
    assert xlsx.name == "FINTAXPRO_GSTRecon_27AAPFU0939F1ZV_FY2025-26_2025-04_to_2025-04_run1.xlsx"
    assert pdf.exists() and pdf.stat().st_size > 1000
    wb = load_workbook(xlsx)
    assert wb.sheetnames == SHEET_NAMES and len(SHEET_NAMES) == 17
    ws = wb["Complete Reconciliation"]
    assert ws.freeze_panes == "C5"
    assert ws.tables, "detail sheets must be Excel tables"
    lines = ReconciliationService(db).load(run_id).result.lines
    d = dashboard(lines)
    assert len(_col_values(ws, "Line No")) == len(lines)
    scope = _col_values(ws, "Books In Scope")
    books_itc = _col_values(ws, "Books ITC")
    g2b_itc = _col_values(ws, "GSTR-2B ITC")
    assert sum(Decimal(str(v)) for v, s in zip(books_itc, scope) if s == "Yes") == d["itc_as_per_books"]
    assert sum(Decimal(str(v)) for v in g2b_itc if v is not None) == d["itc_as_per_gstr2b"]
    # Each category sheet holds exactly the lines of its categories.
    missing = wb["Missing in GSTR-2B"]
    assert len(_col_values(missing, "Line No")) == sum(1 for ln in lines if ln.category == Category.MISSING_IN_2B)
    # Totals row uses SUBTOTAL so it follows filters.
    last = ws.max_row
    assert ws[f"{COL['Books ITC']}{last}"].value.startswith("=SUBTOTAL(109,")
    # Pending register offers the status list.
    assert wb["Pending Action Register"].data_validations.dataValidation
    summary_text = [c for row in wb["Executive Summary"].iter_rows(values_only=True) for c in row if c]
    assert "Potential ITC at Risk" in summary_text


def test_empty_category_sheets_still_valid_tables(db, tmp_path, sample_files):
    c = ClientService(db).create(**CLIENT)
    f = tmp_path / "one.csv"
    f.write_text(f"GSTIN,Invoice No,Invoice Date,Taxable Value,IGST\n{SUPPLIERS['BEL'][0]},A-1,01-04-2025,100,18\n")
    ImportService(db).import_file(c.id, f, Source.BOOKS)
    run_id = ReconciliationService(db).run(c.id, "2025-04")
    xlsx, = export_run(db, run_id, tmp_path, pdf=False)
    wb = load_workbook(xlsx)
    assert wb["Fully Matched Invoices"]["A5"].value == "No records"


def test_cli_end_to_end(tmp_path, sample_files, capsys, monkeypatch):
    monkeypatch.setenv("FINTAX_HOME", str(tmp_path / "home"))
    db = str(tmp_path / "cli.db")
    assert main(["--db", db, "client", "add", "--gstin", CLIENT["gstin"], "--name", CLIENT["name"],
                 "--fy", "2025-26"]) == 0
    assert main(["--db", db, "import", "--client", CLIENT["gstin"], "--source", "books",
                 "--file", str(sample_files["books_april"]), "--preview"]) == 0
    assert "Preview only" in capsys.readouterr().out
    assert main(["--db", db, "import", "--client", "1", "--source", "books",
                 "--file", str(sample_files["books_april"])]) == 0
    assert main(["--db", db, "import", "--client", "1", "--source", "books",
                 "--file", str(sample_files["books_april"])]) == 1  # repeated import refused
    assert main(["--db", db, "import", "--client", "1", "--source", "gstr2b",
                 "--file", str(sample_files["gstr2b_april_json"])]) == 0
    assert main(["--db", db, "reconcile", "--client", "1", "--from", "2025-04", "--date-tolerance", "3",
                 "--export", str(tmp_path / "out"), "--no-pdf"]) == 0
    out = capsys.readouterr().out
    assert "Potential ITC at Risk" in out and "Report:" in out
    assert main(["--db", db, "update-line", "--run", "1", "--line", "2", "--status", "Resolved",
                 "--remarks", "checked"]) == 0
    assert main(["--db", db, "show", "--run", "1", "--category", "missing"]) == 0
    assert main(["--db", db, "backup", "--to", str(tmp_path / "bk.db")]) == 0
    assert (tmp_path / "bk.db").exists()
