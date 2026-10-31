"""Phase 12: direct Tally import."""
import importlib.util
from datetime import date
from pathlib import Path

import pytest

from fintax_gst.cli import main
from fintax_gst.core.models import DocType, Source
from fintax_gst.importers.pipeline import preview_file
from fintax_gst.importers.readers import UnsupportedFileError
from fintax_gst.importers.tally_xml import TallyClient, flatten_vouchers, parse_xml
from fintax_gst.samples import CLIENT
from fintax_gst.services.clients import ClientService
from fintax_gst.services.common import ServiceError
from fintax_gst.services.imports import ImportService
from fintax_gst.services.tally import TallyService

ROOT = Path(__file__).resolve().parents[1]


@pytest.fixture
def tally(sample_files):
    spec = importlib.util.spec_from_file_location("tally_simulator", ROOT / "scripts" / "tally_simulator.py")
    sim = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(sim)
    server = sim.serve(sample_files["books_tally_xml"], port=0, company=CLIENT["name"], background=True)
    yield f"http://127.0.0.1:{server.server_address[1]}"
    server.shutdown()


def _key(r):
    return (r.supplier_gstin, r.invoice_no_norm, r.doc_type, r.invoice_date, r.taxable, r.igst, r.cgst, r.sgst,
            r.period)


def test_tally_xml_matches_the_excel_register(sample_files):
    xml = preview_file(sample_files["books_tally_xml"], Source.BOOKS)
    xlsx = preview_file(sample_files["books_april"], Source.BOOKS)
    assert xml.layout == "tally_xml"
    assert {_key(r) for r in xml.records} == {_key(r) for r in xlsx.records}
    ignored = {i.row: i.message for i in xml.issues if i.severity == "ignored"}
    assert ignored == {"Purchase 99": "cancelled voucher",
                       "Payment 7": "voucher type 'Payment' is not a purchase / sales / note voucher"}
    # a purchase return (debit note voucher) is the supplier's credit note
    assert any(r.doc_type == DocType.CREDIT_NOTE for r in xml.records)
    with pytest.raises(UnsupportedFileError, match="purchase or sales register"):
        preview_file(sample_files["books_tally_xml"], Source.GSTR2B)


def test_tally_xml_details():
    xml = b"""<?xml version="1.0"?><ENVELOPE><BODY><DATA>
      <VOUCHER VCHTYPE="GST Sales"><DATE>20250405</DATE><VOUCHERTYPENAME>GST Sales</VOUCHERTYPENAME>
        <VOUCHERNUMBER>S/1</VOUCHERNUMBER><PARTYLEDGERNAME>Asha &amp; Co</PARTYLEDGERNAME>
        <PARTYGSTIN>27AABCR1122A1Z7</PARTYGSTIN>
        <ALLLEDGERENTRIES.LIST><LEDGERNAME>Asha &amp; Co</LEDGERNAME><ISPARTYLEDGER>Yes</ISPARTYLEDGER>
          <AMOUNT>-11800.40</AMOUNT></ALLLEDGERENTRIES.LIST>
        <ALLLEDGERENTRIES.LIST><LEDGERNAME>Output CGST</LEDGERNAME><AMOUNT>900.00</AMOUNT></ALLLEDGERENTRIES.LIST>
        <ALLLEDGERENTRIES.LIST><LEDGERNAME>Output SGST</LEDGERNAME><AMOUNT>900.00</AMOUNT></ALLLEDGERENTRIES.LIST>
        <ALLLEDGERENTRIES.LIST><LEDGERNAME>Rounding Off</LEDGERNAME><AMOUNT>0.40</AMOUNT></ALLLEDGERENTRIES.LIST>
        <ALLLEDGERENTRIES.LIST><LEDGERNAME>State Bank Charges</LEDGERNAME><AMOUNT>0</AMOUNT></ALLLEDGERENTRIES.LIST>
        <ALLINVENTORYENTRIES.LIST><STOCKITEMNAME>Steel</STOCKITEMNAME><AMOUNT>10000.00</AMOUNT>
          <GSTHSNNAME>7208</GSTHSNNAME><ACTUALQTY>10 MT</ACTUALQTY>
          <ACCOUNTINGALLOCATIONS.LIST><LEDGERNAME>Sales @18%</LEDGERNAME><AMOUNT>10000.00</AMOUNT>
          </ACCOUNTINGALLOCATIONS.LIST></ALLINVENTORYENTRIES.LIST>
      </VOUCHER>
      <VOUCHER VCHTYPE="Purchase"><DATE>20250406</DATE><VOUCHERTYPENAME>Purchase</VOUCHERTYPENAME>
        <VOUCHERNUMBER>P/9</VOUCHERNUMBER><ISOPTIONAL>Yes</ISOPTIONAL></VOUCHER>
    </DATA></BODY></ENVELOPE>"""
    found = flatten_vouchers(parse_xml(xml), "sales")
    (row,) = found.rows
    assert row["invoice_no"] == "S/1" and row["supplier_name"] == "Asha & Co" and row["invoice_date"] == date(2025, 4, 5)
    assert (row["taxable_value"], row["cgst"], row["sgst"], row["igst"]) == ("10000.00", "900.00", "900.00", "0")
    assert row["invoice_value"] == "11800.40" and row["hsn"] == "7208"  # round-off in the value, not the taxable
    assert found.ignored == []  # the optional purchase belongs to the other register
    purchases = flatten_vouchers(parse_xml(xml), "purchase")
    assert purchases.ignored == [("Purchase P/9", "optional (memorandum) voucher")]


def test_tally_xml_refuses_entities():
    with pytest.raises(UnsupportedFileError):
        parse_xml(b'<?xml version="1.0"?><!DOCTYPE x [<!ENTITY a "aaaa">]><ENVELOPE>&a;</ENVELOPE>')
    # control-character references Tally sometimes emits are dropped, not fatal
    assert parse_xml(b"<ENVELOPE><X>a&#4;b</X></ENVELOPE>").findtext("X") == "ab"


def test_live_fetch_and_import(db, tally, sample_files):
    c = ClientService(db).create(**CLIENT)
    client = TallyClient(tally)
    assert client.companies() == [CLIENT["name"]]
    assert client.duty_heads(CLIENT["name"])["Input IGST 18%"] == "igst"
    with pytest.raises(ValueError, match="Could not set"):
        client.vouchers("Another Company", date(2025, 4, 1), date(2025, 4, 30))
    path = TallyService(db).fetch(c.id, CLIENT["name"], "2025-04", url=tally, user="preparer")
    assert path.parent == db.data_dir / "tally" and path.read_bytes().startswith(b"<?xml")
    sf, preview = ImportService(db).import_file(c.id, path, Source.BOOKS, user="preparer")
    assert preview.layout == "tally_xml" and sf.documents_imported == 28
    # a month with no vouchers
    empty = TallyService(db).fetch(c.id, CLIENT["name"], "2025-06", url=tally)
    with pytest.raises(UnsupportedFileError, match="no vouchers"):
        preview_file(empty, Source.BOOKS)


def test_unreachable_tally_explains_the_setting():
    with pytest.raises(ConnectionError, match="Connectivity"):
        TallyClient("http://127.0.0.1:9", timeout=2).companies()


def test_cli_tally(tmp_path, tally, capsys):
    dbp = str(tmp_path / "cli.db")
    assert main(["--db", dbp, "client", "add", "--name", CLIENT["name"], "--gstin", CLIENT["gstin"]]) == 0
    assert main(["--db", dbp, "tally", "companies", "--url", tally]) == 0
    assert CLIENT["name"] in capsys.readouterr().out
    assert main(["--db", dbp, "tally", "import", "--url", tally, "--client", "1", "--company", CLIENT["name"],
                 "--from", "2025-04"]) == 0
    out = capsys.readouterr().out
    assert "Fetched from Tally" in out and "28 documents saved" in out


# ------------------------------------------------------------------ desktop build: no numpy / pandas
def test_program_works_without_numpy_or_pandas(tmp_path, sample_files):
    """numpy failed to load on an office PC, so the desktop build contains neither numpy nor pandas.
    With both blocked (as if not installed), the program must import, reconcile and write its reports."""
    import os
    import subprocess
    import sys

    code = f"""
import sys
class Block:
    def find_spec(self, name, path=None, target=None):
        if name.split(".")[0] in ("numpy", "pandas"):
            raise ImportError(name + " is blocked in this test")
sys.meta_path.insert(0, Block())
import fintax_gst.cli, fintax_gst.ui.app
from fintax_gst.core.models import Source
from fintax_gst.db.database import open_database
from fintax_gst.samples import CLIENT
from fintax_gst.services.clients import ClientService
from fintax_gst.services.common import ServiceError
from fintax_gst.services.imports import ImportService
from fintax_gst.services.reconciliation import ReconciliationService
from fintax_gst.services.reports import export_run
db = open_database(r"{tmp_path / 'nonumpy.db'}")
c = ClientService(db).create(**CLIENT)
svc = ImportService(db)
svc.import_file(c.id, r"{sample_files['books_april']}", Source.BOOKS)
svc.import_file(c.id, r"{sample_files['books_april_csv']}", Source.BOOKS)
svc.import_file(c.id, r"{sample_files['gstr2b_april_json']}", Source.GSTR2B)
run = ReconciliationService(db).run(c.id, "2025-04")
paths = export_run(db, run, r"{tmp_path}", pdf=True)
assert not [m for m in ("numpy", "pandas") if m in sys.modules]
print("OK", len(paths))
"""
    out = subprocess.run([sys.executable, "-c", code], capture_output=True, text=True,
                         env={**os.environ, "QT_QPA_PLATFORM": "offscreen"}, timeout=300)
    assert out.returncode == 0 and "OK 2" in out.stdout, out.stdout + out.stderr


def test_csv_reader_without_pandas(tmp_path):
    from fintax_gst.importers.readers import read_csv

    p = tmp_path / "semi.csv"
    p.write_bytes("GSTIN;Invoice No;Taxable\n27AAPFU0939F1ZV;INV/1;\"1,23,456.50\"\n\n;;\n".encode("cp1252"))
    t = read_csv(p, p.name)
    assert t.rows[0] == (1, ["GSTIN", "Invoice No", "Taxable"])
    assert t.rows[1] == (2, ["27AAPFU0939F1ZV", "INV/1", "1,23,456.50"])  # text kept exactly, no float
    assert t.rows[2] == (3, []) and t.rows[3] == (4, [])
    q = tmp_path / "comma.csv"
    q.write_text('﻿Party,Bill No,Amount\n"Sharma & Sons, Pune",00123,1180\n', encoding="utf-8")
    t = read_csv(q, q.name)
    assert t.rows[1][1] == ["Sharma & Sons, Pune", "00123", "1180"]  # leading zeros kept
    one = tmp_path / "one.csv"
    one.write_text("Only\nvalue\n", encoding="utf-8")
    assert read_csv(one, one.name).rows == [(1, ["Only"]), (2, ["value"])]


# ------------------------------------------------------------------ register templates for users
@pytest.mark.parametrize("name, source, count", [("EXAMPLE_purchase_register", Source.BOOKS, 4),
                                                 ("EXAMPLE_sales_register", Source.SALES, 5),
                                                 ("TEMPLATE_purchase_register", Source.BOOKS, 0),
                                                 ("TEMPLATE_sales_register", Source.SALES, 0)])
def test_register_templates_import_without_mapping(name, source, count):
    p = preview_file(ROOT / "sample_data" / f"{name}.xlsx", source)
    assert p.layout == "tally" and len(p.records) == count
    assert not [i for i in p.issues if i.severity in ("error", "rejected")]
    if name == "EXAMPLE_purchase_register":
        rcm = [r for r in p.records if r.invoice_no == "SR/77"]
        assert rcm[0].reverse_charge and any(r.doc_type == DocType.CREDIT_NOTE for r in p.records)


# ------------------------------------------------------------------ batch processing (Phase 12B)
@pytest.mark.parametrize("name, kind", [
    ("gstr2b_2025-04.json", "GSTR2B"), ("GSTR-2B Apr.xlsx", "GSTR2B"), ("2B.json", "GSTR2B"), ("gstr2a.xlsx", "GSTR2A"),
    ("table8a.json", "GSTR8A"), ("GSTR3B.json", "GSTR3B"), ("gstr1.json", "GSTR1"), ("GSTR-1 April.json", "GSTR1"),
    ("R1.json", "GSTR1"), ("ims.json", "IMS"), ("einvoice_report.csv", "EINV"), ("e-way bills.csv", "EWB"),
    ("TDS credit.xlsx", "TDSTCS"), ("tally_sales_register.xlsx", "SALES"), ("Purchase Register.xlsx", "BOOKS"),
    ("27AAPFU0939F1ZV purchase.xlsx", "BOOKS"), ("claims.xlsx", None), ("notes.xlsx", None), ("a2b3.xlsx", None)])
def test_batch_file_kinds(name, kind):
    from fintax_gst.services.batch import kind_of

    assert kind_of(name) == kind


def test_batch_folder(db, tmp_path):
    from fintax_gst.services.batch import BatchService, write_example

    root = write_example(tmp_path / "batch")
    month = root / f"{CLIENT['name']} {CLIENT['gstin']}" / "2025-04"
    (month / "notes.xlsx").write_bytes(b"x")  # name says nothing about the contents
    (month / "gstr2b_wrong.xlsx").write_bytes((month / "purchase_register.xlsx").read_bytes())  # contents differ
    (month.parent / "loose purchase.xlsx").write_bytes(b"x")  # not in a month folder
    (root / "No GSTIN here" / "2025-04").mkdir(parents=True)
    (root / "No GSTIN here" / "2025-04" / "purchase.xlsx").write_bytes(b"x")
    svc = BatchService(db)

    check = svc.run(root, dry_run=True)  # nothing changes, the missing client is reported
    assert check.summary_path is None and check.count("planned") == 0
    assert any("add the client first" in r.detail for r in check.rows if r.action == "error")
    check = svc.run(root, dry_run=True, create_clients=True)  # the client and 8 files would be done
    assert check.count("planned") == 9 and "client would be added" in [r.detail for r in check.rows]
    with pytest.raises(ServiceError):
        ClientService(db).by_gstin(CLIENT["gstin"])

    res = svc.run(root, user="preparer", create_clients=True)
    by = {(r.item, r.action): r.detail for r in res.rows}
    assert res.count("client") == 1 and res.count("imported") == 7 and res.count("reconciled") == 4
    assert "cannot tell what the file is" in by[("notes.xlsx", "skipped")]
    assert "not in a month folder" in by[("loose purchase.xlsx", "skipped")]
    assert ("gstr2b_wrong.xlsx", "error") in by  # the importer refuses a purchase register as GSTR-2B
    assert any("no valid GSTIN" in r.detail for r in res.rows if r.client == "No GSTIN here")
    reports = sorted(p.name for p in (month / "reports").iterdir())
    assert len(reports) == 4 and reports[0].startswith("Purchase vs GSTR-2B – FINTAXPRO_GSTRecon_")
    assert res.summary_path.exists() and "imported" in res.summary_path.read_text(encoding="utf-8-sig")
    assert ClientService(db).by_gstin(CLIENT["gstin"]).name == CLIENT["name"]

    again = svc.run(root, user="preparer")  # nothing new: nothing imported, nothing re-run
    assert again.count("imported") == 0 and again.count("reconciled") == 0 and again.count("skipped") >= 7
    assert svc.run(root, rerun=True).count("reconciled") == 4


def test_batch_cancel_keeps_a_summary(db, tmp_path):
    from fintax_gst.engine.matcher import CancelToken
    from fintax_gst.services.batch import BatchService, write_example

    root = write_example(tmp_path / "batch")
    token = CancelToken()
    seen = []

    def progress(pct, msg):
        seen.append(msg)
        if len(seen) == 3:
            token.cancel()

    res = BatchService(db).run(root, create_clients=True, progress=progress, cancel=token)
    assert res.cancelled and res.count("imported") == 3 and res.count("reconciled") == 0
    assert "cancelled" in res.summary_path.read_text(encoding="utf-8-sig")


def test_cli_batch(tmp_path, capsys):
    dbp = str(tmp_path / "cli.db")
    folder = str(tmp_path / "batch")
    assert main(["--db", dbp, "batch", folder, "--example"]) == 0
    assert main(["--db", dbp, "batch", folder, "--check"]) == 0
    assert "Check only" in capsys.readouterr().out
    assert main(["--db", dbp, "batch", folder, "--add-clients"]) == 0
    out = capsys.readouterr().out
    assert "7 file(s) imported" in out and "4 reconciliation(s)" in out


# ------------------------------------------------------------------ firm work dashboard (Phase 12C)
def test_firm_work_stages(db, sample_files, tmp_path):
    from fintax_gst.core.models import ExceptionStatus
    from fintax_gst.services.auth import SETUP_SESSION
    from fintax_gst.services.firm import STAGES, UNASSIGNED, FirmService, export_xlsx
    from fintax_gst.services.interest import InterestService
    from fintax_gst.services.notices import NoticeService
    from fintax_gst.services.reconciliation import ReconciliationService
    from fintax_gst.samples import gstin

    today = date(2025, 5, 15)
    clients = ClientService(db)
    demo = clients.create(**CLIENT)
    clients.update(demo.id, "admin", staff="priya")
    clients.create(gstin("29", "AACCB5678L"), None, name="Bangalore Electronics LLP", staff="rahul")
    partial = clients.create(gstin("27", "AABCM1234K"), None, name="Mahalaxmi Steel", return_frequency="Quarterly")
    imports = ImportService(db)
    imports.import_file(partial.id, sample_files["books_april"], Source.BOOKS)  # register without GSTR-2B
    firm = FirmService(db)

    def stage(name, **kw):
        return next(c for c in firm.work("2025-04", today, **kw).clients if c.name == name)

    assert stage("Bangalore Electronics LLP").stage == "Awaiting data"
    m = stage("Mahalaxmi Steel")
    assert m.stage == "Data incomplete" and "GSTR-2B" in m.next_step
    # QRMP, not a quarter end: PMT-06 is due, the IFF is optional
    assert [(r.return_type, r.optional) for r in m.returns] == [("IFF", True), ("PMT-06", False)]

    imports.import_file(demo.id, sample_files["books_april"], Source.BOOKS)
    imports.import_file(demo.id, sample_files["gstr2b_april_json"], Source.GSTR2B)
    assert stage(CLIENT["name"]).stage == "Ready to reconcile"
    recon = ReconciliationService(db)
    run = recon.run(demo.id, "2025-04")
    open_lines = [ln for ln in recon.load(run).result.lines if ln.status.value not in
                  ("No Action Required", "Resolved", "Closed After Review")]
    for ln in open_lines[:3]:
        recon.update_line(run, ln.line_no, "priya", assigned_to="priya", action_date=date(2025, 5, 10))
    w = stage(CLIENT["name"])
    assert w.stage == "Under review" and w.pending == len(open_lines)
    assert w.pending_by_person == {"priya": 3, UNASSIGNED: len(open_lines) - 3} and w.overdue_actions == 3
    work = firm.work("2025-04", today)
    loads = {s.person: s for s in work.staff}
    assert loads["priya"].pending_lines == 3 and loads["priya"].overdue_actions == 3 and loads["priya"].clients == 1
    assert [c.name for c in firm.work("2025-04", today, staff="rahul").clients] == ["Bangalore Electronics LLP"]
    assert work.clients[0].stage_no <= work.clients[-1].stage_no  # sorted by stage

    for ln in open_lines:
        recon.update_line(run, ln.line_no, "priya", status=ExceptionStatus.RESOLVED.value)
    w = stage(CLIENT["name"])
    assert w.stage == "Reviewed – file returns" and "GSTR-1 overdue" in w.next_step
    deadlines = {(d.client, d.item): d.state for d in firm.work("2025-04", today).deadlines}
    assert deadlines[(CLIENT["name"], "GSTR-1 not filed")] == "overdue"
    assert deadlines[(CLIENT["name"], "GSTR-3B not filed")] == "due soon"

    NoticeService(db).add(demo.id, SETUP_SESSION, form="ASMT-10", reference="ZD27", notice_date=date(2025, 5, 1),
                          period_from="2025-04", period_to="2025-04", reply_due=date(2025, 5, 25))
    interest = InterestService(db)
    interest.set_filing(demo.id, "GSTR-1", "2025-04", date(2025, 5, 12), SETUP_SESSION)
    interest.set_filing(demo.id, "GSTR-3B", "2025-04", date(2025, 5, 19), SETUP_SESSION)
    interest.set_filing(partial.id, "PMT-06", "2025-04", date(2025, 5, 20), SETUP_SESSION)  # QRMP payment
    w = stage(CLIENT["name"])
    assert w.stage == "Done" and w.notices_open == 1 and w.next_reply_due == date(2025, 5, 25)
    assert [r.state(today) for r in w.returns] == ["filed late", "filed"]
    items = [(d.client, d.item) for d in firm.work("2025-04", today).deadlines]
    assert (CLIENT["name"], "Reply to ASMT-10 ZD27") in items and (CLIENT["name"], "GSTR-1 not filed") not in items

    extra = tmp_path / "late_bill.csv"  # a bill booked after the reconciliation
    extra.write_text("GSTIN,Invoice No,Invoice Date,Taxable Value,CGST,SGST\n"
                     f"{gstin('27', 'AABCM1234K')},LATE-1,28-04-2025,1000,90,90\n")
    imports.import_file(demo.id, extra, Source.BOOKS)
    w = stage(CLIENT["name"])
    assert w.stage == "Ready to reconcile" and "reconcile again" in w.next_step
    assert STAGES.index(w.stage) == 2

    out = export_xlsx(firm.work("2025-04", today), tmp_path / "firm.xlsx")
    from openpyxl import load_workbook

    wb = load_workbook(out)
    assert wb.sheetnames == ["Clients", "Deadlines", "Staff"] and wb["Clients"].max_row == 5
    with pytest.raises(ServiceError):
        firm.work("April")


def test_cli_firm(tmp_path, capsys):
    dbp = str(tmp_path / "cli.db")
    assert main(["--db", dbp, "client", "add", "--name", CLIENT["name"], "--gstin", CLIENT["gstin"]]) == 0
    capsys.readouterr()
    assert main(["--db", dbp, "firm", "--period", "2025-04", "--excel", str(tmp_path / "f.xlsx")]) == 0
    out = capsys.readouterr().out
    assert "Awaiting data" in out and CLIENT["name"] in out and (tmp_path / "f.xlsx").exists()
