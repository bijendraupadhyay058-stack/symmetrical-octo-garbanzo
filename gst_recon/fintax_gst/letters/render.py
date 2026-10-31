"""Render a supplier letter as Word (.docx), PDF and an unsent e-mail draft (.eml).

Nothing is sent: the .eml file opens as a draft in Outlook/Thunderbird so the user can
review it and send it from their own mailbox.
"""
from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import date
from email.message import EmailMessage
from email.utils import formatdate
from pathlib import Path

from ..core.money import fmt_inr
from .content import SupplierLetter


@dataclass
class LetterContext:
    client_name: str
    client_gstin: str
    client_address: str
    period_label: str
    letter_date: date
    reply_by: date
    reference: str
    signatory: str = ""
    contact_person: str = ""
    to_email: str = ""


def _intro(ctx: LetterContext, letter: SupplierLetter) -> list[str]:
    return [
        f"While reconciling our purchase records with GSTR-2B for {ctx.period_label}, we found the "
        f"following {len(letter.items)} document(s) issued by you that are not reflected, or are reflected "
        "differently, in GSTR-2B.",
        "Under Section 16(2)(aa) of the CGST Act, 2017, input tax credit is available to us only when the "
        "details of the invoice or debit note have been furnished by you in GSTR-1 / IFF and communicated "
        "to us in GSTR-2B.",
        f"We request you to take the action shown against each document and confirm by "
        f"{ctx.reply_by:%d-%m-%Y}. If you have already done so, please share the filing details.",
    ]


def _rows(letter: SupplierLetter) -> list[list[str]]:
    rows = []
    for i, it in enumerate(letter.items, start=1):
        rows.append([str(i), it.invoice_no, f"{it.invoice_date:%d-%m-%Y}" if it.invoice_date else "", it.doc_type,
                     fmt_inr(it.books.taxable) if it.books else "", fmt_inr(it.tax_books),
                     fmt_inr(it.tax_reported) if it.reported else "-", it.issue, it.request])
    return rows


HEADERS = ["#", "Document No", "Date", "Type", "Taxable (our books)", "GST (our books)", "GST in GSTR-2B", "Issue",
           "Action requested"]


def safe_name(text: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]+", "_", text).strip("_")[:60] or "supplier"


def render_docx(ctx: LetterContext, letter: SupplierLetter, path: Path) -> Path:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.shared import Mm, Pt

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Mm(297), Mm(210)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Mm(15))
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)
    head = doc.add_paragraph()
    head.add_run(ctx.client_name).bold = True
    head.add_run(f"\nGSTIN: {ctx.client_gstin}")
    if ctx.client_address:
        head.add_run("\n" + ctx.client_address)
    doc.add_paragraph(f"Ref: {ctx.reference}\nDate: {ctx.letter_date:%d-%m-%Y}")
    doc.add_paragraph(f"To\n{letter.name or 'The Supplier'}\nGSTIN: {letter.gstin}")
    p = doc.add_paragraph()
    p.add_run(f"Subject: Discrepancies between our records and GSTR-2B – {ctx.period_label}").bold = True
    doc.add_paragraph("Dear Sir / Madam,")
    for para in _intro(ctx, letter):
        doc.add_paragraph(para)
    table = doc.add_table(rows=1, cols=len(HEADERS))
    table.style = "Table Grid"
    for c, h in zip(table.rows[0].cells, HEADERS):
        c.text = h
        c.paragraphs[0].runs[0].bold = True
    for row in _rows(letter):
        cells = table.add_row().cells
        for c, v in zip(cells, row):
            c.text = v
    for row in table.rows:
        for c in row.cells:
            for par in c.paragraphs:
                for run in par.runs:
                    run.font.size = Pt(8.5)
    doc.add_paragraph(f"Tax not reflected in GSTR-2B for these documents: ₹{fmt_inr(letter.amount)}")
    if ctx.contact_person:
        doc.add_paragraph(f"For any clarification please contact {ctx.contact_person}.")
    doc.add_paragraph(f"Thanking you,\nFor {ctx.client_name}\n\n\n{ctx.signatory or 'Authorised Signatory'}")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


def render_pdf(ctx: LetterContext, letter: SupplierLetter, path: Path) -> Path:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    from ..reports.pdf_report import NAVY

    styles = getSampleStyleSheet()
    normal = ParagraphStyle("n", parent=styles["Normal"], fontSize=9.5, leading=12)
    cell = ParagraphStyle("c", parent=styles["Normal"], fontSize=7.5, leading=9)
    esc = lambda t: t.replace("&", "&amp;").replace("<", "&lt;").replace("₹", "Rs. ")  # noqa: E731
    story = [Paragraph(f"<b>{esc(ctx.client_name)}</b><br/>GSTIN: {ctx.client_gstin}"
                       + (f"<br/>{esc(ctx.client_address)}" if ctx.client_address else ""), normal),
             Spacer(1, 3 * mm),
             Paragraph(f"Ref: {esc(ctx.reference)}<br/>Date: {ctx.letter_date:%d-%m-%Y}", normal), Spacer(1, 3 * mm),
             Paragraph(f"To<br/>{esc(letter.name or 'The Supplier')}<br/>GSTIN: {letter.gstin}", normal),
             Spacer(1, 3 * mm),
             Paragraph(f"<b>Subject: Discrepancies between our records and GSTR-2B – {esc(ctx.period_label)}</b>",
                       normal), Spacer(1, 2 * mm), Paragraph("Dear Sir / Madam,", normal)]
    for para in _intro(ctx, letter):
        story += [Spacer(1, 1.5 * mm), Paragraph(esc(para), normal)]
    head = ParagraphStyle("h", parent=cell, textColor=colors.white, fontName="Helvetica-Bold")
    data = [[Paragraph(h, head) for h in HEADERS]] + [[Paragraph(esc(v), cell) for v in r] for r in _rows(letter)]
    t = Table(data, colWidths=[8 * mm, 32 * mm, 18 * mm, 18 * mm, 24 * mm, 22 * mm, 22 * mm, 52 * mm, 75 * mm],
              repeatRows=1)
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), NAVY), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                           ("FONTSIZE", (0, 0), (-1, 0), 7.5), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                           ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#90A4AE")),
                           ("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story += [Spacer(1, 3 * mm), t, Spacer(1, 3 * mm),
              Paragraph(f"Tax not reflected in GSTR-2B for these documents: Rs. {fmt_inr(letter.amount)}", normal)]
    if ctx.contact_person:
        story.append(Paragraph(f"For any clarification please contact {esc(ctx.contact_person)}.", normal))
    story += [Spacer(1, 4 * mm), Paragraph(f"Thanking you,<br/>For {esc(ctx.client_name)}<br/><br/><br/>"
                                           f"{esc(ctx.signatory or 'Authorised Signatory')}", normal)]
    path.parent.mkdir(parents=True, exist_ok=True)
    SimpleDocTemplate(str(path), pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm, topMargin=12 * mm,
                      bottomMargin=12 * mm, title=f"GST discrepancy – {letter.gstin}",
                      author=ctx.client_name).build(story)
    return path


def render_eml(ctx: LetterContext, letter: SupplierLetter, path: Path, attachment: Path | None = None) -> Path:
    msg = EmailMessage()
    msg["Subject"] = f"GST discrepancies – {ctx.period_label} – {ctx.client_name} ({ctx.client_gstin})"
    if ctx.to_email:
        msg["To"] = ctx.to_email
    msg["Date"] = formatdate(localtime=True)
    msg["X-Unsent"] = "1"  # opens as a draft in Outlook
    lines = [f"Dear {letter.name or 'Sir / Madam'},", ""] + _intro(ctx, letter) + [""]
    for r in _rows(letter):
        lines.append(f"{r[0]}. {r[1]} dated {r[2]} ({r[3]}) – GST in our books ₹{r[5]}, in GSTR-2B {r[6]}: "
                     f"{r[7]}. {r[8]}.")
    lines += ["", f"Tax not reflected in GSTR-2B: ₹{fmt_inr(letter.amount)}", "",
              "The letter is attached.", "", "Regards,", ctx.signatory or ctx.client_name, ctx.client_name]
    msg.set_content("\n".join(lines))
    if attachment is not None:
        msg.add_attachment(attachment.read_bytes(), maintype="application", subtype="pdf", filename=attachment.name)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(bytes(msg))
    return path
