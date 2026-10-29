from datetime import date
from decimal import Decimal

import pytest
from sqlalchemy import func, select

from fintax_gst.core.models import Category, DocType, Source
from fintax_gst.db.models import AuditLog, Invoice, ItcClaim
from fintax_gst.engine import reconcile
from fintax_gst.importers.gstr3b import Gstr3bValues, preview_gstr3b
from fintax_gst.importers.pipeline import preview_file
from fintax_gst.itc.deadlines import itc_deadline
from fintax_gst.itc.ims import compare_ims
from fintax_gst.itc.review import review_lines
from fintax_gst.itc.rules import DEFAULT_RULES, Rule, provisional_limit_pct, resolve
from fintax_gst.samples import CLIENT, SUPPLIERS
from fintax_gst.services.auth import SETUP_SESSION, AuthService, PermissionDenied, Session, require
from fintax_gst.services.clients import ClientService
from fintax_gst.services.common import ServiceError
from fintax_gst.services.followup import AttachmentService
from fintax_gst.services.gstr3b import Gstr3bService
from fintax_gst.services.imports import ImportService
from fintax_gst.services.itc import ItcService
from fintax_gst.services.reconciliation import ReconciliationService
from fintax_gst.services.security import SecurityService

from .conftest import GSTIN_A, rec

AS_OF = date(2025, 5, 15)


def _demo(db, files, g2b_key="gstr2b_april_json", with_ims=True, with_3b=True):
    c = ClientService(db).create(**CLIENT)
    svc = ImportService(db)
    keys = [("books_march", Source.BOOKS), ("books_april", Source.BOOKS), (g2b_key, Source.GSTR2B)]
    if with_ims:
        keys.append(("ims_april_json", Source.IMS))
    for key, src in keys:
        svc.import_file(c.id, files[key], src)
    if with_3b:
        Gstr3bService(db).import_file(c.id, files["gstr3b_april_json"])
    run = ReconciliationService(db).run(c.id, "2025-04")
    return c, run


def _rules():
    return resolve([], [])


# ------------------------------------------------------------------ rules
def test_section_16_4_deadlines():
    assert itc_deadline(date(2025, 4, 10))[0] == date(2026, 11, 30)
    assert itc_deadline(date(2026, 3, 31))[0] == date(2026, 11, 30)
    assert itc_deadline(date(2021, 4, 1))[0] == date(2022, 11, 30)  # FY 2021-22: amended limit
    assert itc_deadline(date(2020, 6, 1))[0] == date(2021, 10, 20)  # FY 2020-21: September return


def test_rule_effective_dates_and_keywords():
    r = Rule("CUSTOM_X", "x", "ref", "keyword", "blocked", keywords=["cab"], effective_from=date(2024, 1, 1),
             effective_to=date(2024, 12, 31))
    assert r.applies_on(date(2024, 6, 1)) and not r.applies_on(date(2025, 1, 1))
    assert r.keyword_hit("City Cab Services") == "cab"
    assert r.keyword_hit("Cabinet Works") is None  # whole words only
    r.enabled = False
    assert not r.applies_on(date(2024, 6, 1))


def test_rule_36_4_schedule():
    r = next(x for x in DEFAULT_RULES if x.code == "R36_4")
    assert provisional_limit_pct(r, date(2019, 11, 1)) == 20
    assert provisional_limit_pct(r, date(2020, 6, 1)) == 10
    assert provisional_limit_pct(r, date(2021, 6, 1)) == 5
    assert provisional_limit_pct(r, date(2022, 1, 1)) is None  # Section 16(2)(aa) instead


def test_review_flags_with_effective_dates():
    food = rec("B", "F-1", supplier_name="Annapurna Caterers")
    old_missing = rec("B", "OLD-1", dt="2021-06-10", period="2021-06")
    new_missing = rec("B", "NEW-1", dt="2025-04-10")
    lines = reconcile([food, old_missing, new_missing], [rec("G", "F-1")]).lines
    reviews = {r.record.invoice_no: r for r in review_lines(lines, _rules(), AS_OF)}
    assert {h.code for h in reviews["F-1"].hits} == {"S17_5_B_I_FOOD"}
    assert "S16_2_AA" in {h.code for h in reviews["NEW-1"].hits}
    old_codes = {h.code for h in reviews["OLD-1"].hits}
    assert "S16_2_AA" not in old_codes  # section in force only from 01-01-2022
    assert "S16_4" in old_codes and reviews["OLD-1"].days_left < 0  # time limit long passed


def test_time_limit_alert_window():
    doc = rec("B", "T-9", dt="2025-03-20", period="2025-03")
    lines = reconcile([doc], [rec("G", "T-9", dt="2025-03-20")]).lines
    near = review_lines(lines, _rules(), date(2025, 11, 1))[0]
    hit = next(h for h in near.hits if h.code == "S16_4")
    assert hit.severity == "medium" and near.days_left == 29
    assert not any(h.code == "S16_4" for h in review_lines(lines, _rules(), date(2025, 6, 1))[0].hits)


def test_client_specific_rule_overrides_global(db):
    c = ClientService(db).create(**CLIENT)
    itc = ItcService(db)
    rule = next(r for r in itc.rules(c.id) if r.code == "CLIENT_BLOCKED_GSTIN")
    rule.gstins = [GSTIN_A]
    itc.save_rule(rule, c.id, "ca")
    assert next(r for r in itc.rules(c.id) if r.code == "CLIENT_BLOCKED_GSTIN").gstins == [GSTIN_A]
    assert next(r for r in itc.rules(None) if r.code == "CLIENT_BLOCKED_GSTIN").gstins == []
    food = next(r for r in itc.rules(None) if r.code == "S17_5_B_I_FOOD")
    food.enabled = False
    itc.save_rule(food, None, "ca")
    assert not next(r for r in itc.rules(c.id) if r.code == "S17_5_B_I_FOOD").enabled
    with pytest.raises(ServiceError):
        itc.save_rule(Rule("MY_RULE", "x", "y", "keyword", "blocked"), c.id)
    itc.reset_rule("S17_5_B_I_FOOD")
    assert next(r for r in itc.rules(c.id) if r.code == "S17_5_B_I_FOOD").enabled


# -------------------------------------------------------------- importers
def test_gstr3b_json_and_template(sample_files):
    p = preview_gstr3b(sample_files["gstr3b_april_json"])
    t = p.entries["2025-04"]
    assert t["4A1"].igst == Decimal("1800") and t["4A4"].cgst == Decimal("900") and t["4B1"].total == Decimal("1440")
    a = sum((t[c].total for c in ("4A1", "4A2", "4A3", "4A4", "4A5")), Decimal(0))
    assert t["4C"].total == a - t["4B1"].total - t["4B2"].total
    tpl = preview_gstr3b(sample_files["gstr3b_template_csv"])
    assert tpl.periods == ["2025-03"] and tpl.entries["2025-03"]["4A5"].total == Decimal("4770.00")


def test_gstr3b_reimport_supersedes_and_removal_restores(db, sample_files, tmp_path):
    c = ClientService(db).create(**CLIENT)
    g = Gstr3bService(db)
    first, _ = g.import_file(c.id, sample_files["gstr3b_april_json"])
    revised = tmp_path / "revised.csv"
    revised.write_text("Period,Table,IGST,CGST,SGST,Cess\n2025-04,4A5,100,0,0,0\n")
    g.import_file(c.id, revised)
    assert g.values(c.id, ["2025-04"])["2025-04"]["4A5"].total == Decimal("100")
    with pytest.raises(ServiceError):
        g.import_file(c.id, revised)  # same file twice
    second = ImportService(db).list_files(c.id)[-1]
    ImportService(db).remove_import(c.id, second.id)
    assert g.values(c.id, ["2025-04"])["2025-04"]["4A1"].igst == Decimal("1800")
    assert first.id


def test_gstr2b_json_isd_and_imports_stored(db, sample_files):
    c, _run = _demo(db, sample_files, with_ims=False, with_3b=False)
    isd, imports = ItcService(db)._other_totals(c.id, ["2025-04"])
    assert isd == Decimal("1800") and imports == Decimal("1800")
    p = preview_file(sample_files["gstr2b_april_json"], Source.GSTR2B)
    assert len(p.other_itc) == 2


def test_ims_import_and_comparison(db, sample_files):
    p = preview_file(sample_files["ims_april_json"], Source.IMS)
    assert p.layout == "ims_json" and {r.ims_action for r in p.records} == {"A", "R", "N", "P"}
    with pytest.raises(Exception):
        preview_file(sample_files["ims_april_json"], Source.GSTR2B)
    c, _run = _demo(db, sample_files)
    rows, notes = ItcService(db).ims_compare(c.id, ["2025-04"])
    by_no = {r.ims.invoice_no: r for r in rows}
    rej = by_no["GCH-2201"]
    assert rej.action == "R" and rej.worst == "high"
    assert any("recompute GSTR-2B" in t for _s, t in rej.observations)
    assert any("ITC deferred" in t for _s, t in by_no["KTX-2290"].observations)
    assert any("not recorded in books" in t for _s, t in by_no["OSL/377"].observations)
    assert not notes


def test_ims_before_october_2024_is_noted():
    rows, notes = compare_ims([rec("G", "X-1", ims_action="A")], [], [], period="2024-06")
    assert notes and "October 2024" in notes[0]


# ----------------------------------------------------------- ITC review
def test_review_run_flags_and_buckets(db, sample_files):
    c, run = _demo(db, sample_files)
    reviews, b = ItcService(db).review_run(run, AS_OF)
    codes = {r.record.invoice_no: {h.code for h in r.hits} for r in reviews}
    assert {"S17_5_B_I_FOOD", "S17_5_H_GIFTS"} <= codes["ANC/77"]
    assert "RCM" in codes["SRW-310"] and "ITC_NA_2B" in codes["KTX-2301"]
    assert "IMS_REJECTED" in codes["GCH-2201"] and "IMS_PENDING" in codes["KTX-2290"]
    assert "S16_2_AA" in codes["MSL/260/25-26"]
    assert b["blocked_17_5"] == Decimal("1440.00") and b["reverse_charge"] == Decimal("1000")
    assert b["isd"] == Decimal("1800") and b["imports"] == Decimal("1800")
    assert b["itc_3b"] > 0 and b["claimed_register"] == 0


def test_dismissed_flag_needs_reason_and_leaves_bucket(db, sample_files):
    c, run = _demo(db, sample_files)
    itc = ItcService(db)
    reviews, _ = itc.review_run(run, AS_OF)
    anc = next(r for r in reviews if r.record.invoice_no == "ANC/77")
    with pytest.raises(ServiceError):
        itc.set_review(c.id, anc.record.db_id, "S17_5_B_I_FOOD", "DISMISSED", "")
    for code in ("S17_5_B_I_FOOD", "S17_5_H_GIFTS"):
        itc.set_review(c.id, anc.record.db_id, code, "DISMISSED", "Raw material supplier, not food", "ca")
    reviews, b = itc.review_run(run, AS_OF)
    assert b["blocked_17_5"] == 0
    anc = next(r for r in reviews if r.record.invoice_no == "ANC/77")
    assert all(h.decision == "DISMISSED" for h in anc.hits) and anc.worst == ""


# ----------------------------------------------------------------- claims
def test_claim_register_prevents_double_claims(db, sample_files):
    c, run = _demo(db, sample_files)
    itc = ItcService(db)
    res = itc.claim_eligible_from_run(run, "2025-04", "ca", AS_OF)
    claimed_nos = set()
    with db.session() as s:
        for inv_id in res.claimed:
            claimed_nos.add(s.get(Invoice, inv_id).invoice_no)
    assert "BEL-2001" in claimed_nos and "ANC/77" not in claimed_nos and "MSL/260/25-26" not in claimed_nos
    assert any("open review flag" in why for _i, why in res.skipped)
    again = itc.claim(c.id, res.claimed[:3], "2025-05", "ca")
    assert again.claimed == [] and len(again.skipped) == 3 and "already claimed in 2025-04" in again.skipped[0][1]
    # a rerun of the period does not create claims twice either
    res2 = itc.claim_eligible_from_run(run, "2025-04", "ca", AS_OF)
    assert res2.claimed == []
    with db.session() as s:
        assert s.scalar(select(func.count()).select_from(ItcClaim)) == len(res.claimed)


def test_reverse_and_reclaim_limits(db, sample_files):
    c, run = _demo(db, sample_files)
    itc = ItcService(db)
    with db.session() as s:
        inv = s.scalar(select(Invoice).where(Invoice.invoice_no == "BEL-2001"))
    with pytest.raises(ServiceError):
        itc.reverse(c.id, inv.id, "2025-05", "ca", "Rule 37")  # nothing claimed yet
    itc.claim(c.id, [inv.id], "2025-04", "ca")
    with pytest.raises(ServiceError):
        itc.reverse(c.id, inv.id, "2025-05", "ca", "")  # reason required
    with pytest.raises(ServiceError):
        itc.reverse(c.id, inv.id, "2025-05", "ca", "too much", amounts=Gstr3bValues(igst=inv.igst + 1))
    itc.reverse(c.id, inv.id, "2025-10", "ca", "Not paid within 180 days", "RULE_37_180",
                Gstr3bValues(igst=Decimal("1000")))
    with pytest.raises(ServiceError):
        itc.reclaim(c.id, inv.id, "2025-11", "ca", "paid", Gstr3bValues(igst=Decimal("1000.01")))
    itc.reclaim(c.id, inv.id, "2025-11", "ca", "Paid supplier on 02-11-2025")
    st = itc.claim_states(c.id, [inv.id])[inv.id]
    assert st.claimed == inv.igst and st.reversed == Decimal("1000") and st.reclaimed == Decimal("1000")
    assert st.net == inv.igst
    with pytest.raises(ServiceError):
        itc.reclaim(c.id, inv.id, "2025-12", "ca", "again")  # nothing left to reclaim


def test_only_books_documents_can_be_claimed(db, sample_files):
    c, run = _demo(db, sample_files)
    with db.session() as s:
        g = s.scalar(select(Invoice).where(Invoice.source_type == "GSTR2B"))
    with pytest.raises(ServiceError):
        ItcService(db).claim(c.id, [g.id], "2025-04")


# ---------------------------------------------------------- 3B / annual
def test_table4_working_finds_excess_claim(db, sample_files):
    c, run = _demo(db, sample_files)
    w = ItcService(db).table4(c.id, ["2025-04"], AS_OF)[0]
    assert w.row("4A1").auto_2b.igst == Decimal("1800") and w.row("4A4").auto_2b.total == Decimal("1800")
    assert w.row("4D2").auto_2b.total == Decimal("6480")
    excess = [f for f in w.findings if f.kind == "statutory" and "exceeds GSTR-2B" in f.text]
    assert excess and excess[0].amount == Decimal("2000")
    assert w.row("4A5").diff_filed_2b == Decimal("2000")
    kinds = {f.kind for f in w.findings}
    assert {"statutory", "timing", "accounting"} <= kinds
    assert "no claim register" in w.books_basis


def test_table4_uses_claim_register_when_present(db, sample_files):
    c, run = _demo(db, sample_files)
    itc = ItcService(db)
    res = itc.claim_eligible_from_run(run, "2025-04", "ca", AS_OF)
    w = itc.table4(c.id, ["2025-04"], AS_OF)[0]
    with db.session() as s:
        total = Decimal(0)
        for inv_id in res.claimed:
            i = s.get(Invoice, inv_id)
            sign = -1 if i.doc_type == "CRN" else 1
            total += (i.igst + i.cgst + i.sgst + i.cess) * sign
    assert w.books_basis == "Claim register for the period"
    assert w.row("4A5").books.total + w.row("4A3").books.total == total


def test_rule_36_4_allowance_before_2022(db, tmp_path):
    from fintax_gst.itc.table4 import build_period

    g = [rec("G", "Q-1", dt="2021-05-05", period="2021-05", taxable="10000", cgst="900", sgst="900")]
    filed = {"4A5": Gstr3bValues(cgst=Decimal("940"), sgst=Decimal("940"))}  # 4.4 % over GSTR-2B
    w = build_period("2021-05", g, [], filed, [], [], [], _rules())
    assert w.provisional_pct == 5 and not [f for f in w.findings if f.kind == "statutory"]
    w = build_period("2022-05", [rec("G", "Q-1", dt="2022-05-05", period="2022-05", taxable="10000", cgst="900",
                                     sgst="900")], [], filed, [], [], [], _rules())
    assert any("Section 16(2)(aa)" in f.text for f in w.findings)


def test_gstr9_table8(db, sample_files):
    c, run = _demo(db, sample_files)
    itc = ItcService(db)
    t = itc.annual(c.id, "2025-26")
    with db.session() as s:
        docs = s.scalars(select(Invoice).where(Invoice.source_type == "GSTR2B")).all()
    expected = Decimal(0)
    for i in docs:
        if i.invoice_date >= date(2025, 4, 1) and not i.reverse_charge and i.itc_availability != "N":
            expected += (i.igst + i.cgst + i.sgst + i.cess) * (-1 if i.doc_type == "CRN" else 1)
    assert t.a.total == expected
    assert t.b.total == Gstr3bService(db).values(c.id, ["2025-04"])["2025-04"]["4A5"].total
    assert t.d.total == t.a.total - t.b.total - t.c.total
    assert len(t.b_periods_missing) == 11 and len(t.months) == 12
    # ITC of FY 2025-26 documents claimed in the next year (before the cut-off) goes to 8C.
    with db.session() as s:
        inv = s.scalar(select(Invoice).where(Invoice.invoice_no == "BEL-2001"))
    itc.claim(c.id, [inv.id], "2026-05", "ca")
    assert itc.annual(c.id, "2025-26").c.total == inv.igst


# ----------------------------------------------------------------- security
def test_users_roles_and_lockout(db):
    auth = AuthService(db)
    assert not auth.has_users()
    with pytest.raises(ServiceError):
        auth.create_user("asha", "weak", "admin")
    with pytest.raises(ServiceError):
        auth.create_user("asha", "Secret123", "accountant")  # first user must be admin
    auth.create_user("admin", "Admin12345", "admin", "Owner")
    admin = auth.login("admin", "Admin12345")
    with pytest.raises(PermissionDenied):
        auth.create_user("bob", "Secret123", "viewer", by=Session("x", "accountant"))
    auth.create_user("viewer1", "Viewer123", "viewer", by=admin)
    v = auth.login("viewer1", "Viewer123")
    with pytest.raises(PermissionDenied):
        require(v, "reconcile")
    require(v, "export")
    for _ in range(5):
        with pytest.raises(ServiceError):
            auth.login("viewer1", "wrong-pass1")
    with pytest.raises(ServiceError, match="locked"):
        auth.login("viewer1", "Viewer123")
    with db.session() as s:
        assert s.scalar(select(func.count()).select_from(AuditLog).where(AuditLog.action == "LOGIN_FAILED")) >= 6
    with pytest.raises(ServiceError):
        auth.set_role("admin", "viewer", admin)  # last admin cannot be demoted
    u = next(x for x in auth.users() if x.username == "admin")
    assert "Admin12345" not in u.password_hash and u.password_hash.startswith("scrypt$")


def test_encrypted_backup_roundtrip(db, sample_files, tmp_path):
    c = ClientService(db).create(**CLIENT)
    sec = SecurityService(db)
    enc = sec.encrypted_backup(tmp_path / "b.ftpbk", "Backup#2025")
    raw = enc.read_bytes()
    assert raw.startswith(b"FTPBK1") and CLIENT["gstin"].encode() not in raw
    with pytest.raises(ServiceError):
        sec.decrypt_backup(enc, "wrong-password", tmp_path / "x.db")
    ImportService(db).import_file(c.id, sample_files["books_april"], Source.BOOKS)
    sec.restore_encrypted(enc, "Backup#2025", SETUP_SESSION)
    assert ImportService(db).list_files(c.id) == []
    with pytest.raises(PermissionDenied):
        sec.restore_encrypted(enc, "Backup#2025", Session("a", "accountant"))


def test_auto_backup_and_retention(db):
    sec = SecurityService(db)
    first = sec.auto_backup()
    assert first is not None and first.exists()
    assert sec.auto_backup() is None  # within the interval
    sec.set_setting("backup_keep", 2, SETUP_SESSION)
    import os
    import time

    made = []
    for i in range(4):
        p = sec.backup_dir / f"fintaxpro-2025010{i}-000000-auto.db"
        p.write_bytes(b"x")
        os.utime(p, (time.time() - 1000 * (i + 1), time.time() - 1000 * (i + 1)))
        made.append(p)
    removed = sec.apply_retention()
    left = list(sec.backup_dir.glob("*-auto.db"))
    assert len(left) == 2 and first in left and len(removed) == 3


def test_purge_client_removes_everything(db, sample_files, tmp_path):
    c, run = _demo(db, sample_files)
    other = ClientService(db).create(**{**CLIENT, "gstin": SUPPLIERS["BEL"][0], "name": "Other"})
    ev = tmp_path / "ev.pdf"
    ev.write_bytes(b"%PDF evidence")
    lines = ReconciliationService(db).load(run).result.lines
    att = AttachmentService(db).add(run, lines[0].line_no, ev, "ca")
    ItcService(db).claim_eligible_from_run(run, "2025-04", "ca", AS_OF)
    sec = SecurityService(db)
    with pytest.raises(PermissionDenied):
        sec.purge_client(c.id, CLIENT["gstin"], Session("a", "manager"))
    with pytest.raises(ServiceError):
        sec.purge_client(c.id, "WRONG", SETUP_SESSION)
    counts = sec.purge_client(c.id, CLIENT["gstin"], SETUP_SESSION)
    assert counts["invoices"] > 0 and counts["itc_claims"] > 0 and counts["gstr3b_entries"] > 0
    with db.session() as s:
        assert s.scalar(select(func.count()).select_from(Invoice)) == 0
        purge = s.scalars(select(AuditLog).where(AuditLog.action == "PURGE_CLIENT")).one()
        assert CLIENT["gstin"] in purge.details_json
    assert not __import__("pathlib").Path(att.stored_path).exists()
    assert [x.id for x in ClientService(db).list()] == [other.id]
    raw = db.path.read_bytes()
    assert b"Mahalaxmi Steel" not in raw  # secure_delete + VACUUM: no residue in the file


def test_audit_export(db, tmp_path):
    ClientService(db).create(**CLIENT)
    out = SecurityService(db).export_audit(tmp_path / "audit.csv", by=SETUP_SESSION)
    text = out.read_text()
    assert "CREATE" in text and text.splitlines()[0].startswith("id,at_utc,user")


def test_credit_note_in_2b_not_booked_flag():
    cn = rec("G", "CN-9", doc=DocType.CREDIT_NOTE)
    lines = reconcile([], [cn]).lines
    assert lines[0].category == Category.MISSING_IN_BOOKS
    r = review_lines(lines, _rules(), AS_OF)
    assert r and r[0].hits[0].code == "CN_2B_NOT_BOOKED" and r[0].itc == Decimal("-180")


def test_workbook_itc_sheets_are_populated(db, sample_files, tmp_path):
    from openpyxl import load_workbook

    from fintax_gst.services.reports import export_run

    c, run = _demo(db, sample_files)
    xlsx, = export_run(db, run, tmp_path, pdf=False, as_of=AS_OF)
    wb = load_workbook(xlsx)
    t4 = [r for r in wb["GSTR-3B ITC Comparison"].iter_rows(values_only=True)]
    assert any(r[0] == "2025-04" and r[1] == "4A5" and r[9] == 2000 for r in t4)
    assert any(r and r[1] == "Statutory" for r in t4)
    ims = [r for r in wb["IMS Review"].iter_rows(values_only=True)]
    assert any(r[2] == "GCH-2201" and r[4] == "Rejected" for r in ims)
    flags = [r for r in wb["ITC Eligibility Review"].iter_rows(values_only=True)]
    assert any(r[3] == "ANC/77" and r[10] == "S17_5_B_I_FOOD" for r in flags)
    assert any(r[0] == "ITC potentially blocked – Section 17(5)" and r[1] == 1440 for r in flags)


def test_annual_workbook(db, sample_files, tmp_path):
    from openpyxl import load_workbook

    from fintax_gst.reports.annual_report import export_annual

    c, run = _demo(db, sample_files)
    out = export_annual(db, c.id, "2025-26", tmp_path, portal_8a=Decimal("120000"))
    wb = load_workbook(out)
    assert wb.sheetnames == ["Table 8 Summary", "Month-wise ITC", "8A Docs Not in Books", "Books Not in 8A",
                             "Claim Register", "ITC Rules Applied"]
    t = ItcService(db).annual(c.id, "2025-26")
    rows = list(wb["Table 8 Summary"].iter_rows(values_only=True))
    assert any(r[0] and str(r[0]).startswith("8D") and r[5] == t.d.total for r in rows)
    assert any(r[0] and "Portal Table 8A entered" in str(r[0]) for r in rows)


def test_cli_login_permissions_and_itc_commands(tmp_path, sample_files, capsys, monkeypatch):
    from fintax_gst.cli import main

    monkeypatch.setenv("FINTAX_HOME", str(tmp_path / "home"))
    db = str(tmp_path / "cli3.db")
    base = ["--db", db]
    assert main(base + ["client", "add", "--gstin", CLIENT["gstin"], "--name", CLIENT["name"]]) == 0
    assert "setup mode" in capsys.readouterr().err
    monkeypatch.setenv("FINTAX_NEW_PASSWORD", "Admin12345")
    assert main(base + ["user", "add", "--username", "admin", "--role", "admin"]) == 0
    # From now on a login is required.
    assert main(base + ["client", "list"]) == 1
    assert "Log in" in capsys.readouterr().err
    monkeypatch.setenv("FINTAX_PASSWORD", "Admin12345")
    login = base + ["--login", "admin"]
    monkeypatch.setenv("FINTAX_NEW_PASSWORD", "Viewer1234")
    assert main(login + ["user", "add", "--username", "viewer", "--role", "viewer"]) == 0
    for src, key in (("books", "books_april"), ("gstr2b", "gstr2b_april_json"), ("ims", "ims_april_json")):
        assert main(login + ["import", "--client", "1", "--source", src, "--file", str(sample_files[key])]) == 0
    assert main(login + ["gstr3b", "import", "--client", "1", "--file", str(sample_files["gstr3b_april_json"])]) == 0
    assert main(login + ["reconcile", "--client", "1", "--from", "2025-04"]) == 0
    capsys.readouterr()
    assert main(login + ["itc", "review", "--run", "1", "--as-of", "2025-05-15"]) == 0
    out = capsys.readouterr().out
    assert "S17_5_B_I_FOOD" in out and "ITC potentially blocked" in out
    assert main(login + ["itc", "claim", "--run", "1", "--period", "2025-04", "--as-of", "2025-05-15"]) == 0
    assert "Claimed" in capsys.readouterr().out
    assert main(login + ["itc", "table4", "--client", "1", "--from", "2025-04", "--as-of", "2025-05-15"]) == 0
    assert "exceeds GSTR-2B" in capsys.readouterr().out
    assert main(login + ["ims", "--client", "1", "--period", "2025-04"]) == 0
    assert "Rejected in IMS" in capsys.readouterr().out
    assert main(login + ["annual", "--client", "1", "--fy", "2025-26", "--out", str(tmp_path / "ann")]) == 0
    assert list((tmp_path / "ann").glob("*GSTR9*.xlsx"))
    monkeypatch.setenv("FINTAX_BACKUP_PASSWORD", "Backup#2025")
    assert main(login + ["secure-backup", "--to", str(tmp_path / "b.ftpbk")]) == 0
    assert main(login + ["audit-export", "--out", str(tmp_path / "audit.csv")]) == 0
    # A viewer may look and export but not change anything.
    monkeypatch.setenv("FINTAX_PASSWORD", "Viewer1234")
    viewer = base + ["--login", "viewer"]
    assert main(viewer + ["runs", "--client", "1"]) == 0
    assert main(viewer + ["reconcile", "--client", "1", "--from", "2025-04"]) == 1
    assert main(viewer + ["purge-client", "--client", "1", "--confirm", CLIENT["gstin"]]) == 1
    assert "cannot" in capsys.readouterr().err
    monkeypatch.setenv("FINTAX_PASSWORD", "wrong-pass1")
    assert main(viewer + ["runs", "--client", "1"]) == 1
