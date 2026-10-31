"""Phase 14: tax type / rate checks, journal and payment registers, 180-day tracker, PAN view, Tally journals."""
from datetime import date
from decimal import Decimal as D

import pytest

from fintax_gst.compliance.tax_checks import Doc, check, fits_rate, pos_code
from fintax_gst.core.models import Source
from fintax_gst.samples import CLIENT, gstin
from fintax_gst.services.clients import ClientService
from fintax_gst.services.common import ServiceError
from fintax_gst.services.imports import ImportService

MH, KA = gstin("27", "AABCM1234K"), gstin("29", "AACCB5678L")


def _doc(source="BOOKS", g=MH, taxable="1000", igst="0", cgst="0", sgst="0", **kw):
    return Doc(1, source, g, "Party", "INV-1", date(2025, 4, 2), D(taxable), D(igst), D(cgst), D(sgst), **kw)


# ------------------------------------------------------------------ tax type and rate checks (14A)
def test_place_of_supply_and_rates():
    assert [pos_code(x) for x in ("27", "27-Maharashtra", "Karnataka (29)", "tamil nadu", "Orissa", "", "x")] == \
        ["27", "27", "29", "33", "21", None, None]
    assert fits_rate(D("1000"), D("180"), D(0), D(0)) == D("18")
    assert fits_rate(D("1000"), D(0), D("90"), D("90.40")) == D("18")  # rounding
    assert fits_rate(D("1000"), D("400"), D(0), D(0)) == D("40")  # GST 2.0 rate from 22-09-2025
    assert fits_rate(D("1000"), D("150"), D(0), D(0)) is None
    assert fits_rate(D("1000"), D(0), D(0), D(0)) is not None  # nil / exempt: nothing to check


@pytest.mark.parametrize("doc, kinds", [
    (_doc(cgst="90", sgst="90"), []),  # Maharashtra supplier, CGST + SGST: right
    (_doc(igst="180"), ["tax_type"]),  # same state, IGST charged
    (_doc(g=KA, igst="180"), []),  # inter-state, IGST: right
    (_doc(g=KA, cgst="90", sgst="90", place_of_supply="27-Maharashtra"), ["tax_type"]),
    (_doc(g=KA, cgst="90", sgst="90"), ["other_state"]),  # no place of supply: hotel there, or wrong head
    (_doc(g=KA, cgst="90", sgst="90", place_of_supply="Karnataka"), ["other_state"]),
    (_doc(g=KA, cgst="90", sgst="90", reverse_charge=True), []),
    (_doc(igst="100", cgst="40", sgst="40"), ["tax_type"]),  # IGST and CGST/SGST on one document
    (_doc(cgst="75", sgst="75"), ["rate"]),  # 15 %: no such rate
    (_doc(cgst="57.50", sgst="57.50"), ["rate"]),  # 11.5 %: two rates combined, or an error
    (_doc(cgst="57.50", sgst="57.50", lines=((D(500), D(0), D("12.50"), D("12.50")),
                                             (D(500), D(0), D(45), D(45)))), []),  # 5 % + 18 % lines: right
    (_doc("SALES", g=KA, cgst="90", sgst="90"), ["tax_type"]),  # customer in another state
    (_doc("SALES", g="", cgst="90", sgst="90"), []),  # local B2C
    (_doc("SALES", g="", igst="180", place_of_supply="29-Karnataka"), []),
    (_doc("SALES", g="", cgst="90", sgst="90", supply_type="Export (LUT)"), ["tax_type"]),
    (_doc("SALES", g=MH, igst="180"), ["tax_type"]),
])
def test_tax_checks(doc, kinds):
    assert [i.kind for i in check(doc, "27")] == kinds


def test_tax_check_service(db, tmp_path):
    from openpyxl import load_workbook

    from fintax_gst.services.tax_checks import TaxCheckService

    c = ClientService(db).create(**CLIENT)
    reg = tmp_path / "purchases.csv"
    reg.write_text("GSTIN,Invoice No,Invoice Date,Taxable Value,IGST,CGST,SGST,Place of Supply\n"
                   f"{MH},OK-1,02-04-2025,1000,0,90,90,Maharashtra\n"
                   f"{MH},IG-1,03-04-2025,1000,180,0,0,Maharashtra\n"
                   f"{KA},HOT-1,04-04-2025,2000,0,180,180,\n"
                   f"{MH},RT-1,05-04-2025,1000,0,75,75,\n")
    ImportService(db).import_file(c.id, reg, Source.BOOKS)
    issues = TaxCheckService(db).check(c.id, "2025-04")
    assert [(i.doc.invoice_no, i.kind) for i in issues] == [("IG-1", "tax_type"), ("HOT-1", "other_state"),
                                                             ("RT-1", "rate")]
    out = TaxCheckService.export_xlsx(issues, tmp_path / "t.xlsx", "April")
    assert load_workbook(out)["Tax checks"].max_row == 5
    assert TaxCheckService(db).check(c.id, "2025-05") == []
    with pytest.raises(ServiceError):
        TaxCheckService(db).check(c.id, "April")


def test_cli_tax_check(tmp_path, capsys):
    from fintax_gst.cli import main

    dbp = str(tmp_path / "cli.db")
    reg = tmp_path / "p.csv"
    reg.write_text(f"GSTIN,Invoice No,Invoice Date,Taxable Value,IGST\n{MH},IG-1,03-04-2025,1000,180\n")
    assert main(["--db", dbp, "client", "add", "--name", CLIENT["name"], "--gstin", CLIENT["gstin"]]) == 0
    assert main(["--db", dbp, "import", "--client", CLIENT["gstin"], "--source", "books", "--file", str(reg)]) == 0
    capsys.readouterr()
    assert main(["--db", dbp, "tax-check", "--client", CLIENT["gstin"], "--from", "2025-04"]) == 0
    out = capsys.readouterr().out
    assert "Wrong tax type" in out and "IG-1" in out and "1 finding(s)" in out
