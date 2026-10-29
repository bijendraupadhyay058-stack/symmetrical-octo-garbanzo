from datetime import date
from decimal import Decimal

from openpyxl import load_workbook

from fintax_gst.cli import main
from fintax_gst.core.models import DocType, InvoiceRecord, Source
from fintax_gst.gstr9.working import build_gstr9, classify_gstr1_doc, classify_sale, classify_summary
from fintax_gst.importers.gstr3b import Gstr3bValues as V
from fintax_gst.itc.annual import fy_months
from fintax_gst.samples import CLIENT
from fintax_gst.samples_outward import AMENDED_0013, B2CS_UNDER_REPORTED, MAY_INVOICE
from fintax_gst.services.clients import ClientService
from fintax_gst.services.gstr3b import Gstr3bService
from fintax_gst.services.gstr9 import Gstr9Service
from fintax_gst.services.imports import ImportService

D = Decimal


def _doc(no, taxable="1000", tax="180", doc=DocType.INVOICE, supply="R", rcm=False, gstin="27AABCR1122A1Z7",
         amendment=False, original="", period="2025-04", source=Source.GSTR1):
    return InvoiceRecord(uid=no, source=source, supplier_gstin=gstin, doc_type=doc, invoice_no=no,
                         invoice_date=date(2025, 4, 10), taxable=D(taxable), igst=D(tax), period=period,
                         supply_type=supply, reverse_charge=rcm, is_amendment=amendment, original_invoice_no=original)


def test_classification():
    assert classify_gstr1_doc(_doc("a")) == "4B"
    assert classify_gstr1_doc(_doc("b", supply="SEZ")) == "4D"
    assert classify_gstr1_doc(_doc("c", tax="0", supply="SEZ")) == "5B"
    assert classify_gstr1_doc(_doc("d", supply="DE")) == "4E"
    assert classify_gstr1_doc(_doc("e", rcm=True, supply="Reverse charge")) == "5C"
    assert classify_gstr1_doc(_doc("f", doc=DocType.CREDIT_NOTE)) == "4I"
    assert classify_gstr1_doc(_doc("g", doc=DocType.DEBIT_NOTE, tax="0")) == "5I"
    assert classify_sale(_doc("h", gstin="", source=Source.SALES)) == "4A"
    assert classify_sale(_doc("i", tax="0", supply="Export (LUT)", gstin="", source=Source.SALES)) == "5A"
    assert classify_sale(_doc("j", tax="0", supply="Exempt", source=Source.SALES)) == "5D"
    assert classify_summary("EXP", "WPAY", V(igst=D(18), taxable=D(100))) == "4C"
    assert classify_summary("CDNUR", "B2CL", V(igst=D(-18), taxable=D(-100))) == "4I"
    assert classify_summary("NONGST", "", V(taxable=D(5))) == "5F"


def test_amendments_and_subtotals():
    orig = _doc("INV-9", taxable="10000", tax="1800")
    up = _doc("INV-9", taxable="12000", tax="2160", amendment=True, original="INV-9", period="2025-05")
    lost = _doc("INV-7", taxable="500", tax="90", amendment=True, original="INV-7", period="2025-05")
    summaries = [{"period": "2025-04", "section": "B2CS", "supply": "INTRA", "taxable": D(4000), "igst": D(0),
                  "cgst": D(360), "sgst": D(360), "cess": D(0)},
                 {"period": "2025-05", "section": "B2CSA", "supply": "INTRA", "taxable": D(4100), "igst": D(0),
                  "cgst": D(369), "sgst": D(369), "cess": D(0)},
                 {"period": "2025-04", "section": "NIL", "supply": "INTRAB2B", "taxable": D(700), "igst": D(0),
                  "cgst": D(0), "sgst": D(0), "cess": D(0)}]
    filed = {"2025-04": {"31D": V(igst=D(50), taxable=D(1000)), "4A5": V(igst=D(900)), "4B2": V(igst=D(100))}}
    w = build_gstr9("2025-26", fy_months("2025-26"), summaries, [orig, up, lost], [], filed)
    assert w.row("4K").gstr1.taxable == D(2000) and w.row("4K").gstr1.igst == D(360)
    assert w.row("4G").gstr1.taxable == D(1000)
    assert w.row("4H").gstr1.taxable == D(4000 + 10000 + 1000)
    assert w.row("4N").gstr1.taxable == D(15000 + 2000)
    assert w.row("5N").gstr1.taxable == D(17000 + 700 - 1000)
    statuses = {a["document"]: a["status"] for a in w.amendments}
    assert statuses["INV-9"] == "included" and "manually" in statuses["INV-7"] and "manually" in statuses[
        "B2CSA total"]
    assert w.row("6A").gstr1.total == D(900) and w.row("7J").gstr1.total == D(800)
    assert [h for h, _p, _d in w.table9] == ["IGST", "CGST", "SGST", "CESS"]


def test_gstr9_on_demo_data(db, sample_files, tmp_path):
    c = ClientService(db).create(**CLIENT)
    svc = ImportService(db)
    for key, src in (("sales_april", Source.SALES), ("gstr1_april_json", Source.GSTR1),
                     ("gstr1_may_json", Source.GSTR1)):
        svc.import_file(c.id, sample_files[key], src)
    Gstr3bService(db).import_file(c.id, sample_files["gstr3b_april_json"])
    w = Gstr9Service(db).working(c.id, "2025-26")
    four_a = w.row("4A")
    assert four_a.books.taxable - four_a.gstr1.taxable == B2CS_UNDER_REPORTED
    four_l = w.row("4L").gstr1
    assert four_l.taxable == D("66000") - D(AMENDED_0013.taxable)
    assert w.row("4B").gstr1.taxable - w.row("4B").books.taxable == D(MAY_INVOICE.taxable) - D("20000") + D(
        "10000") + D("6000")  # May invoice; 0011 not in GSTR-1; 0012 only in GSTR-1; 0013 reported higher
    assert any(a["section"] == "B2CSA" and a["delta"] is None for a in w.amendments)
    assert w.row("5A").gstr1.taxable == D("250000") and w.row("5D").gstr1.taxable == D("12000")
    assert w.periods_gstr1 == ["2025-04", "2025-05"] and any("GSTR-1 not imported" in n for n in w.notes)
    from fintax_gst.reports.annual_report import export_annual

    path = export_annual(db, c.id, "2025-26", tmp_path)
    assert path.name.startswith("FINTAXPRO_GSTR9_Working_")
    rows = list(load_workbook(path, read_only=True)["GSTR-9 Tables 4-5"].iter_rows(values_only=True))
    assert any(len(r) > 2 and r[0] == "4L" and r[2] == four_l.taxable for r in rows)


def test_cli_gstr9(tmp_path, sample_files, capsys):
    dbp = str(tmp_path / "g9.db")
    assert main(["--db", dbp, "client", "add", "--gstin", CLIENT["gstin"], "--name", CLIENT["name"]]) == 0
    assert main(["--db", dbp, "import", "--client", CLIENT["gstin"], "--source", "gstr1",
                 "--file", str(sample_files["gstr1_may_json"])]) == 0
    capsys.readouterr()
    assert main(["--db", dbp, "gstr9", "--client", CLIENT["gstin"], "--fy", "2025-26",
                 "--out", str(tmp_path / "out")]) == 0
    out = capsys.readouterr().out
    assert "4N" in out and "B2BA" in out and "original not imported" in out
    assert list((tmp_path / "out").glob("FINTAXPRO_GSTR9_Working_*.xlsx"))
