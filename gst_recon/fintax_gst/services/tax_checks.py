"""Run the tax type and rate checks over a client's imported purchase and sales documents."""
from __future__ import annotations

import json
from decimal import Decimal
from pathlib import Path

from sqlalchemy import select

from ..compliance.tax_checks import KINDS, Doc, TaxIssue, check
from ..db.database import Database
from ..db.models import Client, Invoice
from ..importers.parsing import parse_period
from .common import ServiceError


def _d(v) -> Decimal:
    return Decimal(str(v or 0))


class TaxCheckService:
    def __init__(self, db: Database):
        self.db = db

    def check(self, client_id: int, period_from: str, period_to: str | None = None,
              sources: tuple[str, ...] = ("BOOKS", "SALES")) -> list[TaxIssue]:
        p_from, p_to = parse_period(period_from), parse_period(period_to or period_from)
        if not p_from or not p_to:
            raise ServiceError(f"Invalid period {period_from} – {period_to}")
        with self.db.session() as s:
            client = s.get(Client, client_id)
            if client is None:
                raise ServiceError(f"Client {client_id} not found")
            rows = list(s.scalars(select(Invoice).where(
                Invoice.client_id == client_id, Invoice.source_type.in_(sources),
                Invoice.period.between(p_from, p_to)).order_by(Invoice.source_type, Invoice.invoice_date,
                                                               Invoice.id)))
            state = client.state_code
            out: list[TaxIssue] = []
            for r in rows:
                lines = tuple((_d(h.get("taxable")), _d(h.get("igst")), _d(h.get("cgst")), _d(h.get("sgst")))
                              for h in (json.loads(r.hsn_json) if r.hsn_json else []))
                doc = Doc(r.id, r.source_type, r.supplier_gstin, r.supplier_name, r.invoice_no, r.invoice_date,
                          r.taxable, r.igst, r.cgst, r.sgst, r.place_of_supply or "", r.supply_type or "",
                          bool(r.reverse_charge), lines)
                out.extend(check(doc, state))
        return out

    @staticmethod
    def export_xlsx(issues: list[TaxIssue], path: str | Path, title: str = "") -> Path:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        wb = Workbook()
        ws = wb.active
        ws.title = "Tax checks"
        ws.append([title or "Tax type and rate checks"])
        ws["A1"].font = Font(bold=True, size=13)
        ws.append(["Register", "Check", "GSTIN", "Party", "Invoice no", "Date", "Taxable", "IGST", "CGST", "SGST",
                   "Finding"])
        for c in ws[2]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F3864")
        for i in issues:
            d = i.doc
            ws.append(["Purchase" if d.source == "BOOKS" else "Sales", KINDS[i.kind], d.gstin, d.party, d.invoice_no,
                       d.invoice_date, float(d.taxable), float(d.igst), float(d.cgst), float(d.sgst), i.message])
        for col, width in zip("ABCDEFGHIJK", (10, 22, 18, 30, 16, 12, 13, 11, 11, 11, 90)):
            ws.column_dimensions[col].width = width
        for c in ws["F"][2:]:
            c.number_format = "DD-MM-YYYY"
        ws.freeze_panes = "A3"
        out = Path(path)
        wb.save(out)
        return out
