import email
import json
import zipfile
from datetime import date, timedelta
from decimal import Decimal

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from fintax_gst.cli import main
from fintax_gst.compliance.summary import MonthStatus, assess, due_dates
from fintax_gst.core.models import Source
from fintax_gst.db.database import open_database
from fintax_gst.samples import CLIENT
from fintax_gst.samples_outward import SHORT_PAID_3B
from fintax_gst.services.auth import SETUP_SESSION, AuthService, PermissionDenied, Session
from fintax_gst.services.clients import ClientService
from fintax_gst.services.common import ServiceError
from fintax_gst.services.compliance import ComplianceService
from fintax_gst.services.followup import AttachmentService, VendorService
from fintax_gst.services.gstr3b import Gstr3bService
from fintax_gst.services.imports import ImportService
from fintax_gst.services.letters import LetterService
from fintax_gst.services.papers import WorkingPaperService
from fintax_gst.services.reconciliation import ReconciliationService
from fintax_gst.services.sync import SyncService
from fintax_gst.web.app import Signer, create_app

TODAY = date(2025, 5, 25)
PREP = Session("preparer", "accountant", "Preparer")
REVIEWER = Session("reviewer", "manager", "Reviewer")
ADMIN = Session("admin", "admin", "Admin")


def _full(db, files):
    c = ClientService(db).create(**CLIENT)
    svc = ImportService(db)
    for key, src in (("books_march", Source.BOOKS), ("books_april", Source.BOOKS),
                     ("gstr2b_april_json", Source.GSTR2B), ("sales_april", Source.SALES),
                     ("gstr1_april_json", Source.GSTR1)):
        svc.import_file(c.id, files[key], src)
    Gstr3bService(db).import_file(c.id, files["gstr3b_april_json"])
    r = ReconciliationService(db)
    inward = r.run(c.id, "2025-04")
    outward = r.run(c.id, "2025-04", compare_with=Source.GSTR1)
    return c, inward, outward


# ------------------------------------------------------------- compliance
def test_due_dates_and_month_assessment():
    assert due_dates("2025-12") == {"GSTR-1": date(2026, 1, 11), "GSTR-3B": date(2026, 1, 20)}
    m = assess(MonthStatus("2025-04"), TODAY)
    assert m.rag == "grey" and not m.issues
    m = assess(MonthStatus("2025-04", purchase_docs=5, gstr2b_docs=5), TODAY)
    assert m.rag == "red" and "not reconciled" in m.issues[0][1]
    m = MonthStatus("2025-04", purchase_docs=5, gstr2b_docs=5, gstr3b_imported=True, inward_run_id=1,
                    inward={"pending_cases": 0, "itc_at_risk": "0"}, itc_2b=Decimal(100), itc_3b=Decimal(100))
    assert assess(m, TODAY).rag == "green"
    m.itc_3b = Decimal(150)
    assert assess(m, TODAY).rag == "red" and "16(2)(aa)" in m.issues[0][1]


def test_client_summary_and_dashboard(db, sample_files, tmp_path):
    c, _inward, _outward = _full(db, sample_files)
    cc = ComplianceService(db).client_summary(c.id, "2025-26", TODAY)
    april = cc.months[0]
    assert april.period == "2025-04" and april.rag == "red"
    assert april.short_payment == SHORT_PAID_3B
    assert april.inward_run_id and april.outward_run_id and april.gstr3b_imported
    assert all(m.rag == "grey" for m in cc.months[1:])
    assert cc.months_with_data == 1 and cc.pending > 0
    with pytest.raises(ServiceError):
        ComplianceService(db).client_summary(999, "2025-26")
    from fintax_gst.reports.compliance_report import export_client_pdf, export_compliance_excel

    wb = load_workbook(export_compliance_excel([cc], "2025-26", tmp_path / "d.xlsx"), read_only=True)
    assert wb.sheetnames == ["Annual Dashboard", "Month Status", "Issues"]
    assert export_client_pdf(cc, tmp_path / "c.pdf").read_bytes().startswith(b"%PDF")


# ---------------------------------------------------------------- letters
def test_supplier_letters(db, sample_files, tmp_path):
    c, inward, outward = _full(db, sample_files)
    svc = LetterService(db)
    with pytest.raises(ServiceError):
        svc.preview(outward)
    letters = svc.preview(inward)
    assert letters and all(lt.items for lt in letters)
    issues = {i.issue for lt in letters for i in lt.items}
    assert "Not reflected in our GSTR-2B" in issues
    assert not any(i == "Difference in invoice number" for i in issues)  # number variation alone: no letter
    out = svc.generate(inward, tmp_path, letter_date=date(2025, 5, 16), user="tester")
    lt, paths = out[0]
    assert {p.suffix for p in paths} == {".docx", ".pdf", ".eml"}
    text = "\n".join(p.text for p in Document(str(paths[0])).paragraphs)
    assert CLIENT["name"] in text and "Section 16(2)(aa)" in text
    msg = email.message_from_bytes(paths[2].read_bytes())
    assert msg["X-Unsent"] == "1"
    assert any(part.get_filename() == paths[1].name for part in msg.walk())
    v = VendorService(db).get(c.id, lt.gstin)
    assert v.followup_status == "Letter issued" and v.next_followup == date(2025, 5, 31)
    assert len(svc.register(c.id)) == len(out)
    with pytest.raises(ServiceError):
        svc.generate(inward, tmp_path, formats=("rtf",))


# ---------------------------------------------------------- working papers
def test_working_papers_prepare_review_verify(db, sample_files, tmp_path):
    _c, inward, _ = _full(db, sample_files)
    ev = tmp_path / "invoice.pdf"
    ev.write_bytes(b"%PDF-1.4 evidence")
    AttachmentService(db).add(inward, 2, ev, "tester")
    svc = WorkingPaperService(db)
    with pytest.raises(PermissionDenied):
        svc.prepare(inward, tmp_path, Session("v", "viewer", ""))
    wp = svc.prepare(inward, tmp_path / "wp", PREP)
    with zipfile.ZipFile(wp.package_path) as z:
        names = z.namelist()
        manifest = json.loads(z.read("manifest.json"))
    assert names[0] == "00_Index.pdf" and "06_Audit_trail.csv" in names
    assert any(n.startswith("01_Reconciliation/") and n.endswith(".xlsx") for n in names)
    assert any(n.startswith("04_Evidence/") for n in names)
    assert manifest["run_id"] == inward and len(manifest["files"]) == len(names) - 2
    assert svc.verify(wp.id) == (True, "")
    with pytest.raises(ServiceError, match="own working papers"):
        svc.review(wp.id, Session("preparer", "manager", ""))
    with pytest.raises(PermissionDenied):
        svc.review(wp.id, Session("x", "accountant", ""))
    assert svc.review(wp.id, REVIEWER, "Checked").reviewed_by == "reviewer"
    with pytest.raises(ServiceError, match="Already reviewed"):
        svc.review(wp.id, ADMIN)
    with open(wp.package_path, "ab") as fh:  # tampering is detected
        fh.write(b"x")
    ok, problem = svc.verify(wp.id)
    assert not ok and "changed" in problem


# -------------------------------------------------------------------- sync
def test_sync_requires_authorisation_and_round_trips(db, sample_files, tmp_path):
    c, inward, _ = _full(db, sample_files)
    ev = tmp_path / "e.txt"
    ev.write_text("evidence")
    AttachmentService(db).add(inward, 3, ev, "tester")
    svc = SyncService(db)
    with pytest.raises(ServiceError, match="authorisation"):
        svc.export_package(c.id, tmp_path / "cloud", "a-long-password", ADMIN)
    with pytest.raises(ServiceError):
        svc.authorise(c.id, "", date.today() + timedelta(days=30), "ref", ADMIN)
    with pytest.raises(PermissionDenied):
        svc.authorise(c.id, "Proprietor", date.today() + timedelta(days=30), "ref", PREP)
    auth = svc.authorise(c.id, "Ganesh (Proprietor)", date.today() + timedelta(days=365), "Consent 01", ADMIN)
    with pytest.raises(ServiceError, match="10 characters"):
        svc.export_package(c.id, tmp_path / "cloud", "short", ADMIN)
    pkg = svc.export_package(c.id, tmp_path / "cloud", "a-long-password", ADMIN)
    assert pkg.read_bytes().startswith(b"FTPSY1") and CLIENT["gstin"].encode() not in pkg.read_bytes()[100:]

    other = open_database(tmp_path / "other" / "o.db")
    try:
        o = SyncService(other)
        with pytest.raises(ServiceError, match="Wrong password"):
            o.read_package(pkg, "wrong-password!")
        res = o.import_package(pkg, "a-long-password", ADMIN)
        cid = res["client_id"]
        here = ComplianceService(db).client_summary(c.id, "2025-26", TODAY)
        there = ComplianceService(other).client_summary(cid, "2025-26", TODAY)
        assert [m.issues for m in here.months] == [m.issues for m in there.months]
        runs = ReconciliationService(other).list_runs(cid)
        assert len(runs) == 2
        ctx = ReconciliationService(other).load(min(r.id for r in runs))
        assert len(ctx.result.lines) == len(ReconciliationService(db).load(inward).result.lines)
        att = AttachmentService(other).for_line(ctx.run.id, 3)
        assert att and AttachmentService.verify(att[0])
        assert o.active_authorisation(cid) is not None
        with pytest.raises(ServiceError, match="not newer"):
            o.import_package(pkg, "a-long-password", ADMIN)
        VendorService(other).upsert(cid, "27AABCM1234K1ZM", "x", remarks="changed here")
        pkg2 = svc.export_package(c.id, tmp_path / "cloud", "a-long-password", ADMIN)
        with pytest.raises(ServiceError, match="Conflict"):
            o.import_package(pkg2, "a-long-password", ADMIN)
        res2 = o.import_package(pkg2, "a-long-password", ADMIN, replace_local=True)
        assert res2["safety_backup"] and res2["replaced_local_changes"] == 1
        with pytest.raises(ServiceError, match="this computer"):
            svc.import_package(pkg2, "a-long-password", ADMIN)
    finally:
        other.dispose()
    svc.revoke(auth.id, ADMIN)
    with pytest.raises(ServiceError):
        svc.export_package(c.id, tmp_path / "cloud", "a-long-password", ADMIN)


# --------------------------------------------------------------------- web
def test_signer_rejects_tampering_and_expiry():
    s = Signer(b"k" * 32)
    tok = s.dumps({"u": "a", "exp": 4102444800})
    assert s.loads(tok)["u"] == "a"
    assert s.loads(tok[:-1] + ("0" if tok[-1] != "0" else "1")) is None
    assert s.loads(s.dumps({"u": "a", "exp": 1})) is None


def test_web_interface(db, sample_files):
    c, inward, outward = _full(db, sample_files)
    app = create_app(db, secret=b"s" * 32)
    tc = TestClient(app)
    assert tc.get("/").status_code == 200  # setup mode, local client only
    AuthService(db).create_user("admin", "Very#Strong123", "admin", by=SETUP_SESSION)
    AuthService(db).create_user("viewer1", "Very#Strong123", "viewer", by=ADMIN)
    tc = TestClient(app)
    r = tc.get("/", follow_redirects=False)
    assert r.status_code == 303 and r.headers["location"] == "/login"
    assert tc.post("/login", data={"username": "admin", "password": "nope"}).status_code == 401
    r = tc.post("/login", data={"username": "viewer1", "password": "Very#Strong123"}, follow_redirects=False)
    assert r.status_code == 303 and "httponly" in r.headers["set-cookie"].lower()
    home = tc.get("/?fy=2025-26")
    assert home.status_code == 200 and CLIENT["name"] in home.text
    assert "no-store" in home.headers["cache-control"] and "frame-ancestors" in home.headers[
        "content-security-policy"]
    assert "Action required" in tc.get(f"/client/{c.id}?fy=2025-26").text
    page = tc.get(f"/run/{outward}").text
    assert "GSTR-1 vs GSTR-3B" in page and "Customer" in page
    assert "Supplier" in tc.get(f"/run/{inward}").text
    assert tc.get("/run/9999").status_code == 404
    x = tc.get(f"/run/{inward}/report.xlsx")
    assert x.status_code == 200 and x.content[:2] == b"PK"
    assert tc.get(f"/client/{c.id}/compliance.pdf?fy=2025-26").content.startswith(b"%PDF")
    assert tc.get("/dashboard.xlsx?fy=2025-26").status_code == 200
    assert tc.get("/manifest.webmanifest").json()["display"] == "standalone"
    for method in ("post", "put", "delete"):  # read-only
        assert getattr(tc, method)(f"/run/{inward}").status_code == 405
    tc.get("/logout")
    tc.cookies.clear()
    assert tc.get("/", follow_redirects=False).status_code == 303


def test_serve_refuses_unsafe_configurations(db):
    from fintax_gst.web.app import serve

    with pytest.raises(ServiceError, match="user"):
        serve(db, host="0.0.0.0")
    AuthService(db).create_user("admin", "Very#Strong123", "admin", by=SETUP_SESSION)
    with pytest.raises(ServiceError, match="HTTPS"):
        serve(db, host="0.0.0.0")


# --------------------------------------------------------------------- CLI
def test_cli_phase4(tmp_path, sample_files, monkeypatch, capsys):
    dbp = str(tmp_path / "cli.db")
    assert main(["--db", dbp, "client", "add", "--gstin", CLIENT["gstin"], "--name", CLIENT["name"]]) == 0
    for src, key in (("books", "books_april"), ("gstr2b", "gstr2b_april_json"), ("sales", "sales_april"),
                     ("gstr1", "gstr1_april_json")):
        assert main(["--db", dbp, "import", "--client", CLIENT["gstin"], "--source", src,
                     "--file", str(sample_files[key])]) == 0
    assert main(["--db", dbp, "gstr3b", "import", "--client", CLIENT["gstin"],
                 "--file", str(sample_files["gstr3b_april_json"])]) == 0
    assert main(["--db", dbp, "reconcile", "--client", CLIENT["gstin"], "--from", "2025-04"]) == 0
    assert main(["--db", dbp, "reconcile", "--client", CLIENT["gstin"], "--from", "2025-04",
                 "--compare", "gstr1"]) == 0
    capsys.readouterr()
    assert main(["--db", dbp, "liability", "--client", CLIENT["gstin"], "--from", "2025-04"]) == 0
    assert "short payment" in capsys.readouterr().out
    assert main(["--db", dbp, "compliance", "--fy", "2025-26", "--out", str(tmp_path / "c"), "--pdf"]) == 0
    assert (tmp_path / "c" / "GST_Dashboard_2025-26.xlsx").exists()
    assert main(["--db", dbp, "letters", "--run", "1", "--out", str(tmp_path / "l"), "--formats", "pdf"]) == 0
    assert list((tmp_path / "l").rglob("*.pdf"))
    assert main(["--db", dbp, "papers", "prepare", "--run", "1", "--out", str(tmp_path / "wp")]) == 0
    assert main(["--db", dbp, "papers", "verify", "--id", "1"]) == 0
    until = (date.today() + timedelta(days=90)).isoformat()
    assert main(["--db", dbp, "sync", "authorise", "--client", CLIENT["gstin"], "--granted-by", "Proprietor",
                 "--valid-until", until, "--evidence", "Consent 7"]) == 0
    monkeypatch.setenv("FINTAX_SYNC_PASSWORD", "a-long-password")
    assert main(["--db", dbp, "sync", "export", "--client", CLIENT["gstin"], "--folder", str(tmp_path / "s")]) == 0
    pkg = next((tmp_path / "s").glob("*.ftpsync"))
    assert main(["--db", str(tmp_path / "b.db"), "sync", "import", "--file", str(pkg)]) == 0
    assert main(["--db", str(tmp_path / "b.db"), "compliance", "--fy", "2025-26"]) == 0
