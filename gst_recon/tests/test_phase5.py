import os
from datetime import date

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select

from fintax_gst.cli import main
from fintax_gst.compliance.calendar import DueOverride, annual_items, due_items
from fintax_gst.core.models import Category, Source
from fintax_gst.db.database import Database, copy_database, is_server_url, open_database
from fintax_gst.db.models import AuditLog
from fintax_gst.importers.pipeline import preview_file
from fintax_gst.importers.readers import UnsupportedFileError
from fintax_gst.samples import CLIENT, TABLE8A_LATE, TABLE8A_NOT_BOOKED
from fintax_gst.services.auth import SETUP_SESSION, AuthService, PermissionDenied, Session
from fintax_gst.services.calendar import CalendarService
from fintax_gst.services.clients import ClientService
from fintax_gst.services.common import ServiceError, audit
from fintax_gst.services.compliance import ComplianceService
from fintax_gst.services.gstr3b import Gstr3bService
from fintax_gst.services.imports import ImportService
from fintax_gst.services.itc import ItcService
from fintax_gst.services.reconciliation import ReconciliationService
from fintax_gst.services.security import SecurityService
from fintax_gst.web.app import create_app

ADMIN = Session("admin", "admin", "Admin")
SERVER_URL = os.environ.get("FINTAX_TEST_DB_URL")


def _client(db, **over):
    return ClientService(db).create(**{**CLIENT, **over})


# ---------------------------------------------------------------- calendar
def test_monthly_and_qrmp_due_dates():
    m = {i.return_type: i.due for i in due_items("2025-12", "Monthly", "27")}
    assert m == {"GSTR-1": date(2026, 1, 11), "GSTR-3B": date(2026, 1, 20)}
    q1 = due_items("2025-04", "Quarterly", "27")
    assert [(i.return_type, i.due, i.optional) for i in q1] == [("IFF", date(2025, 5, 13), True),
                                                              ("PMT-06", date(2025, 5, 25), False)]
    assert {i.return_type: i.due for i in due_items("2025-06", "Quarterly", "27")}["GSTR-3B"] == date(2025, 7, 22)
    assert {i.return_type: i.due for i in due_items("2025-06", "Quarterly", "09")}["GSTR-3B"] == date(2025, 7, 24)
    assert annual_items("2025-26")[0].due == date(2026, 12, 31)


def test_extensions_apply_only_where_notified():
    o = DueOverride("GSTR-3B", "2025-04", date(2025, 5, 27), "Notif. 9/2025-CT", frozenset({"27"}))
    mh = {i.return_type: i for i in due_items("2025-04", "Monthly", "27", [o])}["GSTR-3B"]
    ka = {i.return_type: i for i in due_items("2025-04", "Monthly", "29", [o])}["GSTR-3B"]
    assert mh.due == date(2025, 5, 27) and mh.extended and mh.basis == "Notif. 9/2025-CT"
    assert ka.due == date(2025, 5, 20) and not ka.extended


def test_calendar_service_and_compliance(db, sample_files):
    c = _client(db)
    q = _client(db, gstin="27AAACQ1234R1Z5", name="Quarterly Traders", return_frequency="Quarterly")
    cal = CalendarService(db)
    with pytest.raises(PermissionDenied):
        cal.add_override("GSTR-3B", "2025-04", date(2025, 5, 27), "ref", Session("a", "accountant"))
    with pytest.raises(ServiceError):
        cal.add_override("GSTR-3B", "2025-04", date(2025, 5, 27), "", ADMIN)
    with pytest.raises(ServiceError):
        cal.add_override("GSTR-3B", "April", date(2025, 5, 27), "ref", ADMIN)
    with pytest.raises(ServiceError, match="state code"):
        cal.add_override("GSTR-3B", "2025-04", date(2025, 5, 27), "ref", ADMIN, ["99X"])
    o = cal.add_override("GSTR-3B", "2025-04", date(2025, 5, 27), "Notif. 9/2025-CT", ADMIN, ["27"])
    ImportService(db).import_file(c.id, sample_files["books_april"], Source.BOOKS)
    ups = cal.upcoming(date(2025, 5, 21), days=10)
    mine = {(u.gstin, u.item.return_type, u.item.period): u for u in ups}
    u3b = mine[(CLIENT["gstin"], "GSTR-3B", "2025-04")]
    assert u3b.item.due == date(2025, 5, 27) and u3b.status == "not imported" and not u3b.overdue
    assert mine[(CLIENT["gstin"], "GSTR-1", "2025-04")].overdue  # due 11-05, no GSTR-1 data
    assert (q.gstin, "PMT-06", "2025-04") in mine and (q.gstin, "GSTR-3B", "2025-04") not in mine
    april = ComplianceService(db).client_summary(c.id, "2025-26", date(2025, 6, 1)).months[0]
    assert any("27-05-2025, extended – Notif. 9/2025-CT" in t for _r, t in april.issues)
    ImportService(db).import_file(q.id, sample_files["books_april"], Source.BOOKS)
    qa = ComplianceService(db).client_summary(q.id, "2025-26", date(2025, 6, 1)).months[0]
    assert not any("GSTR-3B" in t for _r, t in qa.issues)  # QRMP: no GSTR-3B for April
    cal.delete_override(o.id, ADMIN)
    assert cal.overrides() == []


# ---------------------------------------------------------------- Table 8A
def test_table8a_import_annual_and_reconciliation(db, sample_files, tmp_path):
    c = _client(db)
    svc = ImportService(db)
    p = preview_file(sample_files["gstr8a_json"], Source.GSTR8A)
    assert p.layout == "gstr8a_json" and p.period is None
    assert {r.period for r in p.records} == {"2025-04", "2025-05"}
    with pytest.raises(UnsupportedFileError):
        preview_file(sample_files["gstr8a_json"], Source.BOOKS)
    for key, src in (("books_march", Source.BOOKS), ("books_april", Source.BOOKS),
                     ("gstr2b_april_json", Source.GSTR2B)):
        svc.import_file(c.id, sample_files[key], src)
    before = ItcService(db).annual(c.id, "2025-26")
    assert before.basis == "computed"
    svc.import_file(c.id, sample_files["gstr8a_json"], Source.GSTR8A)
    t = ItcService(db).annual(c.id, "2025-26")
    assert t.basis == "portal" and t.a_computed.total == before.a.total
    assert {r.invoice_no for r in t.portal_not_in_2b} == {TABLE8A_LATE, TABLE8A_NOT_BOOKED.no}
    late = next(r for r in t.portal_not_in_2b if r.invoice_no == TABLE8A_LATE)
    assert t.a.total - before.a.total == late.total_tax + sum(
        r.total_tax for r in t.portal_not_in_2b if r.invoice_no != TABLE8A_LATE)
    assert TABLE8A_NOT_BOOKED.no in {r.invoice_no for r in t.docs_8a_not_in_books}
    assert any("Table 8A document details" in n for n in t.basis_notes())
    from fintax_gst.reports.annual_report import export_annual

    wb = load_workbook(export_annual(db, c.id, "2025-26", tmp_path), read_only=True)
    assert "8A Not in GSTR-2B" in wb.sheetnames
    run_id = ReconciliationService(db).run(c.id, "2025-04", "2026-03", compare_with=Source.GSTR8A)
    ctx = ReconciliationService(db).load(run_id)
    assert ctx.run.comparison == "BOOKS_VS_GSTR8A"
    by_no = {(ln.books or ln.gstr2b).invoice_no: ln.category for ln in ctx.result.lines}
    assert by_no[TABLE8A_LATE] == Category.FULLY_MATCHED
    assert by_no[TABLE8A_NOT_BOOKED.no] == Category.MISSING_IN_BOOKS


# ------------------------------------------------------------- database
def test_database_targets_and_copy(tmp_path):
    assert is_server_url("postgresql+psycopg://u:secret@h/db") and not is_server_url(str(tmp_path / "a.db"))
    src = open_database(tmp_path / "src.db")
    try:
        c = _client(src)
        audit_target = "x" * 200
        with src.session() as s:
            audit(s, "EXPORT", "file", audit_target, {}, "u", c.id)
        with src.session() as s:
            row = s.scalars(select(AuditLog).where(AuditLog.action == "EXPORT")).one()
            assert len(row.entity_id) == 60 and audit_target in row.details_json
        counts = copy_database(src, str(tmp_path / "dst.db"))
        assert counts["clients"] == 1
        with pytest.raises(ValueError, match="already contains data"):
            copy_database(src, str(tmp_path / "dst.db"))
        dst = open_database(tmp_path / "dst.db")
        assert ClientService(dst).list()[0].gstin == CLIENT["gstin"]
        _client(dst, gstin="27AAACQ1234R1Z5", name="Second")  # ids continue after copied rows
        dst.dispose()
    finally:
        src.dispose()


@pytest.mark.skipif(not SERVER_URL, reason="set FINTAX_TEST_DB_URL to a PostgreSQL database")
def test_server_backup_restore_and_copy(db, tmp_path):
    assert not db.is_sqlite and db.location.startswith("postgresql")
    c = _client(db)
    snap = db.backup(tmp_path / "snap.db")
    ClientService(db).create(gstin="27AAACQ1234R1Z5", name="After backup")
    db.restore(snap)
    assert [x.gstin for x in ClientService(db).list()] == [c.gstin]
    ClientService(db).create(gstin="27AAACQ1234R1Z5", name="New after restore")  # sequences reset
    local = Database(tmp_path / "local.db")
    local.migrate()
    assert copy_database(db, str(tmp_path / "copy.db"))["clients"] == 2
    local.dispose()


def test_server_location_hides_password(tmp_path):
    pytest.importorskip("psycopg")
    d = Database("postgresql+psycopg://fintax:s3cret@localhost:5432/fintax", data_dir=tmp_path)
    assert "s3cret" not in d.location and "***" in d.location and d.data_dir == tmp_path
    d.dispose()


# --------------------------------------------------------------------- web
def _web(db):
    AuthService(db).create_user("admin", "Very#Strong123", "admin", by=SETUP_SESSION)
    AuthService(db).create_user("viewer1", "Very#Strong123", "viewer", by=ADMIN)
    AuthService(db).create_user("acct1", "Very#Strong123", "accountant", by=ADMIN)
    return TestClient(create_app(db, secret=b"t" * 32))


def _login(tc, user="admin"):
    r = tc.post("/login", data={"username": user, "password": "Very#Strong123"}, follow_redirects=False)
    assert r.status_code == 303


def _csrf(html: str) -> str:
    import re

    return re.search(r'name="csrf" value="([0-9a-f]+)"', html).group(1)


def test_web_upload_preview_and_confirm(db, sample_files, tmp_path):
    c = _client(db)
    tc = _web(db)
    _login(tc)
    assert tc.get("/upload").status_code == 403  # off by default
    SecurityService(db).set_setting("web_uploads", True, ADMIN)
    form = tc.get("/upload")
    assert form.status_code == 200 and "Purchase register" in form.text
    token = _csrf(form.text)
    data = {"csrf": token, "client_id": str(c.id), "source": "BOOKS", "period": ""}
    with open(sample_files["books_april"], "rb") as fh:
        bad = tc.post("/upload", data={**data, "csrf": "0" * 64}, files={"file": ("p.xlsx", fh.read())})
    assert bad.status_code == 400
    with open(sample_files["books_april"], "rb") as fh:
        pv = tc.post("/upload", data=data, files={"file": ("purchase april.xlsx", fh.read())})
    assert pv.status_code == 200 and "Check before saving" in pv.text and "Rows rejected" in pv.text
    import re

    upload_token = re.search(r'/upload/([\w-]+)/commit', pv.text).group(1)
    other = TestClient(tc.app)  # another user, even one allowed to import, cannot use this preview
    _login(other, "acct1")
    other_csrf = _csrf(other.get("/upload").text)
    assert other.post(f"/upload/{upload_token}/commit", data={"csrf": other_csrf}).status_code == 410
    done = tc.post(f"/upload/{upload_token}/commit", data={"csrf": token})
    assert done.status_code == 200 and "Import saved" in done.text
    files = ImportService(db).list_files(c.id)
    assert len(files) == 1 and files[0].file_name == "purchase april.xlsx"
    assert os.path.exists(files[0].file_path) and str(db.data_dir) in files[0].file_path
    assert tc.post(f"/upload/{upload_token}/commit", data={"csrf": token}).status_code == 410
    # the same file again is blocked at preview
    with open(sample_files["books_april"], "rb") as fh:
        again = tc.post("/upload", data=data, files={"file": ("purchase april.xlsx", fh.read())})
    assert "already imported" in again.text and "Save import" not in again.text
    discard = re.search(r'/upload/([\w-]+)/discard', again.text).group(1)
    assert tc.post(f"/upload/{discard}/discard", data={"csrf": token}, follow_redirects=False).status_code == 303
    # GSTR-3B through the web
    with open(sample_files["gstr3b_april_json"], "rb") as fh:
        p3 = tc.post("/upload", data={**data, "source": "GSTR3B"}, files={"file": ("3b.json", fh.read())})
    t3 = re.search(r'/upload/([\w-]+)/commit', p3.text).group(1)
    assert "Import saved" in tc.post(f"/upload/{t3}/commit", data={"csrf": token}).text
    assert Gstr3bService(db).values(c.id, ["2025-04"])["2025-04"]
    # unrecognised headings point to the desktop mapping screen
    weird = tmp_path / "weird.csv"
    weird.write_text("foo,bar\n1,2\n3,4\n", encoding="utf-8")
    r = tc.post("/upload", data=data, files={"file": ("weird.csv", weird.read_bytes())})
    assert r.status_code == 422 and "desktop app" in r.text
    assert tc.post("/upload", data=data, files={"file": ("x.exe", b"MZ")}).status_code == 400
    tc.cookies.clear()
    _login(tc, "viewer1")
    assert tc.get("/upload").status_code == 403  # viewers cannot import


# --------------------------------------------------------------------- CLI
def test_cli_phase5(tmp_path, sample_files, capsys):
    dbp = str(tmp_path / "cli.db")
    assert main(["--db", dbp, "client", "add", "--gstin", CLIENT["gstin"], "--name", CLIENT["name"]]) == 0
    for src, key in (("books", "books_april"), ("gstr8a", "gstr8a_json")):
        assert main(["--db", dbp, "import", "--client", CLIENT["gstin"], "--source", src,
                     "--file", str(sample_files[key])]) == 0
    assert main(["--db", dbp, "reconcile", "--client", CLIENT["gstin"], "--fy", "2025-26",
                 "--compare", "gstr8a"]) == 0
    assert main(["--db", dbp, "calendar", "extend", "--return", "GSTR-3B", "--period", "2025-04",
                 "--due", "2025-05-27", "--reference", "Notif. 9/2025-CT", "--states", "27"]) == 0
    capsys.readouterr()
    assert main(["--db", dbp, "calendar", "--as-of", "2025-05-21", "--days", "10"]) == 0
    out = capsys.readouterr().out
    assert "27-05-2025  GSTR-3B" in out and "Notif. 9/2025-CT" in out
    assert main(["--db", dbp, "setting", "set", "web_uploads", "true"]) == 0
    assert "web_uploads              true" in capsys.readouterr().out
    assert main(["--db", dbp, "setting", "set", "nope", "1"]) == 1
    assert main(["--db", dbp, "annual", "--client", CLIENT["gstin"], "--fy", "2025-26"]) == 0
    assert "Table 8A document details" in capsys.readouterr().out
    assert main(["--db", dbp, "db-copy", "--to", str(tmp_path / "copy.db")]) == 0
