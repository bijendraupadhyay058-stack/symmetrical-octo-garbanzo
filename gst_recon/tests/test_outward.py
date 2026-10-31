from decimal import Decimal

from openpyxl import load_workbook

from fintax_gst.core.models import Source, relabel, side_labels
from fintax_gst.importers.pipeline import preview_file
from fintax_gst.outward.liability import sales_code, summary_code
from fintax_gst.samples import CLIENT
from fintax_gst.samples_outward import OUT_SCENARIOS, SHORT_PAID_3B
from fintax_gst.services.clients import ClientService
from fintax_gst.services.gstr3b import Gstr3bService
from fintax_gst.services.imports import ImportService
from fintax_gst.services.outward import OutwardService
from fintax_gst.services.reconciliation import ReconciliationService
from fintax_gst.services.reports import export_run


def _outward(db, files):
    c = ClientService(db).create(**CLIENT)
    svc = ImportService(db)
    svc.import_file(c.id, files["sales_april"], Source.SALES)
    svc.import_file(c.id, files["gstr1_april_json"], Source.GSTR1)
    Gstr3bService(db).import_file(c.id, files["gstr3b_april_json"])
    run_id = ReconciliationService(db).run(c.id, "2025-04", compare_with=Source.GSTR1)
    return c, run_id


def test_gstr1_json_preview(sample_files):
    p = preview_file(sample_files["gstr1_april_json"], Source.GSTR1)
    assert p.layout == "gstr1_json"
    sections = {s["section"] for s in p.outward_summary}
    assert {"B2B", "B2CS", "EXP", "CDNR"} <= sections


def test_summary_code_mapping():
    assert summary_code("B2B", "REG") == "31A"
    assert summary_code("B2B", "SEZ") == "31B"
    assert summary_code("B2B", "RCM") is None
    assert summary_code("B2CS", "INTRA") == "31A"
    assert summary_code("EXP", "WOPAY") == "31B"
    assert summary_code("NIL", "") == "31C"
    assert summary_code("B2BA", "REG") is None


def test_relabel_outward():
    assert side_labels("SALES_VS_GSTR1") == ("Sales register", "GSTR-1")
    assert relabel("Missing in GSTR-2B", "SALES_VS_GSTR1") == "Missing in GSTR-1"
    assert relabel("Missing in GSTR-2B", "BOOKS_VS_GSTR2B") == "Missing in GSTR-2B"


def test_sales_vs_gstr1_run_and_liability(db, sample_files):
    c, run_id = _outward(db, sample_files)
    ctx = ReconciliationService(db).load(run_id)
    assert ctx.run.comparison == "SALES_VS_GSTR1"
    by_no = {}
    for ln in ctx.result.lines:
        r = ln.books or ln.gstr2b
        by_no.setdefault(r.invoice_no, ln.category)
    for name, books, g1, expected in OUT_SCENARIOS:
        if expected is None:
            continue
        doc = (books or g1)[0].no
        assert by_no[doc] == expected, name
    # B2C, export and exempt never go through the invoice engine
    assert not any(n.startswith(("CS/", "EXP/")) for n in by_no)
    assert "SGT/25-26/0016" not in by_no

    [w] = OutwardService(db).liability(c.id, ["2025-04"])
    a = w.row("31A")
    assert a.diff_filed_gstr1 == -SHORT_PAID_3B
    assert any("short payment" in f.text for f in w.findings)
    b2c = next(s for s in w.segments if s.segment.startswith("B2C"))
    assert b2c.sales.taxable - b2c.gstr1.taxable == Decimal("1000")
    assert w.row("31B").gstr1.taxable > 0


def test_outward_workbook(db, sample_files, tmp_path):
    _c, run_id = _outward(db, sample_files)
    paths = export_run(db, run_id, tmp_path, pdf=True)
    wb = load_workbook(paths[0], read_only=True)
    names = wb.sheetnames
    assert "GSTR-1 vs GSTR-3B" in names and "GSTR-1 Section Totals" in names
    assert "Missing in GSTR-1" in names and "ITC Eligibility Review" not in names
    assert all(len(n) <= 31 for n in names)


def test_sales_code_exempt():
    from datetime import date

    from fintax_gst.core.models import DocType, InvoiceRecord

    r = InvoiceRecord(uid="x", source=Source.SALES, supplier_gstin="", doc_type=DocType.INVOICE, invoice_no="1",
                      invoice_date=date(2025, 4, 1), taxable=Decimal(100), supply_type="Exempt")
    assert sales_code(r) == "31C"
    r.supply_type = "Export (LUT)"
    assert sales_code(r) == "31B"
