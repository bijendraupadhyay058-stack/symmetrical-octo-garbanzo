"""Low-level file readers. They return raw cell values; nothing is interpreted here."""
from __future__ import annotations

import csv
import io
import json
import zipfile
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook

MAX_ZIP_MEMBER_BYTES = 200 * 1024 * 1024
SUPPORTED_EXTENSIONS = {".xlsx", ".xlsm", ".csv", ".json", ".zip", ".xml"}


class UnsupportedFileError(ValueError):
    pass


@dataclass
class RawTable:
    name: str  # sheet name (Excel) or file name (CSV)
    rows: list[tuple[int, list]]  # (1-based row number in the source, cell values)


def read_excel(data: bytes | Path) -> list[RawTable]:
    source = io.BytesIO(data) if isinstance(data, (bytes, bytearray)) else str(data)
    wb = load_workbook(source, read_only=True, data_only=True)
    tables = []
    try:
        for ws in wb.worksheets:
            rows = []
            for r_index, row in enumerate(ws.iter_rows(values_only=True), start=1):
                values = list(row)
                while values and values[-1] in (None, ""):
                    values.pop()
                rows.append((r_index, values))
            tables.append(RawTable(ws.title, rows))
    finally:
        wb.close()
    return tables


def read_csv(data: bytes | Path, name: str) -> RawTable:
    # Python's own csv module (no pandas / numpy): every cell stays text, exactly as in the file.
    raw = data if isinstance(data, (bytes, bytearray)) else Path(data).read_bytes()
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover - latin-1 decodes any bytes
        raise UnsupportedFileError("Could not read CSV: unknown text encoding")
    sample = "\n".join(text.splitlines()[:50])
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel  # a single column, or nothing to tell from: comma-separated
    try:
        parsed = list(csv.reader(io.StringIO(text, newline=""), dialect))
    except csv.Error as exc:
        raise UnsupportedFileError(f"Could not read CSV: {exc}") from None
    rows = []
    for i, values in enumerate(parsed, start=1):
        vals = [v if v != "" else None for v in values]
        while vals and vals[-1] is None:
            vals.pop()
        rows.append((i, vals))
    return RawTable(name, rows)


def read_json(data: bytes | Path) -> dict:
    raw = data if isinstance(data, (bytes, bytearray)) else Path(data).read_bytes()
    try:
        return json.loads(raw.decode("utf-8-sig"), parse_float=str)
    except (UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise UnsupportedFileError(f"Invalid JSON file: {exc}") from exc


def zip_members(path: Path) -> list[tuple[str, bytes]]:
    """Return supported members of a ZIP export (no extraction to disk, no path traversal)."""
    members = []
    with zipfile.ZipFile(path) as zf:
        for info in sorted(zf.infolist(), key=lambda i: i.filename):
            if info.is_dir():
                continue
            suffix = Path(info.filename).suffix.lower()
            if suffix not in {".xlsx", ".xlsm", ".csv", ".json"}:
                continue
            if info.file_size > MAX_ZIP_MEMBER_BYTES:
                raise UnsupportedFileError(f"ZIP member {info.filename} is too large")
            members.append((Path(info.filename).name, zf.read(info)))
    if not members:
        raise UnsupportedFileError("ZIP file contains no supported Excel, CSV or JSON file")
    return members
