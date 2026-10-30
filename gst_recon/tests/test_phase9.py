from datetime import date
from decimal import Decimal

import pytest
from openpyxl import load_workbook

from fintax_gst.cli import main
from fintax_gst.core.models import DocType, InvoiceRecord, Source
from fintax_gst.gstr9.working import build_gstr9
from fintax_gst.importers.pipeline import preview_file
from fintax_gst.itc.annual import fy_months
from fintax_gst.samples import CLIENT, ITEMWISE_HSN, books_docs
from fintax_gst.services.auth import PermissionDenied, Session
from fintax_gst.services.clients import ClientService
from fintax_gst.services.common import ServiceError
from fintax_gst.services.gstr9 import Gstr9Service
from fintax_gst.services.gstr9c import Gstr9cService
from fintax_gst.services.imports import ImportService

D = Decimal
ADMIN = Session("admin", "admin", "Admin")


def test_itemwise_register_merges_hsn_lines(sample_files, tmp_path):
    p = preview_file(sample_files["books_itemwise_csv"], Source.BOOKS)
    assert len(p.records) == 5 and p.count("duplicate") == 0
    first = p.records[0]
    assert [h["hsn"] for h in first.hsn_lines] == [code for code, *_ in ITEMWISE_HSN]
    assert first.taxable == D(books_docs()[0].taxable)
    # the same HSN twice at the same rate is still a repeat, not a second line
    rep = tmp_path / "rep.csv"
    rep.write_text("Supplier Name,GSTIN,Bill No,Bill Date,HSN/SAC,Taxable Value,CGST,SGST\n"
                   "A,27AABCM1234K1ZM,X-1,05-04-2025,7208,100,9,9\nA,27AABCM1234K1ZM,X-1,05-04-2025,7208,100,9,9\n",
                   encoding="utf-8")
    assert preview_file(rep, Source.BOOKS).count("duplicate") == 1


def test_table18_from_purchase_register(db, sample_files):
    c = ClientService(db).create(**CLIENT)
    ImportService(db).import_file(c.id, sample_files["books_itemwise_csv"], Source.BOOKS)
    w = Gstr9Service(db).working(c.id, "2025-26")
    assert [g["hsn"] for g in w.table18] == ["7208", "7318"]
    total = sum(g["taxable"] for g in w.table18)
    assert total == sum(D(d.taxable) for d in books_docs()[:5])
    assert w.table18_without_hsn == (0, D(0))
    # documents without HSN (Tally register) are counted and named in the notes
    ImportService(db).import_file(c.id, sample_files["books_march"], Source.BOOKS)
    w2 = Gstr9Service(db).working(c.id, "2024-25")
    assert w2.table18 == [] and w2.table18_without_hsn[0] == 2
    assert any("have no HSN" in n for n in w2.notes)


def test_table18_credit_note_sign():
    inv = InvoiceRecord(uid="a", source=Source.BOOKS, supplier_gstin="27AABCM1234K1ZM", doc_type=DocType.INVOICE,
                        invoice_no="1", invoice_date=date(2025, 5, 1), taxable=D(100), period="2025-05",
                        hsn_lines=[{"hsn": "7208", "uqc": "KGS", "qty": "10", "taxable": "100", "igst": "18"}])
    cn = InvoiceRecord(uid="b", source=Source.BOOKS, supplier_gstin="27AABCM1234K1ZM", doc_type=DocType.CREDIT_NOTE,
                       invoice_no="2", invoice_date=date(2025, 5, 2), taxable=D(40), period="2025-05",
                       hsn_lines=[{"hsn": "7208", "uqc": "KGS", "qty": "4", "taxable": "40", "igst": "7.2"}])
    w = build_gstr9("2025-26", fy_months("2025-26"), [], [], [], {}, purchases=[inv, cn])
    assert w.table18[0]["taxable"] == D(60) and w.table18[0]["qty"] == D(6) and w.table18[0]["igst"] == D("10.8")


def test_tables_15_16_entered(db, sample_files, tmp_path):
    c = ClientService(db).create(**CLIENT)
    svc = Gstr9cService(db)
    with pytest.raises(PermissionDenied):
        svc.set_input(c.id, "2025-26", "15E-I", "1", "", Session("v", "viewer"))
    with pytest.raises(ServiceError):
        svc.set_input(c.id, "2025-26", "16A-I", "1", "", ADMIN)  # 16A has a taxable value only
    svc.set_input(c.id, "2025-26", "15E-I", "25000", "DRC-07 dated 12-01-2026", ADMIN)
    svc.set_input(c.id, "2025-26", "15F-I", "10000", "", ADMIN)
    svc.set_input(c.id, "2025-26", "16A-T", "45000", "Purchases from composition dealers", ADMIN)
    w = Gstr9Service(db).working(c.id, "2025-26")
    t15 = {x.code: x for x in w.table15}
    assert t15["15E"].values["I"] == D(25000) and t15["15E"].notes == ["DRC-07 dated 12-01-2026"]
    assert set(t15["15A"].values) == {"C", "S", "I", "X", "N", "P", "L"}
    t16 = {x.code: x for x in w.table16}
    assert t16["16A"].values == {"T": D(45000)} and set(t16["16B"].values) == {"T", "C", "S", "I", "X"}
    from fintax_gst.reports.annual_report import export_annual

    rows = list(load_workbook(export_annual(db, c.id, "2025-26", tmp_path), read_only=True)["GSTR-9 Tables 15-16"]
                .iter_rows(values_only=True))
    assert any(len(r) > 4 and r[0] == "15E" and r[4] == D(25000) for r in rows)


def test_cli_phase9(tmp_path, sample_files, capsys):
    dbp = str(tmp_path / "p9.db")
    assert main(["--db", dbp, "client", "add", "--gstin", CLIENT["gstin"], "--name", CLIENT["name"]]) == 0
    assert main(["--db", dbp, "import", "--client", CLIENT["gstin"], "--source", "books",
                 "--file", str(sample_files["books_itemwise_csv"])]) == 0
    capsys.readouterr()
    assert main(["--db", dbp, "gstr9", "set", "--client", CLIENT["gstin"], "--fy", "2025-26", "--item", "15a-c",
                 "--amount", "1200", "--note", "Refund RFD-01"]) == 0
    out = capsys.readouterr().out
    assert "15A-C recorded" in out and "Table 18 HSN 7208" in out and "Table 15A" in out
