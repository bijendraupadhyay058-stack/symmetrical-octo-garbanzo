"""Phase 11: interest and late fees."""
from datetime import date
from decimal import Decimal

import pytest

from fintax_gst.cli import main
from fintax_gst.compliance.interest import gstr9_late_fee, interest, late_fee
from fintax_gst.core.models import Source
from fintax_gst.samples import CLIENT
from fintax_gst.services.auth import SETUP_SESSION, PermissionDenied, Session
from fintax_gst.services.clients import ClientService
from fintax_gst.services.common import ServiceError
from fintax_gst.services.gstr3b import Gstr3bService
from fintax_gst.services.gstr9c import Gstr9cService
from fintax_gst.services.imports import ImportService
from fintax_gst.services.interest import InterestService

AS_OF = date(2025, 10, 31)


def test_interest_and_late_fee_rules():
    assert interest(Decimal("100000"), date(2025, 5, 20), date(2025, 6, 19)) == (30, Decimal("1479.45"))
    assert interest(Decimal("100000"), date(2025, 5, 20), date(2025, 5, 20))[1] == 0  # paid on the due date
    assert interest(Decimal("-5"), date(2025, 5, 20), date(2025, 6, 19))[1] == 0
    due = date(2025, 5, 20)
    fee = late_fee("2025-04", due, date(2025, 6, 9), False, Decimal("12000000"))
    assert (fee.days, fee.cgst, fee.sgst) == (20, 500, 500)
    assert late_fee("2025-04", due, date(2025, 9, 9), False, Decimal("12000000")).cgst == 1000  # ≤ 1.5 cr cap
    assert late_fee("2025-04", due, date(2025, 12, 9), False, Decimal("40000000")).cgst == 2500
    assert late_fee("2025-04", due, date(2026, 3, 9), False, Decimal("90000000")).cgst == 5000
    assert late_fee("2025-04", due, date(2026, 3, 9), True, None).cgst == 250  # nil return
    assert late_fee("2025-04", due, date(2025, 6, 9), True, None).cgst == 200  # 20 days × ₹10
    assert late_fee("2021-05", due, date(2025, 6, 9), False, None) is None  # before June 2021: not computed
    g9 = gstr9_late_fee("2024-25", date(2025, 12, 31), date(2026, 1, 30), Decimal("30000000"), Decimal("31000000"))
    assert (g9.days, g9.cgst) == (30, 750)  # ₹25/day, cap 0.02% of 3.1 cr = 6,200
    assert gstr9_late_fee("2024-25", date(2025, 12, 31), date(2026, 12, 30), Decimal("30000000"),
                          Decimal("31000000")).cgst == 6200
    assert gstr9_late_fee("2024-25", date(2025, 12, 31), date(2026, 1, 30), Decimal("15000000"), Decimal(0)) is None
    assert gstr9_late_fee("2021-22", date(2022, 12, 31), date(2023, 1, 30), None, Decimal(0)) is None


def _client(db, files):
    c = ClientService(db).create(**CLIENT)
    svc = ImportService(db)
    for key, src in (("books_april", Source.BOOKS), ("gstr2b_april_json", Source.GSTR2B),
                     ("sales_april", Source.SALES), ("gstr1_april_json", Source.GSTR1)):
        svc.import_file(c.id, files[key], src)
    Gstr3bService(db).import_file(c.id, files["gstr3b_april_json"])
    return c


def test_interest_working(db, sample_files):
    c = _client(db, sample_files)
    svc = InterestService(db)
    rows, problems = svc.import_filings(c.id, sample_files["filings"], SETUP_SESSION)
    assert len(rows) == 2 and not problems
    w = svc.working(c.id, "2025-26", AS_OF)
    late = {r.return_type: r for r in w.late}
    assert (late["GSTR-1"].days, late["GSTR-1"].fee_total) == (4, 200)
    assert (late["GSTR-3B"].days, late["GSTR-3B"].fee_total) == (8, 400)
    assert late["GSTR-3B"].interest_total == 0 and late["GSTR-3B"].interest_paid == 120  # nothing paid in cash
    assert any("AATO" in n for n in w.notes)
    short = {e.head: e for e in w.exposures if e.kind == "SHORT_PAYMENT"}
    assert set(short) == {"cgst", "sgst"} and short["cgst"].amount == 1800
    assert (short["cgst"].days, short["cgst"].interest, short["cgst"].status) == (164, Decimal("145.58"), "OPEN")
    excess = [e for e in w.exposures if e.kind == "EXCESS_ITC"]
    assert excess and all("UPPER LIMIT" in e.basis and e.start == date(2025, 5, 28) for e in excess)
    d = w.drc03()
    assert d["cgst"]["tax"] == 1800 and d["cgst"]["fee"] == 300

    # the AATO slab lowers the cap; paying the short payment stops the interest
    Gstr9cService(db).set_input(c.id, "2025-26", "AATO", "12000000", "previous year turnover", SETUP_SESSION)
    svc.set_settlement(c.id, "SHORT_PAYMENT", "2025-04", "SETTLED", SETUP_SESSION, date(2025, 6, 19),
                       "DRC-03 AD2706250001")
    svc.set_settlement(c.id, "EXCESS_ITC", "2025-04", "NOT_PAYABLE", SETUP_SESSION,
                       note="ITC on bill of entry, not in GSTR-2B")
    w = svc.working(c.id, "2025-26", AS_OF)
    short = next(e for e in w.exposures if e.kind == "SHORT_PAYMENT" and e.head == "cgst")
    assert (short.days, short.interest, short.status) == (30, Decimal("26.63"), "SETTLED")
    assert all(e.status == "NOT_PAYABLE" and e.interest == 0 for e in w.exposures if e.kind == "EXCESS_ITC")
    assert w.drc03()["cgst"]["tax"] == 0 and w.drc03()["cgst"]["interest"] == Decimal("26.63")
    assert not any("AATO" in n for n in w.notes)

    with pytest.raises(ServiceError):
        svc.set_settlement(c.id, "EXCESS_ITC", "2025-04", "NOT_PAYABLE", SETUP_SESSION)  # reason needed
    with pytest.raises(PermissionDenied):
        svc.set_filing(c.id, "GSTR-3B", "2025-05", date(2025, 6, 20), Session("v", "viewer", "V"))


def test_no_false_excess_without_gstr2b(db, sample_files):
    c = ClientService(db).create(**CLIENT)
    Gstr3bService(db).import_file(c.id, sample_files["gstr3b_april_json"])
    w = InterestService(db).working(c.id, "2025-26", AS_OF)
    assert not w.exposures and any("GSTR-2B 2025-04 not imported" in n for n in w.notes)


def test_cli_interest(tmp_path, sample_files, capsys):
    dbp = str(tmp_path / "cli.db")
    assert main(["--db", dbp, "client", "add", "--name", CLIENT["name"], "--gstin", CLIENT["gstin"]]) == 0
    for src, key in (("sales", "sales_april"), ("gstr1", "gstr1_april_json")):
        assert main(["--db", dbp, "import", "--client", "1", "--source", src, "--file", str(sample_files[key])]) == 0
    assert main(["--db", dbp, "gstr3b", "import", "--client", "1", "--file",
                 str(sample_files["gstr3b_april_json"])]) == 0
    assert main(["--db", dbp, "interest", "filed", "--client", "1", "--return", "GSTR-3B", "--period", "2025-04",
                 "--date", "2025-05-28"]) == 0
    capsys.readouterr()
    assert main(["--db", dbp, "interest", "--client", "1", "--fy", "2025-26", "--as-of", "2025-10-31", "--out",
                 str(tmp_path)]) == 0
    out = capsys.readouterr().out
    assert "8 day(s)" in out and "Tax short paid" in out and "DRC-03" in out
    from openpyxl import load_workbook

    wb = load_workbook(next(tmp_path.glob("FINTAXPRO_Interest_LateFee_*.xlsx")), read_only=True)
    assert wb.sheetnames == ["Interest & Late Fee"]


# ------------------------------------------------------------------ supplier scores
def _rec(gstin, no, d, period, source=Source.BOOKS, itc="1800", **kw):
    from fintax_gst.core.models import DocType, InvoiceRecord

    return InvoiceRecord(uid=no, source=source, supplier_gstin=gstin, doc_type=DocType.INVOICE, invoice_no=no,
                         invoice_date=d, taxable=Decimal("10000"), igst=Decimal(itc), period=period,
                         supplier_name=f"Supplier {gstin[:4]}", **kw)


def test_scoring_rules():
    from fintax_gst.core.models import Category, ReconLine
    from fintax_gst.itc.supplier_score import score_suppliers

    g_bad, g_ok = "27AAAAA0000A1Z5", "27BBBBB0000B1Z5"
    books = [_rec(g_bad, f"B{i}", date(2025, 4, i + 1), "2025-04", db_id=i) for i in range(4)]
    books += [_rec(g_ok, f"K{i}", date(2025, 4, i + 1), "2025-04", db_id=10 + i) for i in range(4)]
    g2b = [_rec(g_bad, "B0", date(2025, 4, 1), "2025-06", Source.GSTR2B),  # two months late
           _rec(g_bad, "B1", date(2025, 4, 2), "2025-04", Source.GSTR2B, itc_availability="N")]
    g2b += [_rec(g_ok, f"K{i}", date(2025, 4, i + 1), "2025-04", Source.GSTR2B) for i in range(4)]
    april = [ReconLine(1, Category.MISSING_IN_2B, books=books[0]), ReconLine(2, Category.MISSING_IN_2B, books=books[1]),
             ReconLine(3, Category.MISSING_IN_2B, books=books[2]), ReconLine(4, Category.TAXABLE_DIFFERENCE,
                                                                              books=books[3])]
    april += [ReconLine(10 + i, Category.FULLY_MATCHED, books=books[4 + i]) for i in range(4)]
    june = [ReconLine(1, Category.PREVIOUS_PERIOD, books=books[0])]  # B0 found in June: no longer missing
    s = {x.gstin: x for x in score_suppliers(books, g2b, [april, june], {g_bad: "Letter sent"})}
    bad = s[g_bad]
    assert (bad.missing_docs, bad.itc_at_risk, bad.late_docs, bad.max_lag) == (2, 3600, 1, 2)
    assert (bad.mismatched, bad.unavailable) == (1, 1)
    # 40×2/4 + 25×1/2 + 20×1/4 + 10×1/2 = 20 + 12 (12.5 rounded half to even) + 5 + 5
    assert bad.score == 100 - 20 - 12 - 5 - 5 and bad.band == "High" and bad.followup == "Letter sent"
    assert len(bad.reasons) == 4 and "ITC ₹3,600.00 at risk" in bad.reasons[0]
    assert (s[g_ok].score, s[g_ok].band, s[g_ok].reasons) == (100, "Low", ["No issues found in the data imported"])
    assert list(s)[0] == g_bad  # highest risk first


def test_supplier_scores_service_and_cli(db, sample_files, tmp_path, capsys):
    from fintax_gst.services.reconciliation import ReconciliationService
    from fintax_gst.services.supplier_score import SupplierScoreService

    c = ClientService(db).create(**CLIENT)
    svc = ImportService(db)
    for key, src in (("books_april", Source.BOOKS), ("gstr2b_april_json", Source.GSTR2B)):
        svc.import_file(c.id, sample_files[key], src)
    rows, notes = SupplierScoreService(db).scores(c.id, "2025-26")
    assert any("No books vs GSTR-2B reconciliation for 2025-04" in n for n in notes)
    assert all(x.missing_docs == 0 for x in rows)
    recon = ReconciliationService(db)
    recon.run(c.id, "2025-04")
    rows, notes = SupplierScoreService(db).scores(c.id, "2025-26")
    by = {x.name: x for x in rows}
    assert not notes and by["Kovai Textiles Pvt Ltd"].missing_docs == 1
    assert by["Delhi Office Supplies"].late_docs == 1
    path = SupplierScoreService(db).export(c.id, "2025-26", tmp_path)
    from openpyxl import load_workbook

    assert load_workbook(path, read_only=True).sheetnames == ["Supplier Scores"]


def test_table4_uses_the_gstr2b_run(db, sample_files):
    """A later books vs GSTR-2A run of the same month must not replace the GSTR-2B run."""
    from fintax_gst.db.models import ReconRun
    from fintax_gst.services.itc import ItcService
    from fintax_gst.services.reconciliation import ReconciliationService

    c = ClientService(db).create(**CLIENT)
    svc = ImportService(db)
    for key, src in (("books_april", Source.BOOKS), ("gstr2b_april_json", Source.GSTR2B)):
        svc.import_file(c.id, sample_files[key], src)
    r2b = ReconciliationService(db).run(c.id, "2025-04")
    with db.session() as s:
        s.add(ReconRun(client_id=c.id, financial_year="2025-26", period_from="2025-04", period_to="2025-04",
                       comparison="BOOKS_VS_GSTR2A", settings_json="{}", fingerprint="x", engine_version="x",
                       summary_json="{}", source_file_ids="[]", status="COMPLETED"))
    with db.session() as s:
        assert ItcService(db)._latest_month_runs(s, c.id, ["2025-04"]) == {"2025-04": r2b}


# ------------------------------------------------------------------ GST notices
def test_notice_register_and_drc01b_reply(db, sample_files, tmp_path):
    from docx import Document
    from docx.enum.text import WD_COLOR_INDEX

    from fintax_gst.services.notices import NoticeService

    c = _client(db, sample_files)
    svc = NoticeService(db)
    with pytest.raises(PermissionDenied):
        svc.add(c.id, Session("v", "viewer", "V"), form="DRC-01B", notice_date=date(2025, 9, 15),
                period_from="2025-04")
    with pytest.raises(ServiceError):
        svc.add(c.id, SETUP_SESSION, form="DRC-99", notice_date=date(2025, 9, 15), period_from="2025-04")
    nid = svc.add(c.id, SETUP_SESSION, form="DRC-01B", reference="ZD2709250001", notice_date=date(2025, 9, 15),
                  period_from="2025-04", amount="3600", reply_due=date(2025, 10, 15))
    assert [n.id for n in svc.due(c.id, on=date(2025, 10, 10))] == [nid]

    draft, _n, _c = svc.build(nid, AS_OF)
    assert draft.difference == 3600
    text = "\n".join(p for s in draft.sections for p in s.paragraphs)
    assert "the figures agree" in text
    # internal findings and the unpaid interest are only working notes, never reply text
    paragraphs = [p for s in draft.sections for p in s.paragraphs]
    notes = [p for p in paragraphs if p.startswith("[[Working note")]
    assert not any("possible short payment" in p or "sales register" in p for p in paragraphs if p not in notes)
    assert any("possible short payment" in p for p in notes) and any("interest under section 50" in p for p in notes)

    # once the payment is recorded, the reply states it with the DRC-03 reference
    from fintax_gst.services.interest import InterestService

    InterestService(db).set_settlement(c.id, "SHORT_PAYMENT", "2025-04", "SETTLED", SETUP_SESSION, date(2025, 6, 19),
                                       "DRC-03 AD2706250001")
    path, draft = svc.draft(nid, tmp_path, AS_OF, signatory="R. Kulkarni")
    text = "\n".join(p for s in draft.sections for p in s.paragraphs)
    assert "was paid on 19-06-2025 vide DRC-03 AD2706250001" in text
    doc = Document(path)
    body = "\n".join(p.text for p in doc.paragraphs)
    assert "Form GST DRC-01B ZD2709250001" in body and "R. Kulkarni" in body and "[[" not in body
    highlighted = [r.text for p in doc.paragraphs for r in p.runs if r.font.highlight_color == WD_COLOR_INDEX.YELLOW]
    assert any(t.startswith("Working note") for t in highlighted)
    assert doc.tables[0].rows[1].cells[3].text == "3,600.00"

    svc.update(nid, SETUP_SESSION, status="REPLIED")
    n = svc.get(nid)
    assert n.status == "REPLIED" and n.replied_on is not None and not svc.due(c.id, on=date(2025, 10, 10))


def test_drc01c_and_asmt10_replies(db, sample_files, tmp_path):
    from fintax_gst.services.notices import NoticeService
    from fintax_gst.services.reconciliation import ReconciliationService

    c = _client(db, sample_files)
    ReconciliationService(db).run(c.id, "2025-04")
    svc = NoticeService(db)
    n1 = svc.add(c.id, SETUP_SESSION, form="DRC-01C", notice_date=date(2025, 9, 15), period_from="2025-04",
                 amount="1")
    d1, _n, _c = svc.build(n1, AS_OF)
    headings = [s.heading for s in d1.sections]
    assert "Input tax credit: GSTR-3B compared with GSTR-2B" in headings
    assert "Output tax: GSTR-1 / IFF compared with GSTR-3B" not in headings
    assert any("[[explain the variation]]" in p for s in d1.sections for p in s.paragraphs)  # amount differs
    assert "[[notice reference / DIN]]" in d1.subject
    n2 = svc.add(c.id, SETUP_SESSION, form="ASMT-10", notice_date=date(2025, 9, 15), period_from="2025-04")
    d2, _n, _c = svc.build(n2, AS_OF)
    headings = [s.heading for s in d2.sections]
    assert {"Output tax: GSTR-1 / IFF compared with GSTR-3B", "Input tax credit: GSTR-3B compared with GSTR-2B",
            "Other discrepancies in the notice"} <= set(headings)
    # a period without returns imported is not invented
    n3 = svc.add(c.id, SETUP_SESSION, form="DRC-01B", notice_date=date(2025, 9, 15), period_from="2025-08")
    d3, _n, _c = svc.build(n3, AS_OF)
    assert any("not imported" in p for s in d3.sections for p in s.paragraphs)


def test_cli_notice(tmp_path, sample_files, capsys):
    dbp = str(tmp_path / "cli.db")
    assert main(["--db", dbp, "client", "add", "--name", CLIENT["name"], "--gstin", CLIENT["gstin"]]) == 0
    assert main(["--db", dbp, "notice", "add", "--client", "1", "--form", "DRC-01B", "--reference", "ZD1",
                 "--date", "2025-09-15", "--from", "2025-04", "--amount", "3600", "--reply-due", "2025-10-15"]) == 0
    assert main(["--db", dbp, "notice", "draft", "--client", "1", "--id", "1", "--out", str(tmp_path)]) == 0
    out = capsys.readouterr().out
    assert "Draft reply:" in out and "highlighted item(s)" in out
    assert main(["--db", dbp, "notice", "--client", "1"]) == 0
    assert "DRC-01B" in capsys.readouterr().out
