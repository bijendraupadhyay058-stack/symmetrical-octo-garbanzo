"""PAN-level view: one business with GST registrations in several states (same PAN), for a financial year.

Per registration (GSTIN): ITC as per books / GSTR-2B (all documents imported for the year, credit notes
deducted, reconciled or not) / GSTR-3B, ITC at risk (latest reconciliation of each
month), output tax as per GSTR-1 / GSTR-3B, months reconciled, returns overdue (no filing date recorded after
the due date) and open notices – with totals for the PAN. Suppliers to follow up are combined across the
registrations, so a supplier that is behind with several of them is seen once.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal
from pathlib import Path

from sqlalchemy import func, select

from ..core.gstin import STATE_CODES
from ..core.money import ZERO
from ..db.database import Database
from ..db.models import Client, GstNotice, Invoice, ReturnFiling
from ..itc.annual import fy_months
from ..engine.summary import supplier_summary
from .common import ServiceError


@dataclass
class Registration:
    client_id: int
    gstin: str
    name: str
    state: str
    months_reconciled: int = 0
    itc_books: Decimal = ZERO
    itc_2b: Decimal = ZERO
    itc_3b: Decimal = ZERO
    itc_at_risk: Decimal = ZERO
    output_gstr1: Decimal = ZERO
    output_3b: Decimal = ZERO
    returns_overdue: int = 0
    notices_open: int = 0
    rag: str = "grey"


@dataclass
class SupplierAcross:
    gstin: str
    name: str
    itc_at_risk: Decimal = ZERO
    missing_in_2b: int = 0
    registrations: set[str] = field(default_factory=set)


@dataclass
class PanView:
    pan: str
    fy: str
    registrations: list[Registration]
    suppliers: list[SupplierAcross]

    def total(self, attr: str) -> Decimal | int:
        return sum((getattr(r, attr) for r in self.registrations), ZERO if attr != "returns_overdue" else 0)


class PanViewService:
    def __init__(self, db: Database):
        self.db = db

    def groups(self, include_single: bool = False) -> list[tuple[str, list[Client]]]:
        """PANs with their active registrations (only PANs with several GSTINs unless ``include_single``)."""
        with self.db.session() as s:
            clients = list(s.scalars(select(Client).where(Client.active.is_(True)).order_by(Client.pan,
                                                                                             Client.state_code)))
        by_pan: dict[str, list[Client]] = {}
        for c in clients:
            by_pan.setdefault(c.pan, []).append(c)
        return [(p, cs) for p, cs in sorted(by_pan.items(), key=lambda x: x[1][0].name.lower())
                if include_single or len(cs) > 1]

    def view(self, pan: str, fy: str, today: date | None = None) -> PanView:
        from .compliance import ComplianceService
        from .reconciliation import ReconciliationService

        today = today or date.today()
        pan = pan.strip().upper()
        with self.db.session() as s:
            clients = list(s.scalars(select(Client).where(Client.pan == pan, Client.active.is_(True))
                                     .order_by(Client.state_code)))
            if not clients:
                raise ServiceError(f"No active client with PAN {pan}")
            ids = [c.id for c in clients]
            notices = dict(s.execute(select(GstNotice.client_id, func.count()).where(
                GstNotice.client_id.in_(ids), GstNotice.status == "OPEN").group_by(GstNotice.client_id)).all())
            filed = {(f.client_id, f.return_type, f.period) for f in s.scalars(
                select(ReturnFiling).where(ReturnFiling.client_id.in_(ids)))}
            itc: dict[tuple[int, str], Decimal] = {}
            for cid, src, doc_type, *taxes in s.execute(  # amounts are exact text: added up here, not in SQL
                    select(Invoice.client_id, Invoice.source_type, Invoice.doc_type, Invoice.igst, Invoice.cgst,
                           Invoice.sgst, Invoice.cess)
                    .where(Invoice.client_id.in_(ids), Invoice.source_type.in_(("BOOKS", "GSTR2B")),
                           Invoice.period.in_(fy_months(fy)))):
                # the buyer's purchase return is the supplier's credit note: it reduces the ITC
                sign = -1 if doc_type == "CRN" else 1
                itc[(cid, src)] = itc.get((cid, src), ZERO) + sign * sum(taxes, ZERO)
            info = [(c.id, c.gstin, c.name, STATE_CODES.get(c.state_code, c.state_code)) for c in clients]
        comp, recon = ComplianceService(self.db), ReconciliationService(self.db)
        regs, suppliers = [], {}
        for cid, g, name, state in info:
            cc = comp.client_summary(cid, fy, today)
            r = Registration(cid, g, name, state, itc_books=itc.get((cid, "BOOKS"), ZERO),
                             itc_2b=itc.get((cid, "GSTR2B"), ZERO), notices_open=notices.get(cid, 0), rag=cc.rag)
            for m in cc.months:
                if m.inward:
                    r.months_reconciled += 1
                    r.itc_at_risk += Decimal(str(m.inward.get("itc_at_risk", 0)))
                r.itc_3b += m.itc_3b or ZERO
                r.output_gstr1 += m.output_gstr1 or ZERO
                r.output_3b += m.output_3b or ZERO
                r.returns_overdue += sum(1 for d in m.due if not d.optional and d.due < today
                                         and (cid, d.return_type, m.period) not in filed)
                if m.inward_run_id:
                    for sr in supplier_summary(recon.load(m.inward_run_id).result.lines, "exposure"):
                        if sr.itc_at_risk <= 0 and not sr.missing_in_2b:
                            continue
                        a = suppliers.setdefault(sr.gstin, SupplierAcross(sr.gstin, sr.name))
                        a.name = a.name or sr.name
                        a.itc_at_risk += sr.itc_at_risk
                        a.missing_in_2b += sr.missing_in_2b
                        a.registrations.add(g)
            regs.append(r)
        ranked = sorted(suppliers.values(), key=lambda a: (-a.itc_at_risk, a.gstin))
        return PanView(pan, fy, regs, ranked)

    @staticmethod
    def export_xlsx(view: PanView, path: str | Path, business: str = "") -> Path:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        head = PatternFill("solid", fgColor="1F3864")
        wb = Workbook()
        ws = wb.active
        ws.title = "Registrations"
        ws.append([f"PAN {view.pan}{' – ' + business if business else ''} – FY {view.fy}"])
        ws["A1"].font = Font(bold=True, size=13)
        cols = ["State", "GSTIN", "Name", "Months reconciled", "ITC as per books", "ITC as per GSTR-2B",
                "ITC claimed in GSTR-3B", "ITC at risk", "Output tax GSTR-1", "Output tax GSTR-3B",
                "Returns overdue", "Open notices"]
        ws.append(cols)
        for c in ws[2]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), head
        for r in view.registrations:
            ws.append([r.state, r.gstin, r.name, r.months_reconciled, float(r.itc_books), float(r.itc_2b),
                       float(r.itc_3b), float(r.itc_at_risk), float(r.output_gstr1), float(r.output_3b),
                       r.returns_overdue, r.notices_open])
        ws.append(["Total", "", "", sum(r.months_reconciled for r in view.registrations),
                   *[float(view.total(a)) for a in ("itc_books", "itc_2b", "itc_3b", "itc_at_risk", "output_gstr1",
                                                    "output_3b")],
                   view.total("returns_overdue"), sum(r.notices_open for r in view.registrations)])
        for c in ws[ws.max_row]:
            c.font = Font(bold=True)
        for i, w in enumerate((18, 17, 32, 10, 15, 15, 15, 13, 15, 15, 10, 10), start=1):
            ws.column_dimensions[ws.cell(2, i).column_letter].width = w
        sup = wb.create_sheet("Suppliers to follow up")
        sup.append(["Supplier", "GSTIN", "ITC at risk", "Invoices not in GSTR-2B", "Your registrations affected"])
        for c in sup[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), head
        for a in view.suppliers:
            sup.append([a.name, a.gstin, float(a.itc_at_risk), a.missing_in_2b, ", ".join(sorted(a.registrations))])
        for col, w in zip("ABCDE", (34, 17, 13, 12, 40)):
            sup.column_dimensions[col].width = w
        out = Path(path)
        wb.save(out)
        return out
