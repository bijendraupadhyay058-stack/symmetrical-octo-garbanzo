"""Phase 10: e-invoice (IRP) and e-way bill reconciliation."""
import csv
from collections import Counter

import pytest

from fintax_gst.cli import main
from fintax_gst.core.models import Category, ExceptionStatus, Source, side_labels
from fintax_gst.importers.pipeline import preview_file
from fintax_gst.importers.readers import UnsupportedFileError
from fintax_gst.samples import CLIENT
from fintax_gst.samples_outward import (
    EINV_CANCELLED, EINV_MATCHED, EINV_NOT_IN_BOOKS, EINV_NOT_REPORTED, EWB_MATCHED, EWB_MISSING, EWB_NOT_IN_BOOKS,
)
from fintax_gst.services.auth import SETUP_SESSION
from fintax_gst.services.clients import ClientService
from fintax_gst.services.imports import ImportService
from fintax_gst.services.reconciliation import ReconciliationService
from fintax_gst.services.reports import export_run
from fintax_gst.services.security import SecurityService


def _client(db, files, *sources):
    c = ClientService(db).create(**CLIENT)
    svc = ImportService(db)
    svc.import_file(c.id, files["sales_april"], Source.SALES)
    keys = {Source.EINV: "einv_april", Source.EWB: "ewb_april", Source.GSTR1: "gstr1_april_json"}
    for src in sources:
        svc.import_file(c.id, files[keys[src]], src)
    return c


def _why(ln):
    return " ".join(ln.reasons)


def _by_no(ctx):
    return {(ln.books or ln.gstr2b).invoice_no: ln for ln in ctx.result.lines}


def test_einvoice_report_import(sample_files):
    pv = preview_file(sample_files["einv_april"], Source.EINV)
    assert pv.layout == "einvoice_report"
    assert {r.invoice_no for r in pv.records} == {*EINV_MATCHED, EINV_CANCELLED, EINV_NOT_IN_BOOKS}
    export = next(r for r in pv.records if r.invoice_no == "EXP/25-26/001")
    assert export.supplier_gstin == "" and len(export.irn) == 64 and export.ref_status == "ACTIVE"
    cancelled = next(r for r in pv.records if r.invoice_no == EINV_CANCELLED)
    assert cancelled.ref_status == "CANCELLED" and cancelled.period == "2025-04"
    assert any("cancelled" in i.message for i in pv.issues if i.severity == "warning")
    # a register is not an e-invoice report, and an e-invoice report is not a GSTR-2B
    with pytest.raises(UnsupportedFileError):
        preview_file(sample_files["gstr2b_april"], Source.EINV)


def test_ewb_report_import_ignores_challans_and_inward(sample_files):
    pv = preview_file(sample_files["ewb_april"], Source.EWB)
    assert pv.layout == "ewb_report"
    ignored = [i.message for i in pv.issues if i.severity == "ignored"]
    assert len(ignored) == 2 and any("Delivery Challan" in m for m in ignored) and any("Inward" in m for m in ignored)
    # the cancelled and the regenerated e-way bill of one invoice are two records, not a duplicate
    two = [r for r in pv.records if r.invoice_no == "SGT/25-26/0002"]
    assert sorted(r.ref_status for r in two) == ["ACTIVE", "CANCELLED"] and not pv.duplicate_rows
    assert all(len(r.ref_no) == 12 for r in pv.records)


def test_sales_vs_einvoice(db, sample_files):
    c = _client(db, sample_files, Source.EINV)
    svc = ReconciliationService(db)
    ctx = svc.load(svc.run(c.id, "2025-04", compare_with=Source.EINV))
    assert ctx.run.comparison == "SALES_VS_EINV"
    assert side_labels(ctx.run.comparison) == ("Sales register", "e-Invoice (IRP)")
    lines = _by_no(ctx)
    assert {n for n, ln in lines.items() if ln.category == Category.FULLY_MATCHED} == set(EINV_MATCHED)
    cancelled = lines[EINV_CANCELLED]
    assert cancelled.category == Category.MISSING_IN_2B and cancelled.status == ExceptionStatus.CLIENT_CLARIFICATION
    assert "cancelled" in _why(cancelled)
    assert lines[EINV_NOT_REPORTED].category == Category.MISSING_IN_2B
    assert "No e-invoice" in _why(lines[EINV_NOT_REPORTED])
    assert lines[EINV_NOT_IN_BOOKS].category == Category.MISSING_IN_BOOKS
    # B2C and exempt sales are not e-invoiced: out of scope
    assert not any(n.startswith("CS/") or n == "SGT/25-26/0016" for n in lines)


def test_sales_vs_ewb_threshold_and_notes(db, sample_files):
    c = _client(db, sample_files, Source.EWB)
    svc = ReconciliationService(db)
    ctx = svc.load(svc.run(c.id, "2025-04", compare_with=Source.EWB))
    lines = _by_no(ctx)
    assert {n for n, ln in lines.items() if ln.category == Category.FULLY_MATCHED} == set(EWB_MATCHED)
    assert lines[EWB_MISSING].status == ExceptionStatus.PENDING_REVIEW and "No e-way bill" in _why(lines[EWB_MISSING])
    assert lines[EWB_NOT_IN_BOOKS.no].category == Category.MISSING_IN_BOOKS
    no_action = [n for n, ln in lines.items() if ln.status == ExceptionStatus.NO_ACTION
                 and ln.category == Category.MISSING_IN_2B]
    assert len(no_action) == 12  # 8 cash sales, 0011, 0014, the exempt sale and the credit note
    assert "threshold" in _why(lines["CS/25-26/001"]) and "Credit" in _why(lines["SGT-CN/25-26/001"])
    # nothing is hidden: every sales document is on a line
    assert Counter(ln.category for ln in ctx.result.lines)[Category.MISSING_IN_2B] == 13

    # a lower threshold makes the small documents exceptions again
    SecurityService(db).set_setting("ewb_threshold", 20000, SETUP_SESSION)
    lines = _by_no(svc.load(svc.run(c.id, "2025-04", compare_with=Source.EWB)))
    assert lines["SGT/25-26/0011"].status != ExceptionStatus.NO_ACTION  # ₹23,600
    assert lines["CS/25-26/001"].status == ExceptionStatus.NO_ACTION

    # a status chosen by the reviewer is kept on a re-run; a rule's status follows the rule
    ctx = svc.load(max(r.id for r in svc.list_runs(c.id)))
    line = _by_no(ctx)["CS/25-26/002"]
    svc.update_line(ctx.run.id, line.line_no, status=ExceptionStatus.CLOSED, remarks="checked", user="preparer")
    SecurityService(db).set_setting("ewb_threshold", 5000, SETUP_SESSION)
    lines = _by_no(svc.load(svc.run(c.id, "2025-04", compare_with=Source.EWB)))
    assert lines["CS/25-26/002"].status == ExceptionStatus.CLOSED and lines["CS/25-26/002"].remarks == "checked"
    assert lines["CS/25-26/003"].status == ExceptionStatus.PENDING_REVIEW  # ₹5,605 ≥ ₹5,000


def test_latest_import_of_a_reference_wins(db, sample_files, tmp_path):
    c = _client(db, sample_files, Source.EWB)
    # a later report shows the e-way bill of SGT/25-26/0001 cancelled
    with open(sample_files["ewb_april"], encoding="utf-8") as fh:
        rows = list(csv.reader(fh))
    later = tmp_path / "ewb_later.csv"
    with later.open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(rows[0])
        w.writerow(rows[1][:-1] + ["Cancelled"])
    ImportService(db).import_file(c.id, later, Source.EWB)
    svc = ReconciliationService(db)
    line = _by_no(svc.load(svc.run(c.id, "2025-04", compare_with=Source.EWB)))["SGT/25-26/0001"]
    assert line.category == Category.MISSING_IN_2B and "cancelled" in _why(line)


def test_review_work_stays_with_its_comparison(db, sample_files):
    c = _client(db, sample_files, Source.EINV, Source.EWB, Source.GSTR1)
    svc = ReconciliationService(db)
    g1 = svc.load(svc.run(c.id, "2025-04", compare_with=Source.GSTR1))
    line = _by_no(g1)["SGT/25-26/0011"]  # not in GSTR-1 either
    svc.update_line(g1.run.id, line.line_no, status=ExceptionStatus.BOOKS_CORRECTION, remarks="GSTR-1 amendment due",
                    user="preparer")
    ewb = _by_no(svc.load(svc.run(c.id, "2025-04", compare_with=Source.EWB)))["SGT/25-26/0011"]
    assert ewb.status == ExceptionStatus.NO_ACTION and not ewb.remarks
    again = _by_no(svc.load(svc.run(c.id, "2025-04", compare_with=Source.GSTR1)))["SGT/25-26/0011"]
    assert again.status == ExceptionStatus.BOOKS_CORRECTION and again.remarks == "GSTR-1 amendment due"


def test_einvoice_report_export(db, sample_files, tmp_path):
    c = _client(db, sample_files, Source.EINV)
    run_id = ReconciliationService(db).run(c.id, "2025-04", compare_with=Source.EINV)
    paths = export_run(db, run_id, tmp_path, pdf=False)
    from openpyxl import load_workbook

    wb = load_workbook(paths[0], read_only=True)
    assert not any("3.1" in n or "Liability" in n for n in wb.sheetnames)


def test_cli_einvoice(tmp_path, sample_files, capsys):
    dbp = str(tmp_path / "cli.db")
    assert main(["--db", dbp, "client", "add", "--name", CLIENT["name"], "--gstin", CLIENT["gstin"]]) == 0
    for src, key in (("sales", "sales_april"), ("einv", "einv_april"), ("ewb", "ewb_april")):
        assert main(["--db", dbp, "import", "--client", "1", "--source", src, "--file",
                     str(sample_files[key])]) == 0
    capsys.readouterr()
    assert main(["--db", dbp, "reconcile", "--client", "1", "--from", "2025-04", "--compare", "ewb"]) == 0
    out = capsys.readouterr().out
    assert "e-Way bill" in out or "Sales register" in out


# ------------------------------------------------------------------ TDS / TCS credits
def test_tds_tcs_read_and_statement(db, sample_files):
    from fintax_gst.outward.tds_tcs import read_credits
    from fintax_gst.samples_outward import ECO, PMC
    from fintax_gst.services.tds_tcs import TdsTcsService

    lines, problems = read_credits(sample_files["tds_tcs"])
    assert not problems and Counter(c.kind for c in lines) == {"TDS": 4, "TCS": 2}
    tcs = next(c for c in lines if c.kind == "TCS" and c.period == "2025-04")
    assert tcs.value == 100000 and tcs.action == "ACCEPTED" and tcs.tax == 500

    c = _client(db, sample_files)
    svc = TdsTcsService(db)
    pv = svc.import_file(c.id, sample_files["tds_tcs"])
    assert len(pv.new) == 6
    st = svc.statement(c.id, "2025-26")
    ven = next(p for p in st.parties if p.name == "Venkatesh Traders")
    assert ven.sales_taxable == ven.value == 120000 and ven.accepted == 1200 and ven.pending == 1200
    texts = [(f.severity, f.party, f.text) for f in st.findings]
    assert texts[0][0] == "high" and texts[0][1] == PMC[0] and "no sales" in texts[0][2]
    assert any("expected 0.5%" in t and p == ECO[0] for _s, p, t in texts)  # May TCS at the old 1% rate
    assert any("CGST/SGST" in t and "another state" in t for _s, _p, t in texts)
    assert sum(1 for s_, _p, t in texts if "not yet accepted" in t) == 3
    assert st.total("TDS", "ACCEPTED") == 1200 and st.total("TCS", "PENDING") == 800

    # a later download: the same credits, the May TDS now accepted – no duplicates, the action is updated
    with pytest.raises(Exception):
        svc.import_file(c.id, sample_files["tds_tcs"])  # nothing new
    from openpyxl import load_workbook

    later = sample_files["tds_tcs"].with_name("tds_later.xlsx")
    wb = load_workbook(sample_files["tds_tcs"])
    wb["TDS Credit"]["H5"] = "Accepted"
    wb.save(later)
    pv = svc.import_file(c.id, later)
    assert not pv.new and len(pv.changed) == 1 and pv.changed[0][1] == "PENDING"
    assert svc.statement(c.id, "2025-26").total("TDS", "PENDING") == 6000

    # removing the import removes its credits; the purge removes them all
    from fintax_gst.db.models import SourceFile
    from sqlalchemy import select

    with db.session() as s:
        first = s.scalar(select(SourceFile.id).where(SourceFile.source_type == "TDSTCS").order_by(SourceFile.id))
    ImportService(db).remove_import(c.id, first)
    assert not svc.statement(c.id, "2025-26").lines


def test_tds_tcs_in_annual_workbook_and_cli(tmp_path, sample_files, capsys):
    from openpyxl import load_workbook

    dbp = str(tmp_path / "cli.db")
    assert main(["--db", dbp, "client", "add", "--name", CLIENT["name"], "--gstin", CLIENT["gstin"]]) == 0
    assert main(["--db", dbp, "import", "--client", "1", "--source", "sales", "--file",
                 str(sample_files["sales_april"])]) == 0
    capsys.readouterr()
    assert main(["--db", dbp, "tds-tcs", "import", "--client", "1", "--fy", "2025-26", "--file",
                 str(sample_files["tds_tcs"]), "--preview"]) == 0
    assert "preview only" in capsys.readouterr().out
    assert main(["--db", dbp, "tds-tcs", "import", "--client", "1", "--fy", "2025-26", "--file",
                 str(sample_files["tds_tcs"])]) == 0
    out = capsys.readouterr().out
    assert "6 new credit(s)" in out and "[HIGH]" in out and "Pune Municipal Corporation" in out
    assert main(["--db", dbp, "gstr9", "--client", "1", "--fy", "2025-26", "--out", str(tmp_path)]) == 0
    wb = load_workbook(next(tmp_path.glob("FINTAXPRO_GSTR9_Working_*.xlsx")), read_only=True)
    assert wb.sheetnames[-1] == "TDS-TCS Credits"


# ------------------------------------------------------------------ GSTR-9 draft JSON
def test_gstr9_draft_json(db, sample_files, tmp_path):
    import json

    from fintax_gst.gstr9.offline_json import STATUS_COMPUTED, STATUS_MANUAL
    from fintax_gst.services.gstr3b import Gstr3bService
    from fintax_gst.services.gstr9 import Gstr9Service
    from fintax_gst.services.gstr9c import Gstr9cService

    c = _client(db, sample_files, Source.GSTR1)
    svc = ImportService(db)
    for key, src in (("books_april", Source.BOOKS), ("gstr2b_april_json", Source.GSTR2B),
                     ("gstr1_may_json", Source.GSTR1)):
        svc.import_file(c.id, sample_files[key], src)
    Gstr3bService(db).import_file(c.id, sample_files["gstr3b_april_json"])
    Gstr9cService(db).set_input(c.id, "2025-26", "15E-I", "25000", "DRC-07", SETUP_SESSION)
    gsvc = Gstr9Service(db)
    w = gsvc.working(c.id, "2025-26")
    js, mp = gsvc.draft_json(c.id, "2025-26", tmp_path)
    data = json.loads(js.read_text(encoding="utf-8"))
    assert data["gstin"] == CLIENT["gstin"] and data["fp"] == "032026"
    rows = {r.code: r for r in w.table4}
    assert data["table4"]["b2b"]["txval"] == float(rows["4B"].gstr1.taxable)
    assert data["table4"]["at"]["txval"] == 50000  # GSTR-1 May advance (4F)
    assert data["table5"]["zero_rtd"]["txval"] == 250000
    assert data["table15"] == {"tax_dmnd": {"camt": 0.0, "samt": 0.0, "iamt": 25000.0, "csamt": 0.0, "intr": 0.0,
                                            "pen": 0.0, "fee": 0.0}}
    assert data["table17"]["items"] and "table18" not in data
    # nothing is guessed: split rows and sub-totals are not in the JSON, the mapping says why
    assert "supp_non_rchrg" not in data["table6"] and "rule42" not in data["table7"]
    text = mp.read_text(encoding="utf-8-sig")
    assert text.startswith('"DRAFT for review')
    status = {r[0]: r[4] for r in csv.reader(text.splitlines()[2:])}
    assert status["6B"] == status["7A–7G"] == STATUS_MANUAL and status["4N"] == STATUS_COMPUTED
    assert status["4B"] == status["15E"] == "included"


def test_gstr9_json_cli(tmp_path, sample_files, capsys):
    dbp = str(tmp_path / "cli.db")
    assert main(["--db", dbp, "client", "add", "--name", CLIENT["name"], "--gstin", CLIENT["gstin"]]) == 0
    assert main(["--db", dbp, "import", "--client", "1", "--source", "gstr1", "--file",
                 str(sample_files["gstr1_april_json"])]) == 0
    capsys.readouterr()
    assert main(["--db", dbp, "gstr9", "json", "--client", "1", "--fy", "2025-26", "--out", str(tmp_path / "j")]) == 0
    out = capsys.readouterr().out
    assert "DRAFT for review" in out and (tmp_path / "j" / f"FINTAXPRO_GSTR9_DRAFT_{CLIENT['gstin']}_FY2025-26.json").exists()


# ------------------------------------------------------------------ multi-user
def _inward_run(db, sample_files):
    c = ClientService(db).create(**CLIENT)
    svc = ImportService(db)
    svc.import_file(c.id, sample_files["books_april"], Source.BOOKS)
    svc.import_file(c.id, sample_files["gstr2b_april_json"], Source.GSTR2B)
    return c, ReconciliationService(db).run(c.id, "2025-04", user="preparer")


def test_concurrent_edit_is_refused_not_lost(db, sample_files):
    from fintax_gst.services.common import ConflictError

    _c, run_id = _inward_run(db, sample_files)
    svc = ReconciliationService(db)
    a_view = svc.load(run_id).result.lines[0]  # two people open the same line
    b_view = svc.load(run_id).result.lines[0]
    assert a_view.version == b_view.version == 1
    v = svc.update_line(run_id, a_view.line_no, "asha", remarks="Asha: called supplier", expected_version=a_view.version)
    assert v == 2
    with pytest.raises(ConflictError, match="changed by asha"):
        svc.update_line(run_id, b_view.line_no, "ravi", remarks="Ravi: overwrite", expected_version=b_view.version)
    line = svc.load(run_id).result.lines[0]
    assert line.remarks == "Asha: called supplier" and line.version == 2
    # after reloading, Ravi can save on top of Asha's change
    assert svc.update_line(run_id, line.line_no, "ravi", remarks="Ravi: noted", expected_version=line.version) == 3
    # manual decisions bump the version too, so a stale save after them is refused
    stale = svc.load(run_id).result.lines[0]
    svc.mark_pending_clarification(run_id, stale.line_no, "asha", "ask client")
    with pytest.raises(ConflictError):
        svc.update_line(run_id, stale.line_no, "ravi", remarks="late", expected_version=stale.version)


def test_racing_saves_keep_every_edit(db, sample_files):
    """Threads saving the same lines: each save either wins (audited, version +1) or is refused."""
    import threading

    from fintax_gst.db.models import AuditLog, ReconLineRow
    from fintax_gst.services.common import ConflictError
    from sqlalchemy import func, select

    _c, run_id = _inward_run(db, sample_files)
    svc = ReconciliationService(db)
    lines = [ln.line_no for ln in svc.load(run_id).result.lines[:2]]
    wins, refused, errors = Counter(), Counter(), []

    def worker(n):
        for i in range(8):
            line = lines[(n + i) % 2]
            with db.session() as s:
                version = s.scalar(select(ReconLineRow.version).where(ReconLineRow.run_id == run_id,
                                                                      ReconLineRow.line_no == line))
            try:
                svc.update_line(run_id, line, f"u{n}", remarks=f"u{n}-{i}", expected_version=version)
                wins[line] += 1
            except ConflictError:
                refused[line] += 1
            except Exception as exc:  # pragma: no cover - reported below
                errors.append(repr(exc))

    threads = [threading.Thread(target=worker, args=(n,)) for n in range(4)]
    for t in threads:
        t.start()
    for t in threads:
        t.join()
    assert not errors
    with db.session() as s:
        versions = dict(s.execute(select(ReconLineRow.line_no, ReconLineRow.version).where(
            ReconLineRow.run_id == run_id, ReconLineRow.line_no.in_(lines))).all())
        audited = s.scalar(select(func.count()).select_from(AuditLog).where(AuditLog.action == "UPDATE",
                                                                            AuditLog.entity == "recon_line"))
    assert all(versions[ln] == 1 + wins[ln] for ln in lines)
    assert audited == sum(wins.values()) and sum(wins.values()) + sum(refused.values()) == 32


def test_user_activity(db, sample_files):
    from fintax_gst.services.auth import PermissionDenied, Session

    _c, run_id = _inward_run(db, sample_files)
    svc = ReconciliationService(db)
    for i, ln in enumerate(svc.load(run_id).result.lines[:3]):
        svc.update_line(run_id, ln.line_no, "asha" if i < 2 else "ravi", remarks="checked")
    sec = SecurityService(db)
    with pytest.raises(PermissionDenied):
        sec.activity(Session("ravi", "accountant", "Ravi"))
    rows = {u.user: u for u in sec.activity(Session("partner", "manager", "Partner"))}
    assert rows["asha"].edits == 2 and rows["ravi"].edits == 1 and rows["preparer"].actions["RECONCILE"] == 1
    assert [u.user for u in sec.activity(Session("p", "manager", "P"), user="ravi")] == ["ravi"]


def test_cli_activity_and_expect_version(tmp_path, sample_files, capsys):
    dbp = str(tmp_path / "cli.db")
    assert main(["--db", dbp, "client", "add", "--name", CLIENT["name"], "--gstin", CLIENT["gstin"]]) == 0
    for src, key in (("books", "books_april"), ("gstr2b", "gstr2b_april_json")):
        assert main(["--db", dbp, "import", "--client", "1", "--source", src, "--file", str(sample_files[key])]) == 0
    assert main(["--db", dbp, "reconcile", "--client", "1", "--from", "2025-04"]) == 0
    capsys.readouterr()
    assert main(["--db", dbp, "show", "--run", "1"]) == 0
    assert "] v1" in capsys.readouterr().out
    assert main(["--db", dbp, "update-line", "--run", "1", "--line", "1", "--remarks", "a", "--expect-version",
                 "1"]) == 0
    assert "version 2" in capsys.readouterr().out
    assert main(["--db", dbp, "update-line", "--run", "1", "--line", "1", "--remarks", "b", "--expect-version",
                 "1"]) != 0
    assert "changed by" in capsys.readouterr().err
    assert main(["--db", dbp, "activity", "--days", "7"]) == 0
    assert "review edits" in capsys.readouterr().out


def test_load_script_sqlite(tmp_path, capsys):
    import importlib.util
    from pathlib import Path

    spec = importlib.util.spec_from_file_location("concurrency_load",
                                                  Path(__file__).resolve().parents[1] / "scripts" / "concurrency_load.py")
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    assert mod.main(["--db", str(tmp_path / "load.db"), "--workers", "3", "--edits", "5"]) == 0
    assert "no lost or duplicated edits" in capsys.readouterr().out
