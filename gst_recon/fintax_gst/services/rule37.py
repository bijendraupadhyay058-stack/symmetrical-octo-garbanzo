"""180-day payment rule (CGST Rule 37): ITC on bills not paid within 180 days of the invoice date.

Payments (payment register or marked by hand) are matched with the supplier's purchase bills: a payment that
names a bill (bill-wise reference = supplier invoice number) settles that bill, the rest settle the oldest
open bills first. The supplier's credit notes (purchase returns) reduce the amount payable like a payment.

For each bill (reverse-charge bills are outside the rule):

* **Not yet due** – the 180 days have not ended.
* **Paid in time** – paid in full by day 180.
* **Reverse ITC** – the part unpaid on day 180 carries ITC that is to be reversed (GSTR-3B Table 4(B)(2)) in
  the return for the month after the month in which the 180 days end, with interest u/s 50.
* **Paid late – re-claim** – paid after day 180: the ITC reversed can be claimed again in the period of
  payment.

The interest shown is an estimate at 18 % a year on the ITC to reverse, from the due date of the GSTR-3B in
which the bill was booked to the due date of the reversal return (or the payment date, if earlier);
interest is payable only to the extent the credit was utilised. All figures are working papers for review.
"""
from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, timedelta
from decimal import Decimal
from pathlib import Path

from sqlalchemy import select

from ..core.invoice_no import normalize_invoice_no
from ..core.models import DocType
from ..core.money import ZERO, q2
from ..db.database import Database
from ..db.models import Client, Invoice, SourceFile, SupplierPayment
from ..importers.payments import party_key, read_payments
from ..importers.pipeline import sha256_of
from .common import ServiceError, audit

DAYS = 180
RATE = Decimal("0.18")
TOLERANCE = Decimal("1.00")
STATUSES = ("Reverse ITC", "Paid late – re-claim", "Not yet due", "Paid in time")


def _next_month(period: str) -> str:
    y, m = int(period[:4]), int(period[5:])
    return f"{y + (m == 12):04d}-{m % 12 + 1:02d}"


def _gstr3b_due(period: str) -> date:
    n = _next_month(period)
    return date(int(n[:4]), int(n[5:]), 20)


@dataclass
class BillStatus:
    invoice_id: int
    supplier_gstin: str
    supplier: str
    invoice_no: str
    invoice_date: date
    booked_period: str
    value: Decimal
    itc: Decimal
    igst: Decimal
    cgst: Decimal
    sgst: Decimal
    cess: Decimal
    payments: list[tuple[date, Decimal]] = field(default_factory=list)
    status: str = ""
    day180: date | None = None
    unpaid_at_day180: Decimal = ZERO
    unpaid_now: Decimal = ZERO
    itc_to_reverse: Decimal = ZERO
    reversal_period: str = ""
    paid_in_full_on: date | None = None
    itc_to_reclaim: Decimal = ZERO
    reclaim_period: str = ""
    interest: Decimal = ZERO

    def share(self, amount: Decimal) -> dict[str, Decimal]:
        """Tax heads of the ITC proportionate to ``amount`` of the bill value."""
        f = amount / self.value if self.value else ZERO
        return {k: q2(getattr(self, k) * f) for k in ("igst", "cgst", "sgst", "cess")}


class Rule37Service:
    def __init__(self, db: Database):
        self.db = db

    # ---------------------------------------------------------------- payments
    def import_payments(self, client_id: int, path: str | Path, user: str | None = None):
        """Import a payment register; returns (source file, the read result)."""
        path = Path(path)
        result = read_payments(path)
        digest = sha256_of(path)
        with self.db.session() as s:
            if s.get(Client, client_id) is None:
                raise ServiceError(f"Client {client_id} not found")
            dup = s.scalar(select(SourceFile.id).where(SourceFile.client_id == client_id, SourceFile.sha256 == digest,
                                                       SourceFile.source_type == "PAYMENTS"))
            if dup:
                raise ServiceError(f"{path.name} was already imported (import #{dup})")
            sf = SourceFile(client_id=client_id, source_type="PAYMENTS", file_name=path.name,
                            file_path=str(path.resolve()), sha256=digest, layout="payment_register", period=None,
                            rows_read=len(result.rows) + len(result.ignored), documents_imported=len(result.rows),
                            rows_rejected=0, warnings_json="[]", imported_by=user or "")
            s.add(sf)
            s.flush()
            for r in result.rows:
                s.add(SupplierPayment(client_id=client_id, source_file_id=sf.id, party_name=r.party[:200],
                                      party_key=party_key(r.party)[:200], supplier_gstin=r.gstin, pay_date=r.pay_date,
                                      amount=r.amount, voucher_no=r.voucher_no[:60], bill_ref=r.bill_ref[:100],
                                      created_by=user or ""))
            audit(s, "IMPORT", "source_file", sf.id, {"file": path.name, "payments": len(result.rows),
                                                      "total": str(result.total), "ignored": len(result.ignored)},
                  user, client_id)
            s.expunge(sf)
        return sf, result

    def mark_paid(self, client_id: int, invoice_id: int, pay_date: date, amount: Decimal | None = None,
                  note: str = "", user: str | None = None) -> None:
        """Record a payment of a bill by hand (default: the amount still unpaid)."""
        with self.db.session() as s:
            inv = s.get(Invoice, invoice_id)
            if inv is None or inv.client_id != client_id or inv.source_type != "BOOKS":
                raise ServiceError("Bill not found for this client")
        if amount is None:
            bill = next((b for b in self.track(client_id, as_of=pay_date) if b.invoice_id == invoice_id), None)
            amount = bill.unpaid_now if bill else ZERO
        if amount <= 0:
            raise ServiceError("Nothing unpaid on this bill")
        with self.db.session() as s:
            s.add(SupplierPayment(client_id=client_id, party_name=inv.supplier_name[:200],
                                  party_key=party_key(inv.supplier_name)[:200], supplier_gstin=inv.supplier_gstin,
                                  pay_date=pay_date, amount=amount, bill_ref=inv.invoice_no[:100],
                                  invoice_id=invoice_id, note=note or "marked paid by hand", created_by=user or ""))
            audit(s, "MARK_PAID", "invoices", invoice_id, {"date": pay_date.isoformat(), "amount": str(amount),
                                                           "note": note}, user, client_id)

    # ---------------------------------------------------------------- tracker
    def track(self, client_id: int, as_of: date | None = None) -> list[BillStatus]:
        as_of = as_of or date.today()
        with self.db.session() as s:
            docs = list(s.scalars(select(Invoice).where(Invoice.client_id == client_id, Invoice.source_type == "BOOKS",
                                                        Invoice.reverse_charge.is_(False))
                                  .order_by(Invoice.invoice_date, Invoice.id)))
            pays = list(s.scalars(select(SupplierPayment).where(SupplierPayment.client_id == client_id)
                                  .order_by(SupplierPayment.pay_date, SupplierPayment.id)))
        # supplier of each payment: GSTIN, else the party name as in the purchase register
        names: dict[str, str] = {}
        for d in docs:
            names.setdefault(party_key(d.supplier_name), d.supplier_gstin)
        bills: dict[str, list[BillStatus]] = defaultdict(list)
        credits: dict[str, list[tuple[date, Decimal]]] = defaultdict(list)
        for d in docs:
            value = d.invoice_value if d.invoice_value else d.taxable + d.igst + d.cgst + d.sgst + d.cess
            if d.doc_type == DocType.CREDIT_NOTE.value:  # supplier's credit note: reduces what is payable
                credits[d.supplier_gstin].append((d.invoice_date, abs(value)))
                continue
            if d.doc_type != DocType.INVOICE.value:
                continue
            bills[d.supplier_gstin].append(BillStatus(
                d.id, d.supplier_gstin, d.supplier_name, d.invoice_no, d.invoice_date, d.period, abs(value),
                d.igst + d.cgst + d.sgst + d.cess, d.igst, d.cgst, d.sgst, d.cess))
        by_supplier: dict[str, list[tuple[date, Decimal, str, int | None]]] = defaultdict(list)
        for p in pays:
            g = p.supplier_gstin or names.get(p.party_key) or ""
            if g:
                by_supplier[g].append((p.pay_date, p.amount, p.bill_ref, p.invoice_id))
        for g, items in credits.items():
            by_supplier[g] += [(dt, amt, "", None) for dt, amt in items]
        out: list[BillStatus] = []
        for g, open_bills in bills.items():
            _allocate(open_bills, sorted(by_supplier.get(g, []), key=lambda x: x[0]))
            for b in open_bills:
                _assess(b, as_of)
                out.append(b)
        return sorted(out, key=lambda b: (STATUSES.index(b.status), b.day180 or as_of, b.supplier))

    def unmatched_payments(self, client_id: int) -> list[SupplierPayment]:
        """Payments whose party is in no purchase bill (name spelt differently, or not a supplier)."""
        with self.db.session() as s:
            keys = {party_key(n) for n in s.scalars(select(Invoice.supplier_name).where(
                Invoice.client_id == client_id, Invoice.source_type == "BOOKS").distinct())}
            gstins = set(s.scalars(select(Invoice.supplier_gstin).where(Invoice.client_id == client_id,
                                                                         Invoice.source_type == "BOOKS").distinct()))
            return [p for p in s.scalars(select(SupplierPayment).where(SupplierPayment.client_id == client_id))
                    if p.supplier_gstin not in gstins and p.party_key not in keys]

    def record_reversals(self, client_id: int, bills: list[BillStatus], user: str | None = None) -> tuple[int, list]:
        """Record the reversals in the ITC claim register (only bills whose ITC is recorded as claimed there)."""
        from ..importers.gstr3b import Gstr3bValues
        from .itc import ItcService

        itc, done, skipped = ItcService(self.db), 0, []
        for b in bills:
            if b.status != "Reverse ITC" or b.itc_to_reverse <= 0:
                continue
            h = b.share(b.unpaid_at_day180)
            try:
                itc.reverse(client_id, b.invoice_id, b.reversal_period, user,
                            f"Not paid within {DAYS} days of the invoice date (unpaid {b.unpaid_at_day180})",
                            "Rule 37 CGST Rules", Gstr3bValues(h["igst"], h["cgst"], h["sgst"], h["cess"]))
                done += 1
            except ServiceError as exc:
                skipped.append((b.invoice_no, str(exc)))
        return done, skipped

    @staticmethod
    def export_xlsx(bills: list[BillStatus], path: str | Path, title: str = "") -> Path:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        wb = Workbook()
        ws = wb.active
        ws.title = "180-day rule"
        ws.append([title or "180-day payment rule (Rule 37)"])
        ws["A1"].font = Font(bold=True, size=13)
        heads = ["Status", "Supplier", "GSTIN", "Invoice no", "Invoice date", "Bill value", "ITC", "Day 180",
                 "Unpaid on day 180", "Unpaid now", "ITC to reverse", "Reverse in", "Paid in full on",
                 "ITC to re-claim", "Re-claim in", "Interest (est.)"]
        ws.append(heads)
        for c in ws[2]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1F3864")
        for b in bills:
            ws.append([b.status, b.supplier, b.supplier_gstin, b.invoice_no, b.invoice_date, float(b.value),
                       float(b.itc), b.day180, float(b.unpaid_at_day180), float(b.unpaid_now),
                       float(b.itc_to_reverse), b.reversal_period, b.paid_in_full_on, float(b.itc_to_reclaim),
                       b.reclaim_period, float(b.interest)])
        for i, w in enumerate((20, 30, 17, 16, 12, 13, 12, 12, 14, 13, 13, 11, 13, 13, 11, 13), start=1):
            ws.column_dimensions[ws.cell(2, i).column_letter].width = w
        for row in ws.iter_rows(min_row=3):
            for c in row:
                if isinstance(c.value, date):
                    c.number_format = "DD-MM-YYYY"
        ws.freeze_panes = "A3"
        out = Path(path)
        wb.save(out)
        return out


def _allocate(bills: list[BillStatus], payments: list) -> None:
    """Payments naming a bill (or marked against it) settle it first; the rest settle the oldest open bills."""
    owed = {id(b): b.value for b in bills}
    loose = []
    for when, amount, ref, inv_id in payments:
        target = next((b for b in bills if (inv_id and b.invoice_id == inv_id) or
                       (ref and normalize_invoice_no(ref) == normalize_invoice_no(b.invoice_no))), None)
        if target is not None and owed[id(target)] > 0:
            use = min(amount, owed[id(target)])
            target.payments.append((when, use))
            owed[id(target)] -= use
            amount -= use
        if amount > 0:
            loose.append((when, amount))
    for when, amount in loose:
        for b in bills:
            if amount <= 0:
                break
            if owed[id(b)] > 0:
                use = min(amount, owed[id(b)])
                b.payments.append((when, use))
                owed[id(b)] -= use
                amount -= use


def _assess(b: BillStatus, as_of: date) -> None:
    b.day180 = b.invoice_date + timedelta(days=DAYS)
    paid_by_due = sum((a for d, a in b.payments if d <= b.day180), ZERO)
    paid_now = sum((a for d, a in b.payments if d <= as_of), ZERO)
    b.unpaid_now = max(b.value - paid_now, ZERO)
    full = ZERO
    for d, a in sorted(b.payments):
        full += a
        if b.value - full <= TOLERANCE:
            b.paid_in_full_on = d
            break
    if as_of <= b.day180:
        b.status = "Paid in time" if b.unpaid_now <= TOLERANCE else "Not yet due"
        return
    b.unpaid_at_day180 = max(b.value - paid_by_due, ZERO)
    if b.unpaid_at_day180 <= TOLERANCE:
        b.unpaid_at_day180 = ZERO
        b.status = "Paid in time"
        return
    f = b.unpaid_at_day180 / b.value if b.value else ZERO
    b.itc_to_reverse = q2(b.itc * f)
    b.reversal_period = _next_month(f"{b.day180:%Y-%m}")
    reversal_due = _gstr3b_due(b.reversal_period)
    late_paid = sum((a for d, a in b.payments if b.day180 < d <= as_of), ZERO)
    if late_paid > 0:
        b.itc_to_reclaim = q2(b.itc * min(late_paid, b.unpaid_at_day180) / b.value)
        b.reclaim_period = f"{max(d for d, a in b.payments if d > b.day180):%Y-%m}"
    b.status = "Paid late – re-claim" if b.unpaid_now <= TOLERANCE else "Reverse ITC"
    end = min(as_of, reversal_due, b.paid_in_full_on or reversal_due)
    days = max((end - _gstr3b_due(b.booked_period)).days, 0)
    b.interest = q2(b.itc_to_reverse * RATE * days / 365)
