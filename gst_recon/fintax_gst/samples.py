"""Demo datasets in realistic layouts, with the expected outcome of every scenario.

Generates, for client "Shree Ganesh Traders" (27AAPFU0939F1ZV), April 2025:

* ``tally_purchase_register_2025-04.xlsx`` – TallyPrime columnar purchase register
  (title rows, ledger-wise tax columns, Debit Note vouchers, a total row,
  plus an unregistered purchase and a bad date that must be rejected)
* ``tally_purchase_register_2025-03.xlsx`` – March books (prior-period search)
* ``custom_purchase_register_2025-04.csv`` – the same April books in a custom layout
* ``gstr2b_2025-04.xlsx`` – GST-portal style GSTR-2B Excel (B2B, B2BA, B2B-CDNR)
* ``gstr2b_2025-04.json`` – the same GSTR-2B as portal JSON
* ``EXPECTED_RESULTS.md`` – scenario list with the expected category
"""
from __future__ import annotations

import csv
import json
from dataclasses import dataclass
from datetime import date
from decimal import Decimal as D
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

from .core.gstin import check_digit
from .core.models import Category


def gstin(state: str, pan: str, entity: str = "1") -> str:
    base = f"{state}{pan}{entity}Z"
    return base + check_digit(base)


CLIENT = {
    "name": "Shree Ganesh Traders", "trade_name": "Shree Ganesh Traders", "gstin": "27AAPFU0939F1ZV",
    "financial_year": "2025-26", "registration_type": "Regular", "return_frequency": "Monthly",
    "address": "Shop 12, Market Yard, Pune 411037", "contact_person": "R. Kulkarni",
    "email": "accounts@example.com", "mobile": "9800000000",
}

SUPPLIERS = {
    "MSL": (gstin("27", "AABCM1234K"), "Mahalaxmi Steel Pvt Ltd"),
    "MSK": (gstin("29", "AABCM1234K"), "Mahalaxmi Steel Pvt Ltd (Karnataka)"),
    "BEL": (gstin("29", "AACCB5678L"), "Bangalore Electronics LLP"),
    "PPK": (gstin("27", "AAFFP4321M"), "Pune Packaging Co"),
    "GCH": (gstin("24", "AADCG8765N"), "Gujarat Chemicals Ltd"),
    "OSL": (gstin("27", "ABMPS2468P"), "Om Sai Logistics"),
    "DOS": (gstin("07", "AAGFD1357Q"), "Delhi Office Supplies"),
    "KTX": (gstin("33", "AAHCK9753R"), "Kovai Textiles Pvt Ltd"),
    "ANC": (gstin("27", "AAJFA2468S"), "Annapurna Caterers"),
    "SRW": (gstin("27", "ABKPS1357T"), "Shree Roadways (GTA)"),
}


@dataclass
class Doc:
    sup: str
    no: str
    dt: date
    taxable: str
    igst: str = "0"
    cgst: str = "0"
    sgst: str = "0"
    cess: str = "0"
    doc: str = "INV"  # supplier view: INV / CRN / DBN
    rate: int = 18
    itc: str = "Yes"
    reason: str = ""
    original_no: str = ""
    original_dt: date | None = None
    narration: str = ""
    rcm: bool = False


def intra(sup, no, dt, taxable, rate=18, **kw):
    half = (D(taxable) * D(rate) / D(200)).quantize(D("0.01"))
    return Doc(sup, no, dt, taxable, cgst=str(half), sgst=str(half), rate=rate, **kw)


def inter(sup, no, dt, taxable, rate=18, **kw):
    tax = (D(taxable) * D(rate) / D(100)).quantize(D("0.01"))
    return Doc(sup, no, dt, taxable, igst=str(tax), rate=rate, **kw)


def A(day):
    return date(2025, 4, day)


# (scenario, books documents, GSTR-2B documents, expected category, note)
SCENARIOS: list[tuple[str, list[Doc], list[Doc], Category, str]] = []


def scenario(name, books, g2b, expected, note):
    SCENARIOS.append((name, books, g2b, expected, note))


for i, (sup, rate) in enumerate([("MSL", 18), ("PPK", 12), ("OSL", 18), ("MSL", 18), ("PPK", 18)], start=1):
    d = intra(sup, f"{sup}/{100 + i}/25-26", A(i + 1), f"{25000 + i * 1750}", rate)
    scenario(f"Perfect match {i}", [d], [d], Category.FULLY_MATCHED, "All fields agree")
for i, sup in enumerate(["BEL", "GCH", "KTX"], start=1):
    d = inter(sup, f"{sup}-{2000 + i}", A(i + 8), f"{48000 + i * 3300}")
    scenario(f"Perfect match inter-state {i}", [d], [d], Category.FULLY_MATCHED, "IGST invoice")

scenario("Invoice number format", [intra("MSL", "MSL/210/25-26", A(10), "18500")],
         [intra("MSL", "MSL-210-25-26", A(10), "18500")], Category.NUMBER_VARIATION, "Slash vs hyphen")
scenario("Leading zeros", [intra("PPK", "PP-7", A(11), "9200", 12)],
         [intra("PPK", "PP-0007", A(11), "9200", 12)], Category.NUMBER_VARIATION, "PP-7 vs PP-0007")
scenario("Date difference", [inter("BEL", "BEL-2101", A(12), "64000")],
         [inter("BEL", "BEL-2101", A(14), "64000")], Category.DATE_DIFFERENCE, "2 days apart")
scenario("Taxable value difference",
         [Doc("OSL", "OSL/310", A(13), "50000", cgst="4500", sgst="4500")],
         [Doc("OSL", "OSL/310", A(13), "50500", cgst="4500", sgst="4500")],
         Category.TAXABLE_DIFFERENCE, "Taxable differs by 500, tax agrees")
scenario("IGST difference", [Doc("GCH", "GCH-2201", A(14), "100000", igst="18000")],
         [Doc("GCH", "GCH-2201", A(14), "100000", igst="12000")], Category.IGST_DIFFERENCE, "18% booked, 12% reported")
scenario("CGST/SGST difference", [Doc("MSL", "MSL/233/25-26", A(15), "40000", cgst="3600", sgst="3600")],
         [Doc("MSL", "MSL/233/25-26", A(15), "40000", cgst="3400", sgst="3400")],
         Category.CGST_SGST_DIFFERENCE, "CGST and SGST each lower by 200")
scenario("Multiple differences", [intra("PPK", "PPK/480", A(16), "30000")],
         [intra("PPK", "PPK/480", A(18), "28000")], Category.MULTIPLE_DIFFERENCES, "Date and values differ")
scenario("Rounding within tolerance", [Doc("OSL", "OSL/322", A(17), "12345.60", cgst="1111.10", sgst="1111.10")],
         [Doc("OSL", "OSL/322", A(17), "12345.60", cgst="1111.60", sgst="1110.60")],
         Category.FULLY_MATCHED, "50 paise differences are within ₹1 tolerance")
scenario("Missing in GSTR-2B 1", [intra("MSL", "MSL/260/25-26", A(25), "22000")], [], Category.MISSING_IN_2B,
         "Supplier has not filed")
scenario("Missing in GSTR-2B 2", [inter("KTX", "KTX-2290", A(27), "75500", 12)], [], Category.MISSING_IN_2B,
         "Supplier has not filed")
scenario("Missing in books", [], [intra("OSL", "OSL/377", A(26), "8800")], Category.MISSING_IN_BOOKS,
         "Invoice not booked")
scenario("Missing in books (dated March)", [], [inter("BEL", "BEL-1950", date(2025, 3, 20), "15000")],
         Category.MISSING_IN_BOOKS, "Old invoice reported late; not in any books")
scenario("Duplicate in books", [intra("PPK", "PPK/455", A(19), "14000"), intra("PPK", "PPK/455", A(19), "14000")],
         [intra("PPK", "PPK/455", A(19), "14000")], Category.DUPLICATE_IN_BOOKS, "Entered twice in Tally")
scenario("Credit note matched", [inter("GCH", "GCH-CN-12", A(20), "5000", doc="CRN")],
         [inter("GCH", "GCH-CN-12", A(20), "5000", doc="CRN")], Category.FULLY_MATCHED,
         "Purchase return (Debit Note voucher in Tally)")
scenario("Credit note difference", [intra("MSL", "MSL-CN-07", A(21), "6000", doc="CRN")],
         [intra("MSL", "MSL-CN-07", A(21), "6500", doc="CRN")], Category.CDN_DIFFERENCE, "Credit note value differs")
scenario("Amended invoice", [intra("OSL", "OSL/45", A(5), "21000")],
         [intra("OSL", "OSL/45A", A(5), "21000", original_no="OSL/45", original_dt=A(5))], Category.AMENDED,
         "Supplier amended invoice number in B2BA")
scenario("Probable match (typo)", [inter("DOS", "DOS-1189", A(22), "11800")],
         [inter("DOS", "DOS-1198", A(22), "11800")], Category.PROBABLE, "Digits transposed – review")
scenario("GSTIN mismatch", [inter("MSK", "MSK-5501", A(23), "42000")],
         [intra("MSL", "MSK-5501", A(23), "42000")], Category.PROBABLE,
         "Booked under supplier's Karnataka GSTIN; supplier reported from Maharashtra registration")
scenario("ITC not available", [inter("KTX", "KTX-2301", A(24), "36000")],
         [inter("KTX", "KTX-2301", A(24), "36000", itc="No", reason="POS and supplier state are same but recipient state is different")],
         Category.FULLY_MATCHED, "Matched, but GSTR-2B marks ITC unavailable – at risk")

scenario("Blocked-credit indicator", [intra("ANC", "ANC/77", A(18), "8000", narration="Staff lunch – Diwali function")],
         [intra("ANC", "ANC/77", A(18), "8000")], Category.FULLY_MATCHED,
         "Matched; ITC review flags Section 17(5)(b)(i) food / gifts for the reviewer")
scenario("Reverse charge (GTA)", [Doc("SRW", "SRW-310", A(8), "20000", cgst="500", sgst="500", rate=5, rcm=True,
                                      narration="Freight inward – GTA under RCM")],
         [Doc("SRW", "SRW-310", A(8), "20000", cgst="500", sgst="500", rate=5, rcm=True)], Category.FULLY_MATCHED,
         "Matched; ITC available only after RCM tax is paid (GSTR-3B 4(A)(3))")

PRIOR_BOOKS = [inter("DOS", "DOS-1001", date(2025, 3, 28), "9500")]
PRIOR_IN_2B = [inter("DOS", "DOS-1001", date(2025, 3, 28), "9500")]
scenario("Previous-period invoice", [], PRIOR_IN_2B, Category.PREVIOUS_PERIOD,
         "Booked in March 2025, reflected in April GSTR-2B")
PRIOR_ONLY_BOOKS = [intra("MSL", "MSL/090/24-25", date(2025, 3, 15), "17000")]  # matched in March already

REJECTED_BOOK_ROWS = [
    {"name": "Local Hardware (URD)", "gstin": "", "no": "LH-55", "dt": "12-04-2025", "taxable": "2500",
     "why": "Unregistered supplier – no GSTIN"},
    {"name": "Pune Packaging Co", "gstin": SUPPLIERS["PPK"][0], "no": "PPK/499", "dt": "31-04-2025",
     "taxable": "4000", "why": "Invalid date 31-04-2025"},
]


def books_docs() -> list[Doc]:
    return [d for _n, b, _g, _c, _t in SCENARIOS for d in b]


def g2b_docs() -> list[Doc]:
    return [d for _n, _b, g, _c, _t in SCENARIOS for d in g]


# ---------------------------------------------------------------- writers
TALLY_HEADERS = [
    "Date", "Particulars", "Voucher Type", "Voucher No.", "Supplier Invoice No.", "Supplier Invoice Date",
    "GSTIN/UIN", "Gross Total", "Purchase @18%", "Purchase @12%", "Input CGST 9%", "Input SGST 9%",
    "Input CGST 6%", "Input SGST 6%", "Input IGST 18%", "Input IGST 12%", "Narration",
]


def _tally_row(n: int, d: Doc, booking: date) -> list:
    sign = -1 if d.doc == "CRN" else 1
    vtype = {"INV": "Purchase", "CRN": "Debit Note", "DBN": "Credit Note"}[d.doc]
    tx, ig, cg, sg = (D(x) * sign for x in (d.taxable, d.igst, d.cgst, d.sgst))
    gross = tx + ig + cg + sg
    r18 = d.rate == 18
    gst, name = SUPPLIERS[d.sup]
    return [
        booking, name, vtype, n, d.no, d.dt, gst, gross,
        tx if r18 else None, None if r18 else tx,
        cg if r18 and cg else None, sg if r18 and sg else None,
        cg if not r18 and cg else None, sg if not r18 and sg else None,
        ig if r18 and ig else None, ig if not r18 and ig else None, d.narration or None,
    ]


def write_tally_register(path: Path, docs: list[Doc], period_label: str, booking_month: tuple[int, int],
                         include_rejects: bool) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Purchase Register"
    ws.append([CLIENT["name"]])
    ws.append(["Purchase Register"])
    ws.append([period_label])
    ws.append([])
    ws.append(TALLY_HEADERS)
    for c in ws[5]:
        c.font = Font(bold=True)
    n = 0
    for d in docs:
        n += 1
        booking = date(booking_month[0], booking_month[1], min(max(d.dt.day, 1), 28)) \
            if (d.dt.year, d.dt.month) != booking_month else d.dt
        ws.append(_tally_row(n, d, booking))
    if include_rejects:
        for rj in REJECTED_BOOK_ROWS:
            n += 1
            tx = D(rj["taxable"])
            ws.append([date(*booking_month, 15), rj["name"], "Purchase", n, rj["no"], rj["dt"], rj["gstin"],
                       tx * D("1.18"), tx, None, tx * D("0.09"), tx * D("0.09"), None, None, None, None])
    last = ws.max_row
    ws.append(["Grand Total", None, None, None, None, None, None, f"=SUM(H6:H{last})"])
    for row in ws.iter_rows(min_row=6):
        for c in row:
            if isinstance(c.value, D):
                c.value = float(c.value)  # Excel stores numbers as binary floats, as Tally exports do
            if isinstance(c.value, date):
                c.number_format = "DD-MMM-YYYY"
    wb.save(path)


def write_custom_csv(path: Path, docs: list[Doc]) -> None:
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["Vendor Name", "Vendor GSTIN", "Bill No", "Bill Date", "Document Type", "Taxable Amount",
                    "IGST Amount", "CGST Amount", "SGST Amount", "Cess Amount", "Total Amount"])
        for d in docs:
            gst, name = SUPPLIERS[d.sup]
            total = sum(D(x) for x in (d.taxable, d.igst, d.cgst, d.sgst, d.cess))
            dtype = {"INV": "Invoice", "CRN": "Credit Note", "DBN": "Debit Note"}[d.doc]
            w.writerow([name, gst, d.no, d.dt.strftime("%d/%m/%Y"), dtype, f"{D(d.taxable):,.2f}", d.igst, d.cgst,
                        d.sgst, d.cess, f"{total:.2f}"])
        for rj in REJECTED_BOOK_ROWS:
            w.writerow([rj["name"], rj["gstin"], rj["no"], rj["dt"].replace("-", "/"), "Invoice", rj["taxable"], 0,
                        D(rj["taxable"]) * D("0.09"), D(rj["taxable"]) * D("0.09"), 0, ""])


B2B_TOP = ["GSTIN of supplier", "Trade/Legal name", "Invoice details", None, None, None, "Place of supply",
           "Supply Attract Reverse Charge", "Taxable Value (₹)", "Tax Amount", None, None, None,
           "GSTR-1/IFF/GSTR-5 Period", "GSTR-1/IFF/GSTR-5 Filing Date", "ITC Availability", "Reason",
           "Applicable % of Tax Rate", "Source", "IRN", "IRN Date"]
B2B_SUB = [None, None, "Invoice number", "Invoice type", "Invoice Date", "Invoice Value(₹)", None, None, None,
           "Integrated Tax(₹)", "Central Tax(₹)", "State/UT Tax(₹)", "Cess(₹)"]
CDNR_TOP = ["GSTIN of supplier", "Trade/Legal name", "Credit note/Debit note details", None, None, None, None,
            "Place of supply", "Supply Attract Reverse Charge", "Taxable Value (₹)", "Tax Amount", None, None, None,
            "GSTR-1/IFF/GSTR-5 Period", "GSTR-1/IFF/GSTR-5 Filing Date", "ITC Availability", "Reason",
            "Applicable % of Tax Rate", "Source", "IRN", "IRN Date"]
CDNR_SUB = [None, None, "Note number", "Note type", "Note Supply type", "Note date", "Note Value (₹)", None, None,
            None, "Integrated Tax(₹)", "Central Tax(₹)", "State/UT Tax(₹)", "Cess(₹)"]
B2BA_TOP = ["Original Details", None, "GSTIN of supplier", "Trade/Legal name", "Invoice details", None, None, None,
            "Place of supply", "Supply Attract Reverse Charge", "Taxable Value (₹)", "Tax Amount", None, None, None,
            "GSTR-1/IFF/GSTR-5 Period", "GSTR-1/IFF/GSTR-5 Filing Date", "ITC Availability", "Reason"]
B2BA_SUB = ["Invoice number", "Invoice Date", None, None, "Invoice number", "Invoice type", "Invoice Date",
            "Invoice Value(₹)", None, None, None, "Integrated Tax(₹)", "Central Tax(₹)", "State/UT Tax(₹)", "Cess(₹)"]


def _value(d: Doc) -> D:
    return sum(D(x) for x in (d.taxable, d.igst, d.cgst, d.sgst, d.cess))


def _pos(d: Doc) -> str:
    return "27-Maharashtra"


def write_gstr2b_excel(path: Path, docs: list[Doc]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Read me"
    ws.append(["Goods and Services Tax - GSTR-2B"])
    ws.append(["Financial Year", "2025-26"])
    ws.append(["Tax Period", "April"])
    ws.append(["GSTIN", CLIENT["gstin"]])
    ws.append(["Legal Name", CLIENT["name"]])
    summary = wb.create_sheet("ITC Available")
    summary.append(["Summary of ITC available (computed on the portal)"])

    def sheet(title, top, sub, caption):
        s = wb.create_sheet(title)
        s.append([caption])
        s.append([])
        s.append([])
        s.append([])
        s.append(top)
        s.append(sub)
        return s

    b2b = sheet("B2B", B2B_TOP, B2B_SUB, "Taxable inward supplies received from registered persons")
    b2ba = sheet("B2BA", B2BA_TOP, B2BA_SUB, "Amendments to previously filed invoices by supplier")
    cdnr = sheet("B2B-CDNR", CDNR_TOP, CDNR_SUB, "Debit/Credit notes (Original)")
    for d in docs:
        gst, name = SUPPLIERS[d.sup]
        common_tail = [_pos(d), "Yes" if d.rcm else "No", float(D(d.taxable))]
        taxes = [float(D(x)) for x in (d.igst, d.cgst, d.sgst, d.cess)]
        filing = ["Apr'25" if d.dt.month == 4 else "Mar'25", "11-05-2025", d.itc, d.reason]
        if d.doc in ("CRN", "DBN"):
            cdnr.append([gst, name, d.no, "Credit Note" if d.doc == "CRN" else "Debit Note", "Regular",
                         d.dt.strftime("%d-%m-%Y"), float(_value(d)), *common_tail, *taxes, *filing, None, "", "",
                         ""])
        elif d.original_no:
            b2ba.append([d.original_no, d.original_dt.strftime("%d-%m-%Y"), gst, name, d.no, "Regular",
                         d.dt.strftime("%d-%m-%Y"), float(_value(d)), *common_tail, *taxes, *filing])
        else:
            b2b.append([gst, name, d.no, "Regular", d.dt.strftime("%d-%m-%Y"), float(_value(d)), *common_tail,
                        *taxes, *filing, None, "e-Invoice", "", ""])
    wb.save(path)


def write_gstr2b_json(path: Path, docs: list[Doc]) -> None:
    sections: dict[str, dict[str, list]] = {"b2b": {}, "b2ba": {}, "cdnr": {}}
    for d in docs:
        gst, name = SUPPLIERS[d.sup]
        amounts = {"txval": float(D(d.taxable)), "igst": float(D(d.igst)), "cgst": float(D(d.cgst)),
                   "sgst": float(D(d.sgst)), "cess": float(D(d.cess))}
        base = {"dt": d.dt.strftime("%d-%m-%Y"), "val": float(_value(d)), "pos": "27", "rev": "Y" if d.rcm else "N",
                "itcavl": "Y" if d.itc == "Yes" else "N", "rsn": "P" if d.reason else "", "diffprcnt": 1,
                **amounts}
        if d.doc in ("CRN", "DBN"):
            sec, key = "cdnr", "nt"
            doc = {"ntnum": d.no, "typ": "C" if d.doc == "CRN" else "D", "suptyp": "R", **base}
        elif d.original_no:
            sec, key = "b2ba", "inv"
            doc = {"oinum": d.original_no, "oidt": d.original_dt.strftime("%d-%m-%Y"), "inum": d.no, "typ": "R",
                   **base}
        else:
            sec, key = "b2b", "inv"
            doc = {"inum": d.no, "typ": "R", **base}
        entry = sections[sec].setdefault(gst, {"ctin": gst, "trdnm": name, "supfildt": "11-05-2025",
                                                "supprd": "042025", key: []})
        entry[key].append(doc)
    docdata = {sec: list(v.values()) for sec, v in sections.items()}
    docdata["isd"] = [{"ctin": ISD_GSTIN, "trdnm": "Shree Ganesh Traders – HO (ISD)", "doclist": [
        {"doctyp": "ISDI", "docnum": "ISD/25/041", "docdt": "15-04-2025", "oinvnum": "", "oinvdt": "",
         "igst": 0, "cgst": 900, "sgst": 900, "cess": 0, "itcelg": "Y"}]}]
    docdata["impg"] = [{"refdt": "20-04-2025", "portcode": "INNSA1", "boenum": "7012345", "boedt": "18-04-2025",
                        "isamd": "N", "txval": 10000, "igst": 1800, "cess": 0}]
    payload = {"chksum": "demo", "data": {
        "gstin": CLIENT["gstin"], "rtnprd": "042025", "version": "1.0", "gendt": "14-05-2025",
        "docdata": docdata,
    }}
    path.write_text(json.dumps(payload, indent=2), encoding="utf-8")


ISD_GSTIN = gstin("27", "AAPFU0939F", "2")  # the client's own Input Service Distributor registration


def _components(docs: list[Doc]) -> dict[str, D]:
    tot = {k: D(0) for k in ("igst", "cgst", "sgst", "cess")}
    for d in docs:
        sign = -1 if d.doc == "CRN" else 1
        for k in tot:
            tot[k] += D(getattr(d, k)) * sign
    return tot


SAMPLE_INTEREST = 120  # interest paid in cash in April's GSTR-3B Table 6.1


def write_gstr3b_json(path: Path) -> None:
    """April 2025 GSTR-3B as filed: 4(A)(5) includes ₹2,000 more than GSTR-2B supports, so the
    Table 4 working shows a possible statutory discrepancy."""
    docs = g2b_docs()
    oth = _components([d for d in docs if d.itc == "Yes" and not d.rcm])
    oth["cgst"] += D(1000)
    oth["sgst"] += D(1000)
    rcm = _components([d for d in docs if d.rcm])
    na = _components([d for d in docs if d.itc != "Yes"])

    def row(ty, v):
        return {"ty": ty, "iamt": float(v.get("igst", 0)), "camt": float(v.get("cgst", 0)),
                "samt": float(v.get("sgst", 0)), "csamt": float(v.get("cess", 0))}

    impg = {"igst": D(1800)}
    isd = {"cgst": D(900), "sgst": D(900)}
    rev = {"cgst": D(720), "sgst": D(720)}  # 17(5) food reversal (ANC/77)
    net = {k: oth.get(k, 0) + rcm.get(k, 0) + impg.get(k, 0) + isd.get(k, 0) - rev.get(k, 0)
           for k in ("igst", "cgst", "sgst", "cess")}
    from .samples_outward import gstr3b_31

    sup = gstr3b_31()
    due = {k: sum(D(str(sup[t].get(k, 0))) for t in ("osup_det", "osup_zero")) for k in ("iamt", "camt", "samt")}
    # Table 6.1: IGST and CGST/SGST liability set off from ITC (same head), ₹120 interest paid in cash
    tx_pmt = {"pditc": {"i_pdi": float(due["iamt"]), "c_pdc": float(due["camt"]), "s_pds": float(due["samt"]),
                        "i_pdc": 0, "i_pds": 0, "c_pdi": 0, "s_pdi": 0, "cs_pdcs": 0},
              "pdcash": [{"ipd": 0, "cpd": 0, "spd": 0, "cspd": 0, "c_intrpd": SAMPLE_INTEREST / 2,
                          "s_intrpd": SAMPLE_INTEREST / 2}]}
    payload = {"gstin": CLIENT["gstin"], "ret_period": "042025", "sup_details": sup, "tx_pmt": tx_pmt, "itc_elg": {
        "itc_avl": [row("IMPG", impg), row("IMPS", {}), row("ISRC", rcm), row("ISD", isd), row("OTH", oth)],
        "itc_rev": [row("RUL", rev), row("OTH", {})],
        "itc_net": {k: v for k, v in row("", net).items() if k != "ty"},
        "itc_inelg": [row("RUL", {}), row("OTH", na)],
    }}
    path.write_text(json.dumps(payload, indent=2), encoding="utf-8")


def write_gstr3b_template(path: Path) -> None:
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["Period", "Table", "Description", "IGST", "CGST", "SGST", "Cess"])
        w.writerow(["2025-03", "4A5", "All other ITC", "1710.00", "1530.00", "1530.00", "0"])
        w.writerow(["2025-03", "4B2", "Others", "0", "0", "0", "0"])


IMS_ACTIONS_DEMO = {"GCH-2201": "R", "DOS-1198": "N"}


def write_ims_json(path: Path) -> None:
    """IMS actions for April 2025: one rejection (still in GSTR-2B), one pending (explains a
    document missing from GSTR-2B) and the rest accepted / no action."""
    sections: dict[str, list] = {"b2b": [], "cn": []}

    def entry(d: Doc, action: str) -> dict:
        gst, name = SUPPLIERS[d.sup]
        key = "ntnum" if d.doc == "CRN" else "inum"
        return {"stin": gst, "trdnm": name, key: d.no, "idt" if d.doc != "CRN" else "ntdt": d.dt.strftime("%d-%m-%Y"),
                "val": float(_value(d)), "txval": float(D(d.taxable)), "iamt": float(D(d.igst)),
                "camt": float(D(d.cgst)), "samt": float(D(d.sgst)), "csamt": float(D(d.cess)), "pos": "27",
                "action": action, "rtnprd": "042025"}

    for d in g2b_docs():
        if d.original_no or d.dt.month != 4:
            continue  # amendments and the late March invoice are not part of this IMS extract
        sections["cn" if d.doc == "CRN" else "b2b"].append(entry(d, IMS_ACTIONS_DEMO.get(d.no, "A")))
    pending = next(d for d in books_docs() if d.no == "KTX-2290")
    sections["b2b"].append(entry(pending, "P"))
    path.write_text(json.dumps({"data": {"gstin": CLIENT["gstin"], "rtnprd": "042025", "invdata": sections}},
                               indent=2), encoding="utf-8")


#: GSTR-9 Table 8A extras: KTX-2290 (missing in April GSTR-2B) was filed late by the supplier and is in 8A;
#: OSL/412 (May) was never booked.
TABLE8A_LATE = "KTX-2290"
TABLE8A_NOT_BOOKED = intra("OSL", "OSL/412", date(2025, 5, 14), "12400")


def table8a_docs() -> list[Doc]:
    late = [d for d in books_docs() if d.no == TABLE8A_LATE]
    in_fy = [d for d in g2b_docs() if not d.rcm and d.dt >= date(2025, 4, 1)]  # 8A: documents dated in the FY
    return in_fy + late + [TABLE8A_NOT_BOOKED]


def write_gstr8a_json(path: Path, docs: list[Doc]) -> None:
    """GSTR-9 Table 8A document details as downloaded from the portal (FY 2025-26)."""
    sections: dict[str, dict[str, dict]] = {"b2b": {}, "b2ba": {}, "cdnr": {}}
    for d in docs:
        gst, name = SUPPLIERS[d.sup]
        item = {"num": 1, "itm_det": {"rt": d.rate, "txval": float(D(d.taxable)), "iamt": float(D(d.igst)),
                                      "camt": float(D(d.cgst)), "samt": float(D(d.sgst)), "csamt": float(D(d.cess))}}
        base = {"val": float(_value(d)), "pos": "27", "rchrg": "N", "itcavl": "Y" if d.itc == "Yes" else "N",
                "rsn": "P" if d.reason else "", "itms": [item]}
        if d.doc in ("CRN", "DBN"):
            sec, key = "cdnr", "nt"
            doc = {"ntty": "C" if d.doc == "CRN" else "D", "nt_num": d.no, "nt_dt": d.dt.strftime("%d-%m-%Y"), **base}
        elif d.original_no:
            sec, key = "b2ba", "inv"
            doc = {"oinum": d.original_no, "oidt": d.original_dt.strftime("%d-%m-%Y"), "inum": d.no,
                   "idt": d.dt.strftime("%d-%m-%Y"), "inv_typ": "R", **base}
        else:
            sec, key = "b2b", "inv"
            doc = {"inum": d.no, "idt": d.dt.strftime("%d-%m-%Y"), "inv_typ": "R", **base}
        filed = "042025" if d.dt.month == 4 and d.no != TABLE8A_LATE else "062025"
        sections[sec].setdefault(gst, {"ctin": gst, "trdnm": name, "fldtr1": filed, key: []})[key].append(doc)
    payload = {"data": {"gstin": CLIENT["gstin"], "fy": "2025-26", "gendt": "15-11-2026",
                        **{sec: list(v.values()) for sec, v in sections.items()}}}
    path.write_text(json.dumps(payload, indent=2), encoding="utf-8")


#: Item-wise purchase register (HSN, unit, quantity): each of the first five April invoices is split into two
#: lines at the same tax rate but different HSN – they merge into one document, and feed GSTR-9 Table 18.
ITEMWISE_HSN = [("7208", "Flat-rolled steel", "KGS", D("0.6")), ("7318", "Screws and bolts", "KGS", D("0.4"))]


def write_itemwise_register(path: Path) -> None:
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["Supplier Name", "GSTIN", "Bill No", "Bill Date", "HSN/SAC", "UQC", "Quantity", "Taxable Value",
                    "IGST", "CGST", "SGST"])
        for d in books_docs()[:5]:
            gst, name = SUPPLIERS[d.sup]
            for code, _desc, uqc, share in ITEMWISE_HSN:
                part = lambda x: (D(x) * share).quantize(D("0.01"))  # noqa: E731
                w.writerow([name, gst, d.no, d.dt.strftime("%d-%m-%Y"), code, uqc, int(D(d.taxable) * share / 50),
                            part(d.taxable), part(d.igst), part(d.cgst), part(d.sgst)])


EXPENSE_LEDGER = [  # (ledger head, value, total ITC, eligible ITC) – FY 2025-26
    ("Purchase Accounts", "760000", "136800", "136800"), ("Freight Inward", "20000", "3600", "3600"),
    ("Electricity Charges", "48000", "0", "0"), ("Office Rent", "240000", "43200", "43200"),
    ("Staff Salaries", "900000", "0", "0"), ("Staff Welfare – Food", "15000", "750", "0"),
    ("Bank Charges", "4000", "720", "720"), ("Printing & Stationery", "12000", "2160", "2160"),
    ("Repairs to Machinery", "30000", "5400", "5400"), ("Computer (Fixed Asset)", "60000", "10800", "10800"),
    ("Audit Fees", "50000", "9000", "9000"),
]


def write_expense_ledger(path: Path) -> None:
    lines = ["Ledger,Amount,ITC,Eligible ITC"] + [f'"{h}",{v},{t},{e}' for h, v, t, e in EXPENSE_LEDGER]
    lines.append(f"Grand Total,{sum(D(v) for _h, v, _t, _e in EXPENSE_LEDGER)},,")
    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_all(out_dir: str | Path) -> dict[str, Path]:
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    files = {
        "books_april": out / "tally_purchase_register_2025-04.xlsx",
        "books_march": out / "tally_purchase_register_2025-03.xlsx",
        "books_april_csv": out / "custom_purchase_register_2025-04.csv",
        "gstr2b_april": out / "gstr2b_2025-04.xlsx",
        "gstr2b_april_json": out / "gstr2b_2025-04.json",
        "gstr3b_april_json": out / "gstr3b_2025-04.json",
        "gstr3b_template_csv": out / "gstr3b_template_2025-03.csv",
        "ims_april_json": out / "ims_2025-04.json",
        "gstr8a_json": out / "gstr9_table8a_2025-26.json",
        "expense_ledger": out / "expense_ledger_2025-26.csv",
        "books_itemwise_csv": out / "purchase_register_itemwise_2025-04.csv",
    }
    write_tally_register(files["books_april"], books_docs(), "1-Apr-2025 to 30-Apr-2025", (2025, 4), True)
    write_tally_register(files["books_march"], PRIOR_BOOKS + PRIOR_ONLY_BOOKS, "1-Mar-2025 to 31-Mar-2025",
                         (2025, 3), False)
    write_custom_csv(files["books_april_csv"], books_docs())
    write_gstr2b_excel(files["gstr2b_april"], g2b_docs())
    write_gstr2b_json(files["gstr2b_april_json"], g2b_docs())
    write_gstr3b_json(files["gstr3b_april_json"])
    write_gstr3b_template(files["gstr3b_template_csv"])
    write_ims_json(files["ims_april_json"])
    write_gstr8a_json(files["gstr8a_json"], table8a_docs())
    write_expense_ledger(files["expense_ledger"])
    write_itemwise_register(files["books_itemwise_csv"])
    from .samples_outward import write_outward

    files.update(write_outward(out))
    lines = ["# Demo dataset – expected reconciliation outcome", "",
             f"Client: {CLIENT['name']} ({CLIENT['gstin']}), FY 2025-26, period April 2025.", "",
             "| # | Scenario | Expected category | Note |", "|---|---|---|---|"]
    for i, (name, _b, _g, cat, note) in enumerate(SCENARIOS, start=1):
        lines.append(f"| {i} | {name} | {cat.value} | {note} |")
    lines += ["", "Rows that must be rejected at import (books):", ""]
    lines += [f"* {r['no']}: {r['why']}" for r in REJECTED_BOOK_ROWS]
    lines += ["", "## ITC module (Phase 3)", "",
              "* `gstr2b_2025-04.json` also carries an ISD credit (₹1,800) and a bill of entry (IGST ₹1,800).",
              "* `gstr3b_2025-04.json`: 4(A)(5) is ₹2,000 more than GSTR-2B supports → possible statutory "
              "discrepancy; 4(B)(1) reverses ₹1,440 for ANC/77 (Section 17(5)).",
              "* `ims_2025-04.json`: GCH-2201 rejected but present in GSTR-2B (recompute GSTR-2B); KTX-2290 kept "
              "pending (explains why it is missing in GSTR-2B); OSL/377 accepted but not in books.",
              "* ITC review flags ANC/77 (Section 17(5)(b)(i) food and (h) gifts indicators), SRW-310 (reverse "
              "charge), KTX-2301 (GSTR-2B ITC not available) and all books documents not in GSTR-2B "
              "(Section 16(2)(aa))."]
    lines += ["", "## Outward supplies (Phase 4)", "",
              "* `tally_sales_register_2025-04.xlsx` vs `gstr1_2025-04.json`: 5 perfect B2B matches, one invoice not "
              "in GSTR-1, one in GSTR-1 only, a taxable-value difference, a wrong customer GSTIN (probable), a "
              "credit note and an SEZ supply; B2C, export and exempt sales are compared at summary level.",
              "* GSTR-1 B2CS leaves out ₹1,000 taxable value; GSTR-3B 3.1(a) tax is ₹3,600 lower than GSTR-1 → "
              "possible short payment in the liability working."]
    lines += ["", "## GSTR-9 Table 8A (Phase 5)", "",
              "* `gstr9_table8a_2025-26.json` holds the April GSTR-2B documents (excluding reverse charge), "
              f"{TABLE8A_LATE} (missing in April GSTR-2B, filed late by the supplier) and "
              f"{TABLE8A_NOT_BOOKED.no} dated May 2025, which is not in the books.",
              "* Books vs Table 8A for FY 2025-26: KTX-2290 matches; OSL/412 and OSL/377 are in 8A but not in the "
              "books; KTX-2301 (ITC not available) stays out of the 8A total."]
    lines += ["", "## GSTR-9 and GSTR-9C (Phases 6–7)", "",
              "* `gstr1_2025-05.json`: a new B2B invoice (SGT/25-26/0021), a B2BA reducing SGT/25-26/0013 from "
              "66,000 to 60,000 (GSTR-9 4L), a B2CSA revising April's B2C total (listed for manual entry) and an "
              "advance of 50,000 received (Table 11A → GSTR-9 4F, GSTR-3B 3.1(a)).",
              "* `gstr3b_2025-04.json` Table 6.1: the whole April liability set off from ITC, and 120 of interest paid "
              "in cash (GSTR-9 Table 9).",
              "* GSTR-9C: 5A defaults to the sales-register turnover (761,000) until the audited figure is entered; 5Q "
              "is GSTR-9 5N."]
    lines += ["", "## GSTR-9 Tables 10–17 and GSTR-9C Table 14 (Phase 8)", "",
              "* `gstr1_2026-04.json` (next year): B2BA raising SGT/25-26/0005 from 65,000 to 70,000 and SGT/25-26/0099 "
              "of 30-03-2026 reported late → GSTR-9 2025-26 Table 10 = 13,000; SGT/26-27/0001 belongs to 2026-27.",
              "* GSTR-1 HSN summaries (Table 12): April includes the 1,000 B2C sale left out of B2CS, so Table 17 "
              "exceeds the GSTR-1 sections by 1,000.",
              "* `expense_ledger_2025-26.csv`: ledger heads placed in GSTR-9C Table 14 rows by keyword (salaries and "
              "staff welfare in 14H, repairs to machinery in 14M, the computer in 14O, audit fees in 14N)."]
    lines += ["", "## GSTR-9 Tables 15, 16 and 18 (Phase 9)", "",
              "* `purchase_register_itemwise_2025-04.csv`: the first five April invoices, each in two lines at the same "
              "rate with different HSN (7208, 7318) – merged into one document each; GSTR-9 Table 18 shows both HSN.",
              "* Tables 15 and 16 are entered (items such as 15E-I, 16A-T)."]
    lines += ["", "## e-Invoice and e-way bill (Phase 10)", "",
              "* `einvoice_report_2025-04.csv` vs the sales register: 9 documents match (B2B, credit note, SEZ, export); "
              "the IRN of SGT/25-26/0011 is cancelled although the sale is still in the register (client "
              "clarification); SGT/25-26/0014 has no IRN; SGT/25-26/0012 has an IRN but is not in the register. B2C "
              "and exempt sales are not e-invoiced and are left out.",
              "* `ewaybill_report_2025-04.csv`: 7 documents match (SGT/25-26/0002 through its regenerated e-way bill); "
              "SGT/25-26/0005 (₹76,700) has none; 12 documents below ₹50,000 or credit notes need none (no action); "
              "SGT/25-26/0018 has an e-way bill but is not in the register; a delivery challan and an inward e-way "
              "bill are not compared.",
              "* `tds_tcs_credit_2025-26.xlsx` (TDS and TCS credit received): Venkatesh Traders' TDS on ₹1,20,000 "
              "agrees with the sales to it (May credit not yet accepted); Pune Municipal Corporation's TDS has no sales "
              "in the register; Mehta Polymers (Gujarat) deducted CGST/SGST instead of IGST (rejected); the May TCS is "
              "1% of net value where 0.5% applies."]
    lines += ["", "## Interest and late fees (Phase 11)", "",
              "* `returns_filed_2025-26.csv`: April GSTR-1 filed 4 days late and GSTR-3B 8 days late → late fees "
              "₹200 and ₹400 at the highest cap (AATO not entered); no interest on the delay, as no tax was paid in "
              "cash (Table 6.1).",
              "* The April GSTR-1 tax exceeds GSTR-3B by ₹3,600 (CGST/SGST ₹1,800 each) → interest at 18% from 20-05-2025 "
              "until paid; ITC in GSTR-3B above GSTR-2B is listed with interest as an upper limit."]
    (out / "EXPECTED_RESULTS.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    return files
